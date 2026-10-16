"""问题三通信链路计算与独立验证。

本模块只负责通信物理层计算，不构造中继任务，也不运行 baseline/ALNS。
所有端点高度均为绝对海拔；FSPL 的频率单位为 MHz、距离单位为 km。
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import sys
from typing import Callable, Dict, FrozenSet, Iterable, Mapping, Optional, Tuple
import math

import numpy as np
from affine import Affine
from openpyxl import load_workbook
from pyproj import Geod
import rasterio
from rasterio.coords import BoundingBox


# Preserve historical module names used by trusted communication caches.
_PROJECT_ROOT = Path(__file__).resolve().parents[2]
if str(_PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(_PROJECT_ROOT))
import project_paths as paths
paths.bootstrap()
MODULE_DIR = Path(__file__).resolve().parent
PROJECT_DIR = paths.ROOT

COMMUNICATION_FILE = paths.raw_file("通信链路参数.xlsx")
NODE_FILE = paths.raw_file("调度中心与服务区.xlsx")
DEM_FILE = paths.dem_file()

GEOD = Geod(ellps="WGS84")
TOLERANCE = 1e-9

# 后续连续通信检查统一使用：1 s确定性细网格，状态切换细化至0.01 s。
COMMUNICATION_TIME_STEP_S = 1.0
STATE_TRANSITION_TOLERANCE_S = 0.01

# PREPARE_LOAD 明确不在连续通信约束内。
COMMUNICATION_REQUIRED_PHASES: FrozenSet[str] = frozenset(
    {"CLIMB", "CRUISE", "DESCENT", "HANDOVER"}
)

TRANSPORT_RADIO = "运输无人机"
RELAY_ACCESS_RADIO = "中继接入端"
RELAY_BACKHAUL_RADIO = "中继回传端"
GATEWAY_RADIO = "固定网关 G01"


@dataclass(frozen=True)
class Point3D:
    """WGS84经纬度与绝对海拔构成的通信端点。"""

    longitude_deg: float
    latitude_deg: float
    altitude_m: float


@dataclass(frozen=True)
class RadioProfile:
    name: str
    transmit_power_dbm: float
    antenna_gain_dbi: float


@dataclass(frozen=True)
class CommunicationParameters:
    frequency_mhz: float
    system_loss_db: float
    obstruction_loss_db: float
    receiver_sensitivity_dbm: float
    fade_margin_db: float
    gateway_antenna_height_agl_m: float
    radios: Mapping[str, RadioProfile]

    @property
    def required_received_power_dbm(self) -> float:
        """计入衰落裕量后的有效接收门限 Psens+M。"""
        return self.receiver_sensitivity_dbm + self.fade_margin_db


@dataclass(frozen=True)
class GatewayEndpoint:
    point: Point3D
    ground_elevation_m: float
    antenna_height_agl_m: float


@dataclass(frozen=True)
class DemData:
    elevation_m: np.ndarray
    transform: Affine
    bounds: BoundingBox
    crs: str


@dataclass(frozen=True)
class LineOfSightResult:
    obstructed: bool
    minimum_clearance_m: float
    checked_cell_count: int
    first_blocking_cell: Optional[Tuple[int, int]]


@dataclass(frozen=True)
class LinkEvaluation:
    distance_3d_m: float
    distance_3d_km: float
    fspl_db: float
    obstructed: bool
    obstruction_loss_db: float
    total_path_loss_db: float
    maximum_path_loss_db: float
    available: bool


def load_communication_parameters(
    path: Path = COMMUNICATION_FILE,
) -> CommunicationParameters:
    """从附件读取传播参数和四类通信端点参数。"""
    workbook = load_workbook(paths.require_file(path), data_only=True, read_only=True)
    worksheet = workbook["数据"]
    values: Dict[Tuple[str, str], float] = {}
    for row in worksheet.iter_rows(min_row=3, values_only=True):
        category, parameter_name, _, _, value = row[:5]
        if category is None or parameter_name is None or value is None:
            continue
        values[(str(category).strip(), str(parameter_name).strip())] = float(value)

    def parameter(category: str, name: str) -> float:
        return values[(category, name)]

    radios = {
        name: RadioProfile(
            name=name,
            transmit_power_dbm=parameter(name, "发射功率（dBm）"),
            antenna_gain_dbi=parameter(name, "天线增益（dBi）"),
        )
        for name in (
            TRANSPORT_RADIO,
            RELAY_ACCESS_RADIO,
            RELAY_BACKHAUL_RADIO,
            GATEWAY_RADIO,
        )
    }
    result = CommunicationParameters(
        frequency_mhz=parameter("传播参数", "载波频率（MHz）"),
        system_loss_db=parameter("传播参数", "系统损耗（dB）"),
        obstruction_loss_db=parameter("传播参数", "地形遮挡附加损耗（dB）"),
        receiver_sensitivity_dbm=parameter("接收参数", "接收灵敏度（dBm）"),
        fade_margin_db=parameter("接收参数", "衰落裕量（dB）"),
        gateway_antenna_height_agl_m=parameter(
            GATEWAY_RADIO, "天线离地高度（m）"
        ),
        radios=radios,
    )
    if result.frequency_mhz <= 0.0:
        raise ValueError("通信载波频率必须为正数")
    return result


def load_g01_endpoint(
    parameters: CommunicationParameters,
    node_file: Path = NODE_FILE,
) -> GatewayEndpoint:
    """G01使用O01的Excel地面海拔，禁止以DEM高程替代。"""
    workbook = load_workbook(paths.require_file(node_file), data_only=True, read_only=True)
    worksheet = workbook["数据"]
    for row in worksheet.iter_rows(values_only=True):
        node_id = str(row[0]).strip().upper() if row[0] is not None else ""
        if node_id != "O01":
            continue
        longitude = float(row[2])
        latitude = float(row[3])
        ground = float(row[4])
        antenna_altitude = ground + parameters.gateway_antenna_height_agl_m
        return GatewayEndpoint(
            point=Point3D(longitude, latitude, antenna_altitude),
            ground_elevation_m=ground,
            antenna_height_agl_m=parameters.gateway_antenna_height_agl_m,
        )
    raise ValueError("调度中心与服务区附件中未找到O01")


def load_dem(path: Path = DEM_FILE) -> DemData:
    """读取30 m DEM，并统一屏蔽无效高程占位值。"""
    with rasterio.open(paths.require_file(path)) as source:
        elevation = source.read(1).astype(np.float64)
        invalid = ~np.isfinite(elevation) | (elevation <= -10000.0)
        if source.nodata is not None:
            invalid |= elevation == float(source.nodata)
        elevation[invalid] = np.nan
        return DemData(
            elevation_m=elevation,
            transform=source.transform,
            bounds=source.bounds,
            crs=str(source.crs),
        )


def dem_elevation_at(dem: DemData, longitude_deg: float, latitude_deg: float) -> float:
    """读取坐标所在DEM像元的地面海拔。"""
    col_float, row_float = (~dem.transform) * (longitude_deg, latitude_deg)
    row = int(math.floor(row_float))
    col = int(math.floor(col_float))
    height, width = dem.elevation_m.shape
    if not (0 <= row < height and 0 <= col < width):
        raise ValueError("通信端点位于DEM覆盖范围之外")
    value = float(dem.elevation_m[row, col])
    if not math.isfinite(value):
        raise ValueError("通信端点所在DEM像元高程无效")
    return value


def horizontal_distance_m(first: Point3D, second: Point3D) -> float:
    """WGS84椭球上的水平测地距离，单位m。"""
    _, _, distance_m = GEOD.inv(
        first.longitude_deg,
        first.latitude_deg,
        second.longitude_deg,
        second.latitude_deg,
    )
    return float(abs(distance_m))


def distance_3d_m(first: Point3D, second: Point3D) -> float:
    """水平测地距离与绝对海拔差合成三维直线距离，单位m。"""
    horizontal_m = horizontal_distance_m(first, second)
    vertical_m = second.altitude_m - first.altitude_m
    return math.hypot(horizontal_m, vertical_m)


def distance_3d_km(first: Point3D, second: Point3D) -> float:
    return distance_3d_m(first, second) / 1000.0


def _axis_interval(
    origin: float, direction: float, lower: float, upper: float
) -> Optional[Tuple[float, float]]:
    if abs(direction) <= TOLERANCE:
        if lower - TOLERANCE <= origin <= upper + TOLERANCE:
            return -math.inf, math.inf
        return None
    first = (lower - origin) / direction
    second = (upper - origin) / direction
    return min(first, second), max(first, second)


def _line_cell_interval(
    col0: float,
    row0: float,
    col1: float,
    row1: float,
    row: int,
    col: int,
) -> Optional[Tuple[float, float]]:
    """返回二维线段位于指定DEM像元内的参数区间。"""
    col_interval = _axis_interval(col0, col1 - col0, col, col + 1.0)
    row_interval = _axis_interval(row0, row1 - row0, row, row + 1.0)
    if col_interval is None or row_interval is None:
        return None
    lower = max(0.0, col_interval[0], row_interval[0])
    upper = min(1.0, col_interval[1], row_interval[1])
    if upper < lower - TOLERANCE:
        return None
    return lower, upper


def _crossed_dem_cells(
    first: Point3D, second: Point3D, dem: DemData
) -> Tuple[Tuple[int, int], ...]:
    """按all-touched口径枚举二维视线穿过或接触的DEM像元。"""
    inverse = ~dem.transform
    col0, row0 = inverse * (first.longitude_deg, first.latitude_deg)
    col1, row1 = inverse * (second.longitude_deg, second.latitude_deg)
    height, width = dem.elevation_m.shape
    if not (
        0.0 <= row0 < height
        and 0.0 <= col0 < width
        and 0.0 <= row1 < height
        and 0.0 <= col1 < width
    ):
        raise ValueError("通信链路端点位于DEM覆盖范围之外")

    breakpoints = {0.0, 1.0}
    delta_col = col1 - col0
    delta_row = row1 - row0
    if abs(delta_col) > TOLERANCE:
        for boundary in range(
            math.floor(min(col0, col1)) + 1,
            math.floor(max(col0, col1)) + 1,
        ):
            value = (boundary - col0) / delta_col
            if TOLERANCE < value < 1.0 - TOLERANCE:
                breakpoints.add(float(value))
    if abs(delta_row) > TOLERANCE:
        for boundary in range(
            math.floor(min(row0, row1)) + 1,
            math.floor(max(row0, row1)) + 1,
        ):
            value = (boundary - row0) / delta_row
            if TOLERANCE < value < 1.0 - TOLERANCE:
                breakpoints.add(float(value))

    def touching_indices(value: float, limit: int) -> Tuple[int, ...]:
        indices = {int(math.floor(value))}
        nearest = int(round(value))
        if math.isclose(value, nearest, rel_tol=0.0, abs_tol=1e-10):
            indices.add(nearest - 1)
        return tuple(sorted(index for index in indices if 0 <= index < limit))

    ordered = sorted(breakpoints)
    cells = set()
    sample_parameters = set(ordered)
    sample_parameters.update(
        (lower + upper) / 2.0 for lower, upper in zip(ordered, ordered[1:])
    )
    for parameter in sample_parameters:
        col_value = col0 + delta_col * parameter
        row_value = row0 + delta_row * parameter
        for row in touching_indices(row_value, height):
            for col in touching_indices(col_value, width):
                if _line_cell_interval(col0, row0, col1, row1, row, col) is not None:
                    cells.add((row, col))
    return tuple(sorted(cells))


def analyze_line_of_sight(
    first: Point3D,
    second: Point3D,
    dem: DemData,
    clearance_tolerance_m: float = 1e-6,
) -> LineOfSightResult:
    """比较三维视线海拔与沿途每个DEM像元的地面海拔。

    对每个二维视线穿过的像元，先求线段在该像元内部的参数区间，
    再计算该区间内三维视线的最低绝对海拔。若像元地面高程达到或
    超过该视线高度，则判为遮挡。端点本身不作为中间地形遮挡点。
    """
    if (
        abs(first.longitude_deg - second.longitude_deg) <= TOLERANCE
        and abs(first.latitude_deg - second.latitude_deg) <= TOLERANCE
    ):
        return LineOfSightResult(False, math.inf, 0, None)

    inverse = ~dem.transform
    col0, row0 = inverse * (first.longitude_deg, first.latitude_deg)
    col1, row1 = inverse * (second.longitude_deg, second.latitude_deg)
    cells = _crossed_dem_cells(first, second, dem)
    altitude_delta = second.altitude_m - first.altitude_m
    minimum_clearance = math.inf
    first_blocking_cell: Optional[Tuple[int, int]] = None
    endpoint_epsilon = 1e-9

    for row, col in cells:
        interval = _line_cell_interval(col0, row0, col1, row1, row, col)
        if interval is None:
            continue
        lower = max(interval[0], endpoint_epsilon)
        upper = min(interval[1], 1.0 - endpoint_epsilon)
        if upper < lower:
            continue
        ground_m = float(dem.elevation_m[row, col])
        if not math.isfinite(ground_m):
            raise ValueError(f"视线经过无效DEM像元: row={row}, col={col}")
        altitude_at_lower = first.altitude_m + altitude_delta * lower
        altitude_at_upper = first.altitude_m + altitude_delta * upper
        minimum_los_altitude = min(altitude_at_lower, altitude_at_upper)
        clearance = minimum_los_altitude - ground_m
        minimum_clearance = min(minimum_clearance, clearance)
        if clearance <= clearance_tolerance_m and first_blocking_cell is None:
            first_blocking_cell = (row, col)

    return LineOfSightResult(
        obstructed=first_blocking_cell is not None,
        minimum_clearance_m=minimum_clearance,
        checked_cell_count=len(cells),
        first_blocking_cell=first_blocking_cell,
    )


def terrain_obstruction(first: Point3D, second: Point3D, dem: DemData) -> int:
    """返回题目定义的地形遮挡变量b_ij。"""
    return int(analyze_line_of_sight(first, second, dem).obstructed)


def fspl_db(frequency_mhz: float, distance_km: float) -> float:
    """自由空间传播损耗；f必须为MHz，D必须为km。"""
    if frequency_mhz <= 0.0:
        raise ValueError("FSPL频率必须大于0 MHz")
    if distance_km < 0.0:
        raise ValueError("FSPL距离不能为负")
    if distance_km == 0.0:
        return -math.inf
    return 32.45 + 20.0 * math.log10(frequency_mhz) + 20.0 * math.log10(
        distance_km
    )


def total_path_loss_db(
    frequency_mhz: float,
    distance_km: float,
    obstructed: bool,
    obstruction_loss_db: float,
) -> float:
    return fspl_db(frequency_mhz, distance_km) + obstruction_loss_db * int(
        obstructed
    )


def one_way_maximum_path_loss_db(
    transmitter: RadioProfile,
    receiver: RadioProfile,
    parameters: CommunicationParameters,
) -> float:
    """Pt+Gt+Gr-Lsys-(Psens+M)。"""
    return (
        transmitter.transmit_power_dbm
        + transmitter.antenna_gain_dbi
        + receiver.antenna_gain_dbi
        - parameters.system_loss_db
        - parameters.required_received_power_dbm
    )


def bidirectional_maximum_path_loss_db(
    first: RadioProfile,
    second: RadioProfile,
    parameters: CommunicationParameters,
) -> float:
    """双向链路门限取两个通信方向允许损耗的较小值。"""
    return min(
        one_way_maximum_path_loss_db(first, second, parameters),
        one_way_maximum_path_loss_db(second, first, parameters),
    )


def standard_link_thresholds_db(
    parameters: CommunicationParameters,
) -> Dict[str, float]:
    radios = parameters.radios
    return {
        "UG": bidirectional_maximum_path_loss_db(
            radios[TRANSPORT_RADIO], radios[GATEWAY_RADIO], parameters
        ),
        "UR": bidirectional_maximum_path_loss_db(
            radios[TRANSPORT_RADIO], radios[RELAY_ACCESS_RADIO], parameters
        ),
        "RG": bidirectional_maximum_path_loss_db(
            radios[RELAY_BACKHAUL_RADIO], radios[GATEWAY_RADIO], parameters
        ),
    }


def evaluate_bidirectional_link(
    first_point: Point3D,
    first_radio: RadioProfile,
    second_point: Point3D,
    second_radio: RadioProfile,
    parameters: CommunicationParameters,
    dem: DemData,
) -> LinkEvaluation:
    distance_m = distance_3d_m(first_point, second_point)
    distance_km = distance_m / 1000.0
    line_of_sight = analyze_line_of_sight(first_point, second_point, dem)
    free_space_loss = fspl_db(parameters.frequency_mhz, distance_km)
    extra_loss = parameters.obstruction_loss_db * int(line_of_sight.obstructed)
    path_loss = free_space_loss + extra_loss
    threshold = bidirectional_maximum_path_loss_db(
        first_radio, second_radio, parameters
    )
    return LinkEvaluation(
        distance_3d_m=distance_m,
        distance_3d_km=distance_km,
        fspl_db=free_space_loss,
        obstructed=line_of_sight.obstructed,
        obstruction_loss_db=extra_loss,
        total_path_loss_db=path_loss,
        maximum_path_loss_db=threshold,
        available=path_loss <= threshold + TOLERANCE,
    )


def evaluate_ug_link(
    transport_point: Point3D,
    gateway_point: Point3D,
    parameters: CommunicationParameters,
    dem: DemData,
) -> LinkEvaluation:
    return evaluate_bidirectional_link(
        transport_point,
        parameters.radios[TRANSPORT_RADIO],
        gateway_point,
        parameters.radios[GATEWAY_RADIO],
        parameters,
        dem,
    )


def evaluate_ur_link(
    transport_point: Point3D,
    relay_point: Point3D,
    parameters: CommunicationParameters,
    dem: DemData,
) -> LinkEvaluation:
    return evaluate_bidirectional_link(
        transport_point,
        parameters.radios[TRANSPORT_RADIO],
        relay_point,
        parameters.radios[RELAY_ACCESS_RADIO],
        parameters,
        dem,
    )


def evaluate_rg_link(
    relay_point: Point3D,
    gateway_point: Point3D,
    parameters: CommunicationParameters,
    dem: DemData,
) -> LinkEvaluation:
    return evaluate_bidirectional_link(
        relay_point,
        parameters.radios[RELAY_BACKHAUL_RADIO],
        gateway_point,
        parameters.radios[GATEWAY_RADIO],
        parameters,
        dem,
    )


def deterministic_time_grid(
    start_s: float,
    end_s: float,
    stage_boundaries_s: Iterable[float] = (),
    step_s: float = COMMUNICATION_TIME_STEP_S,
) -> Tuple[float, ...]:
    """细时间网格中强制加入全部轨迹阶段边界。"""
    if end_s < start_s:
        raise ValueError("时间区间终点不能早于起点")
    if step_s <= 0.0:
        raise ValueError("通信采样时间步长必须为正")
    values = {float(start_s), float(end_s)}
    count = int(math.floor((end_s - start_s) / step_s))
    values.update(float(start_s + index * step_s) for index in range(count + 1))
    values.update(
        float(value)
        for value in stage_boundaries_s
        if start_s - TOLERANCE <= value <= end_s + TOLERANCE
    )
    return tuple(sorted(values))


def refine_binary_state_transition(
    state_function: Callable[[float], int],
    left_s: float,
    right_s: float,
    tolerance_s: float = STATE_TRANSITION_TOLERANCE_S,
) -> float:
    """对A_ij(t)的0/1单次状态切换区间进行二分细化。"""
    if right_s <= left_s:
        raise ValueError("二分区间必须满足right_s > left_s")
    if tolerance_s <= 0.0:
        raise ValueError("状态切换时间精度必须为正")
    left_state = int(bool(state_function(left_s)))
    right_state = int(bool(state_function(right_s)))
    if left_state == right_state:
        raise ValueError("二分区间两端的通信状态必须不同")
    left = float(left_s)
    right = float(right_s)
    while right - left > tolerance_s:
        middle = (left + right) / 2.0
        middle_state = int(bool(state_function(middle)))
        if middle_state == left_state:
            left = middle
        else:
            right = middle
    return (left + right) / 2.0


def _synthetic_dem(elevation: np.ndarray) -> DemData:
    return DemData(
        elevation_m=elevation.astype(float),
        transform=Affine(1.0, 0.0, 0.0, 0.0, -1.0, 3.0),
        bounds=BoundingBox(left=0.0, bottom=0.0, right=3.0, top=3.0),
        crs="EPSG:4326",
    )


def run_independent_tests() -> None:
    """只验证通信模块，不构造问题三中继或baseline方案。"""
    parameters = load_communication_parameters()
    gateway = load_g01_endpoint(parameters)
    dem = load_dem()

    assert math.isclose(gateway.ground_elevation_m, 127.7, abs_tol=1e-12)
    assert math.isclose(gateway.point.altitude_m, 147.7, abs_tol=1e-12)
    thresholds = standard_link_thresholds_db(parameters)
    assert math.isclose(thresholds["UG"], 122.0, abs_tol=1e-12)
    assert math.isclose(thresholds["UR"], 116.0, abs_tol=1e-12)
    assert math.isclose(thresholds["RG"], 126.0, abs_tol=1e-12)
    radios = parameters.radios
    one_way_thresholds = {
        "U_to_G": one_way_maximum_path_loss_db(
            radios[TRANSPORT_RADIO], radios[GATEWAY_RADIO], parameters
        ),
        "G_to_U": one_way_maximum_path_loss_db(
            radios[GATEWAY_RADIO], radios[TRANSPORT_RADIO], parameters
        ),
        "U_to_R": one_way_maximum_path_loss_db(
            radios[TRANSPORT_RADIO], radios[RELAY_ACCESS_RADIO], parameters
        ),
        "R_to_U": one_way_maximum_path_loss_db(
            radios[RELAY_ACCESS_RADIO], radios[TRANSPORT_RADIO], parameters
        ),
        "R_to_G": one_way_maximum_path_loss_db(
            radios[RELAY_BACKHAUL_RADIO], radios[GATEWAY_RADIO], parameters
        ),
        "G_to_R": one_way_maximum_path_loss_db(
            radios[GATEWAY_RADIO], radios[RELAY_BACKHAUL_RADIO], parameters
        ),
    }
    expected_one_way = {
        "U_to_G": 122.0,
        "G_to_U": 129.0,
        "U_to_R": 116.0,
        "R_to_U": 116.0,
        "R_to_G": 126.0,
        "G_to_R": 134.0,
    }
    assert all(
        math.isclose(one_way_thresholds[name], expected, abs_tol=1e-12)
        for name, expected in expected_one_way.items()
    )

    vertical_first = Point3D(109.23, 23.00, 100.0)
    vertical_second = Point3D(109.23, 23.00, 200.0)
    assert math.isclose(distance_3d_m(vertical_first, vertical_second), 100.0)

    expected_fspl = 32.45 + 20.0 * math.log10(2400.0)
    assert math.isclose(fspl_db(2400.0, 1.0), expected_fspl, abs_tol=1e-12)

    clear_dem = _synthetic_dem(np.zeros((3, 3)))
    first = Point3D(0.5, 1.5, 10.0)
    second = Point3D(2.5, 1.5, 20.0)
    clear_result = analyze_line_of_sight(first, second, clear_dem)
    assert not clear_result.obstructed
    blocked_elevation = np.zeros((3, 3))
    blocked_elevation[1, 1] = 16.0
    blocked_result = analyze_line_of_sight(
        first, second, _synthetic_dem(blocked_elevation)
    )
    assert blocked_result.obstructed
    # 若只做二维高地相交检查，上述10->20 m倾斜视线无法正确区分；
    # 中间像元内视线最低为12.5 m，将整格地形降到12 m后应恢复无遮挡。
    lower_elevation = np.zeros((3, 3))
    lower_elevation[1, 1] = 12.0
    assert not analyze_line_of_sight(
        first, second, _synthetic_dem(lower_elevation)
    ).obstructed

    grid = deterministic_time_grid(0.0, 2.5, (0.4, 1.7))
    assert 0.4 in grid and 1.7 in grid and grid[-1] == 2.5
    transition = refine_binary_state_transition(
        lambda time_s: int(time_s >= 1.234567), 1.0, 2.0
    )
    assert abs(transition - 1.234567) <= STATE_TRANSITION_TOLERANCE_S

    # 使用真实DEM分别调用UG、UR、RG接口，验证完整计算链路。
    service_lon = 109.2432319
    service_lat = 23.0335927
    transport = Point3D(service_lon, service_lat, 154.0 + 30.0)
    relay_ground = dem_elevation_at(dem, service_lon, service_lat)
    relay = Point3D(service_lon, service_lat, relay_ground + 100.0)
    actual_links = {
        "UG": evaluate_ug_link(transport, gateway.point, parameters, dem),
        "UR": evaluate_ur_link(transport, relay, parameters, dem),
        "RG": evaluate_rg_link(relay, gateway.point, parameters, dem),
    }

    print(f"G01 ground elevation = {gateway.ground_elevation_m:.1f} m")
    print(f"G01 antenna altitude = {gateway.point.altitude_m:.1f} m")
    print(f"Lmax_UG = {thresholds['UG']:.0f} dB")
    print(f"Lmax_UR = {thresholds['UR']:.0f} dB")
    print(f"Lmax_RG = {thresholds['RG']:.0f} dB")
    print(
        "One-way budgets: "
        + ", ".join(
            f"{name}={value:.0f} dB" for name, value in one_way_thresholds.items()
        )
    )
    print(
        "FSPL unit check: "
        f"f={parameters.frequency_mhz:.0f} MHz, D=1.000 km, "
        f"L_FSPL={fspl_db(parameters.frequency_mhz, 1.0):.6f} dB"
    )
    print(
        "Time grid: "
        f"step={COMMUNICATION_TIME_STEP_S:.2f} s, "
        f"transition tolerance={STATE_TRANSITION_TOLERANCE_S:.3f} s"
    )
    print(
        "Synthetic 3D LOS: "
        f"clear={not clear_result.obstructed}, "
        f"blocked={blocked_result.obstructed}"
    )
    for name, result in actual_links.items():
        print(
            f"{name}: D3={result.distance_3d_km:.6f} km, "
            f"blocked={int(result.obstructed)}, "
            f"Lpath={result.total_path_loss_db:.6f} dB, "
            f"Lmax={result.maximum_path_loss_db:.0f} dB, "
            f"A={int(result.available)}"
        )
    print("problem3_communication independent tests: PASS")


if __name__ == "__main__":
    run_independent_tests()
