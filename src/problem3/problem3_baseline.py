"""问题三 Step 5：固定问题二运输方案下的首个完整中继调度 baseline。

本模块不执行 ALNS，不改变运输架次、运输资源或运输开始时刻。中继候选位置
及通信覆盖关系来自 Step 4；中继往返飞行采用方案 B 的统一实际飞行海拔。
"""

from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass, field
from itertools import permutations
from pathlib import Path
from time import perf_counter
from typing import Dict, Iterable, List, Mapping, Optional, Sequence, Tuple
import math
import sys

import numpy as np
import pandas as pd
from openpyxl import load_workbook


# Preserve historical module names used by trusted communication caches.
_PROJECT_ROOT = Path(__file__).resolve().parents[2]
if str(_PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(_PROJECT_ROOT))
import project_paths as paths
paths.bootstrap()
MODULE_DIR = Path(__file__).resolve().parent
PROJECT_DIR = paths.ROOT

import problem3_communication as communication  # noqa: E402
import problem3_core as core  # noqa: E402
import problem3_relay as relay  # noqa: E402


RELAY_DATA_FILE = paths.raw_file("中继无人机数据.xlsx")
PROBLEM2_FILE = paths.input_file("problem2", "problem2_Z1_priority_result.xlsx")
STEP3_FILE = paths.input_file("problem3", "problem3_transport_communication_check.xlsx")
STEP4_FILE = paths.input_file("problem3", "problem3_relay_coverage.xlsx")
OUTPUT_FILE = paths.output_file("problem3", "problem3_baseline.xlsx")

GRAVITY_M_S2 = 9.8
JOULES_PER_KWH = 3_600_000.0
TIME_TOLERANCE_S = 1e-6
ENERGY_TOLERANCE_KWH = 1e-9
# Step 3 的中断边界本身由 0.01 s 二分得到。中继服务窗在两端各留同量裕度，
# 避免两个独立二分过程在同一真实切换点两侧取代表值时产生毫秒级空隙。
RELAY_SERVICE_BOUNDARY_BUFFER_S = communication.STATE_TRANSITION_TOLERANCE_S

# 原文只定义“通信服务能耗由悬停功率与通信附加功率共同确定”，没有明确说明
# 30 s 建链阶段的能耗。本 baseline 严格只对 SERVICE 计 1.10 kW，SETUP 记0。
LINK_SETUP_ENERGY_POLICY = (
    "附件未明确建链30s是否消耗悬停及通信附加功率；baseline按题目明确的"
    "通信服务能耗定义，仅对SERVICE阶段计1.10kW，LINK_SETUP能耗记0。"
)


@dataclass(frozen=True)
class RelayParameters:
    model_id: str
    planned_mass_kg: float
    cruise_speed_m_s: float
    cruise_power_kw: float
    usable_energy_kwh: float
    reserve_ratio: float
    prepare_time_s: float
    link_setup_time_s: float
    turnaround_time_s: float
    climb_speed_m_s: float
    descent_speed_m_s: float
    climb_efficiency: float
    descent_efficiency: float
    hover_power_kw: float
    communication_power_kw: float
    maximum_hover_agl_m: float
    relay_uav_ids: Tuple[str, ...]
    energy_component_ids: Tuple[str, ...]
    full_charge_time_s: float

    @property
    def service_power_kw(self) -> float:
        return self.hover_power_kw + self.communication_power_kw

    @property
    def sortie_energy_limit_kwh(self) -> float:
        return (1.0 - self.reserve_ratio) * self.usable_energy_kwh


@dataclass(frozen=True)
class RelayPosition:
    position_id: str
    candidate_id: str
    longitude_deg: float
    latitude_deg: float
    ground_elevation_m: float
    hover_height_agl_m: float
    hover_absolute_altitude_m: float
    segment_maximum_ground_elevation_m: float
    terrain_clearance_altitude_m: float
    actual_flight_altitude_m: float
    horizontal_distance_m: float
    outbound_climb_height_m: float
    outbound_descent_height_m: float
    return_climb_height_m: float
    return_descent_height_m: float
    outbound_flight_time_s: float
    return_flight_time_s: float
    outbound_energy_kwh: float
    return_energy_kwh: float
    maximum_service_time_s: float
    endpoint: communication.Point3D = field(repr=False, compare=False)


@dataclass(frozen=True)
class AssignedDemand:
    demand: core.CommunicationDemand
    position_id: str
    eligible_position_ids: Tuple[str, ...]


@dataclass(frozen=True)
class RelayServiceTask:
    task_id: str
    position_id: str
    required_start_s: float
    required_end_s: float
    duration_s: float
    covered_transport_sorties: Tuple[str, ...]
    covered_transport_uavs: Tuple[str, ...]
    covered_outages: Tuple[str, ...]
    covered_demand_ids: Tuple[str, ...]
    demand_duration_weight_s: float


@dataclass(frozen=True)
class RelaySortiePlan:
    plan_id: str
    position_id: str
    service_start_s: float
    service_end_s: float
    task_ids: Tuple[str, ...]
    covered_transport_sorties: Tuple[str, ...]
    covered_transport_uavs: Tuple[str, ...]
    covered_outages: Tuple[str, ...]
    covered_demand_ids: Tuple[str, ...]

    @property
    def service_duration_s(self) -> float:
        return self.service_end_s - self.service_start_s


@dataclass(frozen=True)
class ScheduledRelaySortie:
    sortie_id: str
    relay_uav_id: str
    energy_component_id: str
    position_id: str
    task_ids: Tuple[str, ...]
    covered_transport_sorties: Tuple[str, ...]
    covered_transport_uavs: Tuple[str, ...]
    covered_outages: Tuple[str, ...]
    covered_demand_ids: Tuple[str, ...]
    prepare_start_s: float
    prepare_end_s: float
    outbound_start_s: float
    outbound_end_s: float
    setup_start_s: float
    setup_end_s: float
    service_start_s: float
    service_end_s: float
    return_start_s: float
    return_end_s: float
    turnaround_end_s: float
    outbound_energy_kwh: float
    setup_energy_kwh: float
    service_energy_kwh: float
    return_energy_kwh: float
    total_energy_kwh: float
    soc_after: float
    charge_start_s: float
    charge_end_s: float

    @property
    def service_duration_s(self) -> float:
        return self.service_end_s - self.service_start_s


@dataclass(frozen=True)
class CommunicationCheckResult:
    timeline: pd.DataFrame
    uncovered_duration_s: float
    uncovered_demand_ids: Tuple[str, ...]
    all_samples_available: bool
    all_relay_gateway_links_available: bool


@dataclass
class BaselineResult:
    parameters: RelayParameters
    positions: Tuple[RelayPosition, ...]
    assignments: Tuple[AssignedDemand, ...]
    service_tasks: Tuple[RelayServiceTask, ...]
    relay_sorties: Tuple[ScheduledRelaySortie, ...]
    communication: CommunicationCheckResult
    objective: Mapping[str, float]
    feasibility: pd.DataFrame
    chosen_position_priority: Tuple[str, ...]
    runtime_s: float


class SchedulingConflict(RuntimeError):
    pass


def load_relay_parameters(path: Path = RELAY_DATA_FILE) -> RelayParameters:
    model = pd.read_excel(paths.require_file(path), sheet_name="数据", header=1).iloc[0]
    raw = pd.read_excel(paths.require_file(path), sheet_name="数据", header=None)
    relay_ids = tuple(
        str(value).strip()
        for value in raw.iloc[:, 0]
        if str(value).strip() in {"R01", "R02"}
    )
    component_row = raw.loc[
        raw.iloc[:, 0].astype(str).str.strip().eq("R")
        & pd.to_numeric(raw.iloc[:, 1], errors="coerce").notna()
    ].iloc[-1]
    component_count = int(component_row.iloc[1])
    return RelayParameters(
        model_id=str(model["机型编号"]).strip(),
        planned_mass_kg=float(model["计划起飞总质量（kg）"]),
        cruise_speed_m_s=float(model["计划巡航速度（m/s）"]),
        cruise_power_kw=float(model["巡航功率（kW）"]),
        usable_energy_kwh=float(model["能源组件可用能量（kWh）"]),
        reserve_ratio=float(model["返航电量下限（%）"]) / 100.0,
        prepare_time_s=float(model["工位固定准备时间（s）"]),
        link_setup_time_s=float(model["建链时间（s）"]),
        turnaround_time_s=float(model["架次周转时间（s）"]),
        climb_speed_m_s=float(model["最大爬升速度（m/s）"]),
        descent_speed_m_s=float(model["最大下降速度（m/s）"]),
        climb_efficiency=float(model["爬升能耗效率"]),
        descent_efficiency=float(model["下降能耗效率"]),
        hover_power_kw=float(model["悬停功率（kW）"]),
        communication_power_kw=float(model["通信附加功率（kW）"]),
        maximum_hover_agl_m=float(model["最大悬停离地高度（m）"]),
        relay_uav_ids=relay_ids,
        energy_component_ids=tuple(
            f"REC-{index:02d}" for index in range(1, component_count + 1)
        ),
        full_charge_time_s=float(component_row.iloc[2]),
    )


def _climb_energy_kwh(height_m: float, parameters: RelayParameters) -> float:
    return (
        parameters.planned_mass_kg
        * GRAVITY_M_S2
        * height_m
        / (parameters.climb_efficiency * JOULES_PER_KWH)
    )


def load_relay_positions(
    parameters: RelayParameters,
    gateway: communication.GatewayEndpoint,
    path: Path = STEP4_FILE,
) -> Tuple[RelayPosition, ...]:
    table = pd.read_excel(paths.require_file(path), sheet_name="Greedy_Selection")
    table = table.loc[
        table["Solution"].eq("FINAL_WITH_LOCAL_REFINEMENT")
    ].sort_values("Greedy_Iteration")
    positions: List[RelayPosition] = []
    for order, row in enumerate(table.itertuples(index=False), start=1):
        hover_height = float(row.Hover_Height_m)
        if not (0.0 < hover_height <= parameters.maximum_hover_agl_m):
            raise ValueError("Step 4 入选位置违反附件悬停离地高度上限")
        horizontal_distance = communication.horizontal_distance_m(
            gateway.point,
            communication.Point3D(
                float(row.Longitude),
                float(row.Latitude),
                float(row.Absolute_Altitude_m),
            ),
        )
        outbound_climb = float(row.Outbound_Climb_Height_m)
        outbound_descent = float(row.Outbound_Descent_Height_m)
        return_climb = float(row.Return_Climb_Height_m)
        return_descent = float(row.Return_Descent_Height_m)
        outbound_time = (
            outbound_climb / parameters.climb_speed_m_s
            + horizontal_distance / parameters.cruise_speed_m_s
            + outbound_descent / parameters.descent_speed_m_s
        )
        return_time = (
            return_climb / parameters.climb_speed_m_s
            + horizontal_distance / parameters.cruise_speed_m_s
            + return_descent / parameters.descent_speed_m_s
        )
        horizontal_energy = (
            parameters.cruise_power_kw
            * (horizontal_distance / parameters.cruise_speed_m_s)
            / 3600.0
        )
        outbound_energy = horizontal_energy + _climb_energy_kwh(
            outbound_climb, parameters
        )
        return_energy = horizontal_energy + _climb_energy_kwh(
            return_climb, parameters
        )
        remaining_service_energy = (
            parameters.sortie_energy_limit_kwh - outbound_energy - return_energy
        )
        if remaining_service_energy <= 0.0:
            raise ValueError(f"位置 {row.Candidate_ID} 往返飞行已超过能量上限")
        maximum_service_time = (
            remaining_service_energy / parameters.service_power_kw * 3600.0
        )
        position = RelayPosition(
            position_id=f"P{order}",
            candidate_id=str(row.Candidate_ID),
            longitude_deg=float(row.Longitude),
            latitude_deg=float(row.Latitude),
            ground_elevation_m=float(row.Ground_Elevation_m),
            hover_height_agl_m=hover_height,
            hover_absolute_altitude_m=float(row.Absolute_Altitude_m),
            segment_maximum_ground_elevation_m=float(
                row.Segment_Maximum_Ground_Elevation_m
            ),
            terrain_clearance_altitude_m=float(row.Terrain_Clearance_Altitude_m),
            actual_flight_altitude_m=float(row.Actual_Flight_Altitude_m),
            horizontal_distance_m=horizontal_distance,
            outbound_climb_height_m=outbound_climb,
            outbound_descent_height_m=outbound_descent,
            return_climb_height_m=return_climb,
            return_descent_height_m=return_descent,
            outbound_flight_time_s=outbound_time,
            return_flight_time_s=return_time,
            outbound_energy_kwh=outbound_energy,
            return_energy_kwh=return_energy,
            maximum_service_time_s=maximum_service_time,
            endpoint=communication.Point3D(
                float(row.Longitude),
                float(row.Latitude),
                float(row.Absolute_Altitude_m),
            ),
        )
        expected_flight_altitude = max(
            position.terrain_clearance_altitude_m,
            gateway.ground_elevation_m,
            position.hover_absolute_altitude_m,
        )
        if not math.isclose(
            position.actual_flight_altitude_m,
            expected_flight_altitude,
            abs_tol=1e-6,
        ):
            raise AssertionError("Step 4 实际飞行海拔不符合方案B")
        positions.append(position)
    return tuple(positions)


def load_demands_and_coverage(
    positions: Sequence[RelayPosition],
) -> Tuple[Tuple[core.CommunicationDemand, ...], Dict[str, Tuple[str, ...]]]:
    demands, _ = relay.load_step3_demands(STEP3_FILE)
    coverage = pd.read_excel(paths.require_file(STEP4_FILE), sheet_name="Demand_Coverage")
    coverage_map_raw = coverage.set_index("Demand_ID")[
        "Selected_Covering_Candidate_IDs"
    ].fillna("").to_dict()
    candidate_to_position = {
        position.candidate_id: position.position_id for position in positions
    }
    coverage_map: Dict[str, Tuple[str, ...]] = {}
    for demand in demands:
        candidate_ids = tuple(
            value
            for value in str(coverage_map_raw[demand.demand_id]).split(";")
            if value
        )
        eligible = tuple(
            candidate_to_position[value]
            for value in candidate_ids
            if value in candidate_to_position
        )
        if not eligible:
            raise ValueError(f"需求 {demand.demand_id} 不被最终3个位置覆盖")
        coverage_map[demand.demand_id] = eligible
    return demands, coverage_map


def assign_demands_by_priority(
    demands: Sequence[core.CommunicationDemand],
    coverage_map: Mapping[str, Tuple[str, ...]],
    position_priority: Sequence[str],
) -> Tuple[AssignedDemand, ...]:
    rank = {position_id: index for index, position_id in enumerate(position_priority)}
    assignments = []
    for demand in demands:
        eligible = coverage_map[demand.demand_id]
        selected = min(eligible, key=lambda position_id: rank[position_id])
        assignments.append(
            AssignedDemand(
                demand=demand,
                position_id=selected,
                eligible_position_ids=eligible,
            )
        )
    return tuple(assignments)


def build_service_tasks(
    assignments: Sequence[AssignedDemand],
) -> Tuple[RelayServiceTask, ...]:
    """把同一位置连续或重叠的需求区间合并为 RelayServiceTask。"""
    grouped: Dict[str, List[AssignedDemand]] = defaultdict(list)
    for assignment in assignments:
        grouped[assignment.position_id].append(assignment)
    raw_tasks = []
    for position_id, rows in grouped.items():
        rows.sort(
            key=lambda row: (
                row.demand.interval_start_s,
                row.demand.interval_end_s,
                row.demand.demand_id,
            )
        )
        current: List[AssignedDemand] = []
        current_start = math.nan
        current_end = math.nan
        for row in rows:
            if not current:
                current = [row]
                current_start = row.demand.interval_start_s
                current_end = row.demand.interval_end_s
                continue
            if row.demand.interval_start_s <= current_end + TIME_TOLERANCE_S:
                current.append(row)
                current_end = max(current_end, row.demand.interval_end_s)
                continue
            raw_tasks.append((position_id, current_start, current_end, tuple(current)))
            current = [row]
            current_start = row.demand.interval_start_s
            current_end = row.demand.interval_end_s
        if current:
            raw_tasks.append((position_id, current_start, current_end, tuple(current)))

    raw_tasks.sort(key=lambda row: (row[1], row[2], row[0]))
    tasks = []
    for index, (position_id, start, end, rows) in enumerate(raw_tasks, start=1):
        demands = tuple(row.demand for row in rows)
        tasks.append(
            RelayServiceTask(
                task_id=f"RST-{index:03d}",
                position_id=position_id,
                required_start_s=float(start),
                required_end_s=float(end),
                duration_s=float(end - start),
                covered_transport_sorties=tuple(
                    sorted({row.sortie_id for row in demands})
                ),
                covered_transport_uavs=tuple(
                    sorted({row.transport_uav_id for row in demands})
                ),
                covered_outages=tuple(sorted({row.outage_id for row in demands})),
                covered_demand_ids=tuple(sorted(row.demand_id for row in demands)),
                demand_duration_weight_s=sum(
                    row.duration_weight_s for row in demands
                ),
            )
        )
    return tuple(tasks)


def _split_task(
    task: RelayServiceTask,
    maximum_duration_s: float,
) -> Tuple[RelayServiceTask, ...]:
    if task.duration_s <= maximum_duration_s + TIME_TOLERANCE_S:
        return (task,)
    pieces = []
    start = task.required_start_s
    while start < task.required_end_s - TIME_TOLERANCE_S:
        end = min(task.required_end_s, start + maximum_duration_s)
        pieces.append(
            RelayServiceTask(
                task_id=f"{task.task_id}-S{len(pieces) + 1:02d}",
                position_id=task.position_id,
                required_start_s=start,
                required_end_s=end,
                duration_s=end - start,
                covered_transport_sorties=task.covered_transport_sorties,
                covered_transport_uavs=task.covered_transport_uavs,
                covered_outages=task.covered_outages,
                covered_demand_ids=task.covered_demand_ids,
                demand_duration_weight_s=task.demand_duration_weight_s,
            )
        )
        start = end
    return tuple(pieces)


def build_sortie_plans(
    tasks: Sequence[RelayServiceTask],
    positions: Mapping[str, RelayPosition],
) -> Tuple[RelaySortiePlan, ...]:
    """同位置任务在能量允许时由一个持续 SERVICE 窗口跨越空档统一保障。"""
    grouped: Dict[str, List[RelayServiceTask]] = defaultdict(list)
    for task in tasks:
        grouped[task.position_id].extend(
            _split_task(
                task,
                positions[task.position_id].maximum_service_time_s
                - 2.0 * RELAY_SERVICE_BOUNDARY_BUFFER_S,
            )
        )

    provisional = []
    for position_id, position_tasks in grouped.items():
        position_tasks.sort(key=lambda task: (task.required_start_s, task.required_end_s))
        batch: List[RelayServiceTask] = []
        batch_start = math.nan
        batch_end = math.nan
        for task in position_tasks:
            if not batch:
                batch = [task]
                batch_start = task.required_start_s
                batch_end = task.required_end_s
                continue
            proposed_end = max(batch_end, task.required_end_s)
            if (
                proposed_end
                - batch_start
                + 2.0 * RELAY_SERVICE_BOUNDARY_BUFFER_S
                <= positions[position_id].maximum_service_time_s + TIME_TOLERANCE_S
            ):
                batch.append(task)
                batch_end = proposed_end
                continue
            provisional.append((position_id, batch_start, batch_end, tuple(batch)))
            batch = [task]
            batch_start = task.required_start_s
            batch_end = task.required_end_s
        if batch:
            provisional.append((position_id, batch_start, batch_end, tuple(batch)))

    provisional.sort(key=lambda row: (row[1], row[2], row[0]))
    plans = []
    for index, (position_id, start, end, batch) in enumerate(provisional, start=1):
        plans.append(
            RelaySortiePlan(
                plan_id=f"RSP-{index:03d}",
                position_id=position_id,
                service_start_s=max(
                    0.0, float(start) - RELAY_SERVICE_BOUNDARY_BUFFER_S
                ),
                service_end_s=float(end) + RELAY_SERVICE_BOUNDARY_BUFFER_S,
                task_ids=tuple(task.task_id for task in batch),
                covered_transport_sorties=tuple(
                    sorted(set().union(*(set(task.covered_transport_sorties) for task in batch)))
                ),
                covered_transport_uavs=tuple(
                    sorted(set().union(*(set(task.covered_transport_uavs) for task in batch)))
                ),
                covered_outages=tuple(
                    sorted(set().union(*(set(task.covered_outages) for task in batch)))
                ),
                covered_demand_ids=tuple(
                    sorted(set().union(*(set(task.covered_demand_ids) for task in batch)))
                ),
            )
        )
    return tuple(plans)


def charge_time_s(soc_after: float, parameters: RelayParameters) -> float:
    soc = min(1.0, max(0.0, soc_after))
    if soc < 0.90:
        return parameters.full_charge_time_s * (
            0.65 * (0.90 - soc) / 0.90 + 0.35
        )
    return parameters.full_charge_time_s * 0.35 * (1.0 - soc) / 0.10


def schedule_relay_sorties(
    plans: Sequence[RelaySortiePlan],
    positions: Mapping[str, RelayPosition],
    parameters: RelayParameters,
    timing_accumulator: Optional[Dict[str, float]] = None,
) -> Tuple[ScheduledRelaySortie, ...]:
    schedule_started = perf_counter()
    component_selection_s = 0.0

    def record_timing() -> None:
        if timing_accumulator is None:
            return
        total_s = perf_counter() - schedule_started
        timing_accumulator["relay_component_scheduling"] = component_selection_s
        timing_accumulator["relay_uav_scheduling"] = max(
            0.0, total_s - component_selection_s
        )

    uav_available = {relay_id: 0.0 for relay_id in parameters.relay_uav_ids}
    component_available = {
        component_id: 0.0 for component_id in parameters.energy_component_ids
    }
    scheduled = []
    for plan in sorted(plans, key=lambda row: (row.service_start_s, row.position_id)):
        position = positions[plan.position_id]
        prepare_start = (
            plan.service_start_s
            - parameters.link_setup_time_s
            - position.outbound_flight_time_s
            - parameters.prepare_time_s
        )
        if prepare_start < -TIME_TOLERANCE_S:
            record_timing()
            raise SchedulingConflict(
                f"{plan.plan_id} 缺少飞行提前量: prepare_start={prepare_start:.6f}s"
            )
        eligible_uavs = [
            relay_id
            for relay_id, available in uav_available.items()
            if available <= prepare_start + TIME_TOLERANCE_S
        ]
        if not eligible_uavs:
            detail = ", ".join(
                f"{relay_id}:{available:.3f}"
                for relay_id, available in sorted(uav_available.items())
            )
            record_timing()
            raise SchedulingConflict(
                f"{plan.plan_id} 在 {prepare_start:.3f}s 无可用实体中继；{detail}"
            )
        relay_uav_id = min(
            eligible_uavs, key=lambda relay_id: (uav_available[relay_id], relay_id)
        )

        component_started = perf_counter()
        eligible_components = [
            component_id
            for component_id, available in component_available.items()
            if available <= prepare_start + TIME_TOLERANCE_S
        ]
        if not eligible_components:
            component_selection_s += perf_counter() - component_started
            detail = ", ".join(
                f"{component_id}:{available:.3f}"
                for component_id, available in sorted(component_available.items())
            )
            record_timing()
            raise SchedulingConflict(
                f"{plan.plan_id} 在 {prepare_start:.3f}s 无已充满能源组件；{detail}"
            )
        component_id = min(
            eligible_components,
            key=lambda item: (component_available[item], item),
        )
        component_selection_s += perf_counter() - component_started

        service_energy = (
            parameters.service_power_kw * plan.service_duration_s / 3600.0
        )
        setup_energy = 0.0
        total_energy = (
            position.outbound_energy_kwh
            + setup_energy
            + service_energy
            + position.return_energy_kwh
        )
        if total_energy > parameters.sortie_energy_limit_kwh + ENERGY_TOLERANCE_KWH:
            record_timing()
            raise SchedulingConflict(
                f"{plan.plan_id} 能耗 {total_energy:.6f}kWh 超过2.56kWh"
            )
        soc_after = 1.0 - total_energy / parameters.usable_energy_kwh
        if soc_after < parameters.reserve_ratio - ENERGY_TOLERANCE_KWH:
            record_timing()
            raise SchedulingConflict(f"{plan.plan_id} 返回SOC低于20%")

        prepare_end = prepare_start + parameters.prepare_time_s
        outbound_start = prepare_end
        outbound_end = outbound_start + position.outbound_flight_time_s
        setup_start = outbound_end
        setup_end = setup_start + parameters.link_setup_time_s
        if not math.isclose(setup_end, plan.service_start_s, abs_tol=1e-6):
            raise AssertionError("中继架次未在服务开始前完成建链")
        return_start = plan.service_end_s
        return_end = return_start + position.return_flight_time_s
        turnaround_end = return_end + parameters.turnaround_time_s
        charging = charge_time_s(soc_after, parameters)
        charge_start = return_end
        charge_end = charge_start + charging
        sortie_id = f"RS{len(scheduled) + 1:03d}"
        scheduled.append(
            ScheduledRelaySortie(
                sortie_id=sortie_id,
                relay_uav_id=relay_uav_id,
                energy_component_id=component_id,
                position_id=plan.position_id,
                task_ids=plan.task_ids,
                covered_transport_sorties=plan.covered_transport_sorties,
                covered_transport_uavs=plan.covered_transport_uavs,
                covered_outages=plan.covered_outages,
                covered_demand_ids=plan.covered_demand_ids,
                prepare_start_s=prepare_start,
                prepare_end_s=prepare_end,
                outbound_start_s=outbound_start,
                outbound_end_s=outbound_end,
                setup_start_s=setup_start,
                setup_end_s=setup_end,
                service_start_s=plan.service_start_s,
                service_end_s=plan.service_end_s,
                return_start_s=return_start,
                return_end_s=return_end,
                turnaround_end_s=turnaround_end,
                outbound_energy_kwh=position.outbound_energy_kwh,
                setup_energy_kwh=setup_energy,
                service_energy_kwh=service_energy,
                return_energy_kwh=position.return_energy_kwh,
                total_energy_kwh=total_energy,
                soc_after=soc_after,
                charge_start_s=charge_start,
                charge_end_s=charge_end,
            )
        )
        uav_available[relay_uav_id] = turnaround_end
        component_available[component_id] = charge_end
    record_timing()
    return tuple(scheduled)


def choose_service_assignment(
    demands: Sequence[core.CommunicationDemand],
    coverage_map: Mapping[str, Tuple[str, ...]],
    positions: Sequence[RelayPosition],
    parameters: RelayParameters,
) -> Tuple[
    Tuple[str, ...],
    Tuple[AssignedDemand, ...],
    Tuple[RelayServiceTask, ...],
    Tuple[ScheduledRelaySortie, ...],
]:
    position_map = {position.position_id: position for position in positions}
    feasible = []
    conflicts = []
    for priority in permutations(position_map):
        assignments = assign_demands_by_priority(demands, coverage_map, priority)
        tasks = build_service_tasks(assignments)
        plans = build_sortie_plans(tasks, position_map)
        try:
            scheduled = schedule_relay_sorties(plans, position_map, parameters)
        except SchedulingConflict as error:
            conflicts.append((priority, str(error)))
            continue
        score = (
            len(scheduled),
            sum(row.total_energy_kwh for row in scheduled),
            max(row.return_end_s for row in scheduled),
            priority,
        )
        feasible.append((score, priority, assignments, tasks, scheduled))
    if not feasible:
        details = "\n".join(
            f"priority={priority}: {message}" for priority, message in conflicts
        )
        raise SchedulingConflict("全部位置优先顺序均不可行:\n" + details)
    _, priority, assignments, tasks, scheduled = min(feasible, key=lambda row: row[0])
    return priority, assignments, tasks, scheduled


def rebuild_transport_trajectories() -> Tuple[core.TransportTrajectory, ...]:
    data = core.p2_core.load_data()
    nodes = core.load_nodes()
    _, cruise_altitude = core.load_segment_altitudes()
    records, delivery_table = core.load_problem2_transport_plan(data)
    return core.build_all_transport_trajectories(
        records, data, nodes, cruise_altitude, delivery_table
    )


def verify_continuous_communication(
    trajectories: Sequence[core.TransportTrajectory],
    demands: Sequence[core.CommunicationDemand],
    scheduled: Sequence[ScheduledRelaySortie],
    positions: Mapping[str, RelayPosition],
    parameters: communication.CommunicationParameters,
    gateway: communication.GatewayEndpoint,
    dem: communication.DemData,
) -> CommunicationCheckResult:
    """在1 s网格、阶段边界和状态二分点上重新调用标准链路函数。"""
    rg_evaluations = {
        position_id: communication.evaluate_rg_link(
            position.endpoint, gateway.point, parameters, dem
        )
        for position_id, position in positions.items()
    }
    if not all(row.available for row in rg_evaluations.values()):
        failed = [key for key, value in rg_evaluations.items() if not value.available]
        raise AssertionError(f"最终位置RG标准复核失败: {failed}")

    ug_cache: Dict[Tuple[float, float, float], communication.LinkEvaluation] = {}
    ur_cache: Dict[
        Tuple[str, float, float, float], communication.LinkEvaluation
    ] = {}

    def point_key(point: communication.Point3D) -> Tuple[float, float, float]:
        return (
            round(point.longitude_deg, 10),
            round(point.latitude_deg, 10),
            round(point.altitude_m, 6),
        )

    def evaluate_state(
        trajectory: core.TransportTrajectory, time_s: float
    ) -> Dict[str, object]:
        phase = trajectory.phase_at(time_s)
        point = phase.position_at(time_s)
        key = point_key(point)
        if key not in ug_cache:
            ug_cache[key] = communication.evaluate_ug_link(
                point, gateway.point, parameters, dem
            )
        ug = ug_cache[key]
        active = sorted(
            (
                row
                for row in scheduled
                if row.service_start_s - TIME_TOLERANCE_S
                <= time_s
                <= row.service_end_s + TIME_TOLERANCE_S
            ),
            key=lambda row: (row.relay_uav_id, row.sortie_id),
        )
        source = "G01" if ug.available else "UNAVAILABLE"
        selected_sortie: Optional[ScheduledRelaySortie] = None
        selected_ur: Optional[communication.LinkEvaluation] = None
        selected_rg: Optional[communication.LinkEvaluation] = None
        if not ug.available:
            for relay_sortie in active:
                position = positions[relay_sortie.position_id]
                ur_key = (relay_sortie.position_id, *key)
                if ur_key not in ur_cache:
                    ur_cache[ur_key] = communication.evaluate_ur_link(
                        point, position.endpoint, parameters, dem
                    )
                ur = ur_cache[ur_key]
                rg = rg_evaluations[relay_sortie.position_id]
                if ur.available and rg.available:
                    source = relay_sortie.relay_uav_id
                    selected_sortie = relay_sortie
                    selected_ur = ur
                    selected_rg = rg
                    break
        return {
            "time_s": time_s,
            "phase": phase.phase,
            "point": point,
            "ug": ug,
            "source": source,
            "relay_sortie": selected_sortie,
            "ur": selected_ur,
            "rg": selected_rg,
            "available": source != "UNAVAILABLE",
        }

    timeline_rows = []
    all_available = True
    uncovered_duration = 0.0
    for trajectory in trajectories:
        boundaries = [
            value
            for phase in trajectory.communication_phases
            for value in (phase.start_time_s, phase.end_time_s)
        ]
        boundaries.extend(
            value
            for row in scheduled
            for value in (row.service_start_s, row.service_end_s)
            if trajectory.communication_start_s - TIME_TOLERANCE_S
            <= value
            <= trajectory.communication_end_s + TIME_TOLERANCE_S
        )
        base_times = list(
            communication.deterministic_time_grid(
                trajectory.communication_start_s,
                trajectory.communication_end_s,
                boundaries,
                communication.COMMUNICATION_TIME_STEP_S,
            )
        )
        state_cache = {
            time_s: evaluate_state(trajectory, time_s) for time_s in base_times
        }
        refined_times = set()
        for left_time, right_time in zip(base_times, base_times[1:]):
            left = state_cache[left_time]
            right = state_cache[right_time]
            if left["ug"].available != right["ug"].available:
                transition = communication.refine_binary_state_transition(
                    lambda value: communication.evaluate_ug_link(
                        trajectory.phase_at(value).position_at(value),
                        gateway.point,
                        parameters,
                        dem,
                    ).available,
                    left_time,
                    right_time,
                    communication.STATE_TRANSITION_TOLERANCE_S,
                )
                refined_times.add(transition)
            if bool(left["available"]) != bool(right["available"]):
                transition = communication.refine_binary_state_transition(
                    lambda value: bool(evaluate_state(trajectory, value)["available"]),
                    left_time,
                    right_time,
                    communication.STATE_TRANSITION_TOLERANCE_S,
                )
                refined_times.add(transition)
        for time_s in refined_times:
            state_cache[time_s] = evaluate_state(trajectory, time_s)
        ordered_times = sorted(state_cache)
        for left_time, right_time in zip(ordered_times, ordered_times[1:]):
            midpoint = (left_time + right_time) / 2.0
            midpoint_state = evaluate_state(trajectory, midpoint)
            if not bool(midpoint_state["available"]):
                uncovered_duration += right_time - left_time
                all_available = False
        for time_s in ordered_times:
            state = state_cache[time_s]
            ug = state["ug"]
            ur = state["ur"]
            rg = state["rg"]
            relay_sortie = state["relay_sortie"]
            point = state["point"]
            available = bool(state["available"])
            all_available &= available
            timeline_rows.append(
                {
                    "Time_s": time_s,
                    "Transport_UAV_ID": trajectory.record.uav_id,
                    "Transport_Sortie_ID": trajectory.record.sortie_id,
                    "Phase": state["phase"],
                    "Longitude": point.longitude_deg,
                    "Latitude": point.latitude_deg,
                    "Absolute_Altitude_m": point.altitude_m,
                    "A_UG": int(ug.available),
                    "UG_Obstruction": int(ug.obstructed),
                    "UG_Path_Loss_dB": ug.total_path_loss_db,
                    "Communication_Source": state["source"],
                    "Relay_Sortie_ID": relay_sortie.sortie_id if relay_sortie else "",
                    "Relay_Position_ID": relay_sortie.position_id if relay_sortie else "",
                    "A_UR": int(ur.available) if ur is not None else np.nan,
                    "UR_Obstruction": int(ur.obstructed) if ur is not None else np.nan,
                    "UR_Path_Loss_dB": ur.total_path_loss_db if ur is not None else np.nan,
                    "A_RG": int(rg.available) if rg is not None else np.nan,
                    "Combined_Available": int(available),
                    "Is_Refined_Transition": int(time_s in refined_times),
                    "Is_Service_Boundary": int(
                        any(
                            abs(time_s - value) <= TIME_TOLERANCE_S
                            for row in scheduled
                            for value in (row.service_start_s, row.service_end_s)
                        )
                    ),
                }
            )

    trajectory_map = {row.record.sortie_id: row for row in trajectories}
    uncovered_demand_ids = []
    for demand in demands:
        state = evaluate_state(
            trajectory_map[demand.sortie_id], demand.sample_time_s
        )
        if not bool(state["available"]):
            uncovered_demand_ids.append(demand.demand_id)
    if uncovered_demand_ids:
        demand_map = {row.demand_id: row for row in demands}
        uncovered_duration = max(
            uncovered_duration,
            sum(demand_map[value].duration_weight_s for value in uncovered_demand_ids),
        )
        all_available = False
    if uncovered_duration <= TIME_TOLERANCE_S:
        uncovered_duration = 0.0
    timeline = pd.DataFrame(timeline_rows).sort_values(
        ["Transport_Sortie_ID", "Time_s"]
    )
    return CommunicationCheckResult(
        timeline=timeline,
        uncovered_duration_s=uncovered_duration,
        uncovered_demand_ids=tuple(uncovered_demand_ids),
        all_samples_available=all_available,
        all_relay_gateway_links_available=all(
            row.available for row in rg_evaluations.values()
        ),
    )


def build_feasibility_table(
    positions: Sequence[RelayPosition],
    tasks: Sequence[RelayServiceTask],
    scheduled: Sequence[ScheduledRelaySortie],
    parameters: RelayParameters,
    communication_result: CommunicationCheckResult,
    trajectories: Sequence[core.TransportTrajectory],
    transport_table: pd.DataFrame,
) -> pd.DataFrame:
    position_map = {row.position_id: row for row in positions}
    task_map = {row.task_id: row for row in tasks}
    checks: List[Tuple[str, bool, str]] = []
    checks.append((
        "运输方案与问题二正式结果完全一致",
        len(transport_table) == 20
        and max(row.maximum_time_error_s for row in trajectories) <= 1e-6,
        "20个运输架次；轨迹重建时刻与正式结果逐项一致",
    ))
    checks.append((
        "中继悬停离地高度满足0<h<=300m",
        all(0.0 < row.hover_height_agl_m <= parameters.maximum_hover_agl_m for row in positions),
        "",
    ))
    checks.append((
        "方案B实际飞行海拔及往返高度差正确",
        all(
            math.isclose(
                row.actual_flight_altitude_m,
                max(
                    row.terrain_clearance_altitude_m,
                    127.7,
                    row.hover_absolute_altitude_m,
                ),
                abs_tol=1e-6,
            )
            and min(
                row.outbound_climb_height_m,
                row.outbound_descent_height_m,
                row.return_climb_height_m,
                row.return_descent_height_m,
            ) >= -TIME_TOLERANCE_S
            for row in positions
        ),
        "",
    ))
    checks.append((
        "所有RelayServiceTask均被同位置SERVICE窗口完整包含",
        all(
            any(
                task.position_id == sortie.position_id
                and task.task_id in sortie.task_ids
                and sortie.service_start_s <= task.required_start_s + TIME_TOLERANCE_S
                and sortie.service_end_s + TIME_TOLERANCE_S >= task.required_end_s
                for sortie in scheduled
            )
            for task in tasks
        ),
        "",
    ))
    checks.append((
        "每架次准备、飞行、建链、服务、返航阶段顺序正确",
        all(
            row.prepare_start_s <= row.prepare_end_s + TIME_TOLERANCE_S
            and math.isclose(row.prepare_end_s, row.outbound_start_s, abs_tol=1e-6)
            and math.isclose(row.outbound_end_s, row.setup_start_s, abs_tol=1e-6)
            and math.isclose(row.setup_end_s, row.service_start_s, abs_tol=1e-6)
            and math.isclose(row.service_end_s, row.return_start_s, abs_tol=1e-6)
            and row.return_end_s <= row.turnaround_end_s + TIME_TOLERANCE_S
            for row in scheduled
        ),
        "",
    ))
    checks.append((
        "每架次SERVICE不超过该位置最大允许时长",
        all(
            row.service_duration_s
            <= position_map[row.position_id].maximum_service_time_s + TIME_TOLERANCE_S
            for row in scheduled
        ),
        "",
    ))
    checks.append((
        "每架次总能耗不超过2.56kWh",
        all(
            row.total_energy_kwh
            <= parameters.sortie_energy_limit_kwh + ENERGY_TOLERANCE_KWH
            for row in scheduled
        ),
        "",
    ))
    checks.append((
        "每架次返回SOC不低于20%",
        all(row.soc_after >= parameters.reserve_ratio - ENERGY_TOLERANCE_KWH for row in scheduled),
        "",
    ))
    uav_ok = True
    for relay_id in parameters.relay_uav_ids:
        rows = sorted(
            (row for row in scheduled if row.relay_uav_id == relay_id),
            key=lambda row: row.prepare_start_s,
        )
        uav_ok &= all(
            left.turnaround_end_s <= right.prepare_start_s + TIME_TOLERANCE_S
            for left, right in zip(rows, rows[1:])
        )
    checks.append(("R01/R02无任务重叠且满足300s周转", uav_ok, ""))
    component_ok = True
    for component_id in parameters.energy_component_ids:
        rows = sorted(
            (row for row in scheduled if row.energy_component_id == component_id),
            key=lambda row: row.prepare_start_s,
        )
        component_ok &= all(
            left.charge_end_s <= right.prepare_start_s + TIME_TOLERANCE_S
            for left, right in zip(rows, rows[1:])
        )
    checks.append(("能源组件再次使用前已充至100%且无占用/充电冲突", component_ok, ""))
    checks.append((
        "30s LINK_SETUP能耗口径已显式记录",
        all(abs(row.setup_energy_kwh) <= ENERGY_TOLERANCE_KWH for row in scheduled),
        LINK_SETUP_ENERGY_POLICY,
    ))
    checks.append((
        "最终标准链路逐时连续通信复核",
        communication_result.all_samples_available
        and communication_result.uncovered_duration_s <= TIME_TOLERANCE_S,
        f"uncovered={communication_result.uncovered_duration_s:.9f}s; "
        f"uncovered_demands={len(communication_result.uncovered_demand_ids)}",
    ))
    checks.append((
        "全部中继位置RG标准链路可用",
        communication_result.all_relay_gateway_links_available,
        "由最终通信复核中的标准evaluate_rg_link确认",
    ))
    return pd.DataFrame(
        [
            {
                "Check_ID": index,
                "Feasibility_Check": description,
                "Result": "PASS" if passed else "FAIL",
                "Detail": detail,
            }
            for index, (description, passed, detail) in enumerate(checks, start=1)
        ]
    )


def _positions_table(positions: Sequence[RelayPosition]) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "Position_ID": row.position_id,
                "Candidate_ID": row.candidate_id,
                "Longitude": row.longitude_deg,
                "Latitude": row.latitude_deg,
                "DEM_Ground_Elevation_m": row.ground_elevation_m,
                "Hover_Height_AGL_m": row.hover_height_agl_m,
                "Hover_Absolute_Altitude_m": row.hover_absolute_altitude_m,
                "Segment_Maximum_Ground_Elevation_m": row.segment_maximum_ground_elevation_m,
                "Terrain_Clearance_Altitude_m": row.terrain_clearance_altitude_m,
                "Actual_Flight_Altitude_m": row.actual_flight_altitude_m,
                "Horizontal_Distance_m": row.horizontal_distance_m,
                "Outbound_Climb_m": row.outbound_climb_height_m,
                "Outbound_Descent_m": row.outbound_descent_height_m,
                "Return_Climb_m": row.return_climb_height_m,
                "Return_Descent_m": row.return_descent_height_m,
                "Outbound_Flight_Time_s": row.outbound_flight_time_s,
                "Return_Flight_Time_s": row.return_flight_time_s,
                "Outbound_Energy_kWh": row.outbound_energy_kwh,
                "Return_Energy_kWh": row.return_energy_kwh,
                "Maximum_Service_Time_s": row.maximum_service_time_s,
            }
            for row in positions
        ]
    )


def _tasks_table(tasks: Sequence[RelayServiceTask]) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "Task_ID": row.task_id,
                "Position_ID": row.position_id,
                "Required_Service_Start_s": row.required_start_s,
                "Required_Service_End_s": row.required_end_s,
                "Required_Service_Duration_s": row.duration_s,
                "Demand_Duration_Weight_s": row.demand_duration_weight_s,
                "Covered_Transport_Sorties": ";".join(row.covered_transport_sorties),
                "Covered_Transport_UAVs": ";".join(row.covered_transport_uavs),
                "Covered_Outages": ";".join(row.covered_outages),
                "Covered_Demand_Count": len(row.covered_demand_ids),
                "Covered_Demand_IDs": ";".join(row.covered_demand_ids),
            }
            for row in tasks
        ]
    )


def _sorties_table(sorties: Sequence[ScheduledRelaySortie]) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "Sortie_ID": row.sortie_id,
                "Relay_UAV_ID": row.relay_uav_id,
                "Energy_Component_ID": row.energy_component_id,
                "Position_ID": row.position_id,
                "Task_IDs": ";".join(row.task_ids),
                "Prepare_Start_s": row.prepare_start_s,
                "Prepare_End_s": row.prepare_end_s,
                "Outbound_Start_s": row.outbound_start_s,
                "Outbound_End_s": row.outbound_end_s,
                "Setup_Start_s": row.setup_start_s,
                "Setup_End_s": row.setup_end_s,
                "Service_Start_s": row.service_start_s,
                "Service_End_s": row.service_end_s,
                "Service_Duration_s": row.service_duration_s,
                "Return_Start_s": row.return_start_s,
                "Return_End_s": row.return_end_s,
                "Turnaround_End_s": row.turnaround_end_s,
                "Outbound_Energy_kWh": row.outbound_energy_kwh,
                "Setup_Energy_kWh": row.setup_energy_kwh,
                "Service_Energy_kWh": row.service_energy_kwh,
                "Return_Energy_kWh": row.return_energy_kwh,
                "Total_Energy_kWh": row.total_energy_kwh,
                "SOC_After": row.soc_after,
                "Charge_Start_s": row.charge_start_s,
                "Charge_End_s": row.charge_end_s,
                "Covered_Transport_Sorties": ";".join(row.covered_transport_sorties),
                "Covered_Transport_UAVs": ";".join(row.covered_transport_uavs),
                "Covered_Outages": ";".join(row.covered_outages),
                "Covered_Demand_Count": len(row.covered_demand_ids),
            }
            for row in sorties
        ]
    )


def _energy_table(
    sorties: Sequence[ScheduledRelaySortie],
    parameters: RelayParameters,
) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "Sortie_ID": row.sortie_id,
                "Position_ID": row.position_id,
                "Energy_Component_ID": row.energy_component_id,
                "Outbound_Energy_kWh": row.outbound_energy_kwh,
                "Link_Setup_Energy_kWh": row.setup_energy_kwh,
                "Service_Energy_kWh": row.service_energy_kwh,
                "Return_Energy_kWh": row.return_energy_kwh,
                "Total_Energy_kWh": row.total_energy_kwh,
                "Energy_Limit_kWh": parameters.sortie_energy_limit_kwh,
                "SOC_Before": 1.0,
                "SOC_After": row.soc_after,
                "Charge_Duration_s": row.charge_end_s - row.charge_start_s,
                "Link_Setup_Energy_Policy": LINK_SETUP_ENERGY_POLICY,
            }
            for row in sorties
        ]
    )


def _resource_schedule_table(
    sorties: Sequence[ScheduledRelaySortie],
    parameters: RelayParameters,
) -> pd.DataFrame:
    rows = []
    for sortie in sorties:
        rows.extend(
            [
                {
                    "Resource_Type": "Relay_UAV",
                    "Resource_ID": sortie.relay_uav_id,
                    "Sortie_ID": sortie.sortie_id,
                    "Position_ID": sortie.position_id,
                    "Activity": "TASK_OCCUPANCY",
                    "Start_s": sortie.prepare_start_s,
                    "End_s": sortie.return_end_s,
                },
                {
                    "Resource_Type": "Relay_UAV",
                    "Resource_ID": sortie.relay_uav_id,
                    "Sortie_ID": sortie.sortie_id,
                    "Position_ID": sortie.position_id,
                    "Activity": "TURNAROUND",
                    "Start_s": sortie.return_end_s,
                    "End_s": sortie.turnaround_end_s,
                },
                {
                    "Resource_Type": "Energy_Component",
                    "Resource_ID": sortie.energy_component_id,
                    "Sortie_ID": sortie.sortie_id,
                    "Position_ID": sortie.position_id,
                    "Activity": "TASK_OCCUPANCY",
                    "Start_s": sortie.prepare_start_s,
                    "End_s": sortie.return_end_s,
                },
                {
                    "Resource_Type": "Energy_Component",
                    "Resource_ID": sortie.energy_component_id,
                    "Sortie_ID": sortie.sortie_id,
                    "Position_ID": sortie.position_id,
                    "Activity": "CHARGING_TO_100_PERCENT",
                    "Start_s": sortie.charge_start_s,
                    "End_s": sortie.charge_end_s,
                },
            ]
        )
    for component_id in parameters.energy_component_ids:
        if not any(row.energy_component_id == component_id for row in sorties):
            rows.append(
                {
                    "Resource_Type": "Energy_Component",
                    "Resource_ID": component_id,
                    "Sortie_ID": "",
                    "Position_ID": "",
                    "Activity": "UNUSED",
                    "Start_s": 0.0,
                    "End_s": 0.0,
                }
            )
    return pd.DataFrame(rows).sort_values(
        ["Resource_Type", "Resource_ID", "Start_s", "Activity"]
    )


def export_baseline(
    result: BaselineResult,
    transport_table: pd.DataFrame,
    output_file: Path = OUTPUT_FILE,
) -> None:
    sorties = result.relay_sorties
    position_counts = Counter(row.position_id for row in sorties)
    uav_counts = Counter(row.relay_uav_id for row in sorties)
    component_counts = Counter(row.energy_component_id for row in sorties)
    summary = pd.DataFrame(
        [
            {
                "Status": "FEASIBLE" if result.feasibility["Result"].eq("PASS").all() else "INFEASIBLE",
                "Transport_Sorties": len(transport_table),
                "Relay_Sorties": len(sorties),
                "P1_Visits": position_counts["P1"],
                "P2_Visits": position_counts["P2"],
                "P3_Visits": position_counts["P3"],
                "R01_Sorties": uav_counts["R01"],
                "R02_Sorties": uav_counts["R02"],
                "Total_Relay_Service_Time_s": sum(row.service_duration_s for row in sorties),
                "Maximum_Relay_Service_Time_s": max(row.service_duration_s for row in sorties),
                "Average_Relay_Service_Time_s": np.mean([row.service_duration_s for row in sorties]),
                "Total_Relay_Energy_kWh": sum(row.total_energy_kwh for row in sorties),
                "Maximum_Relay_Sortie_Energy_kWh": max(row.total_energy_kwh for row in sorties),
                "Minimum_Return_SOC": min(row.soc_after for row in sorties),
                "Uncovered_Communication_Duration_s": result.communication.uncovered_duration_s,
                "Transport_Schedule_Adjusted": "NO",
                "Chosen_Position_Priority": ">".join(result.chosen_position_priority),
                "Energy_Component_Usage": ";".join(
                    f"{component}:{component_counts[component]}"
                    for component in result.parameters.energy_component_ids
                ),
                "Link_Setup_Energy_Policy": LINK_SETUP_ENERGY_POLICY,
                "Runtime_s": result.runtime_s,
            }
        ]
    )
    objective = pd.DataFrame([dict(result.objective)])
    tables = {
        "Summary": summary,
        "Transport_Sorties": transport_table,
        "Relay_Positions": _positions_table(result.positions),
        "Relay_Service_Tasks": _tasks_table(result.service_tasks),
        "Relay_Sorties": _sorties_table(sorties),
        "Relay_Energy": _energy_table(sorties, result.parameters),
        "Relay_Resource_Schedule": _resource_schedule_table(sorties, result.parameters),
        "Communication_Timeline": result.communication.timeline,
        "Objective": objective,
        "Feasibility": result.feasibility,
    }
    output_file.parent.mkdir(parents=True, exist_ok=True)
    paths.ensure_output_path(output_file)
    with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
        for sheet_name, table in tables.items():
            table.to_excel(writer, sheet_name=sheet_name, index=False)
        for worksheet in writer.book.worksheets:
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            for column in worksheet.columns:
                values = [
                    "" if cell.value is None else str(cell.value)
                    for cell in column[:1000]
                ]
                width = min(max(max(map(len, values), default=10) + 2, 12), 42)
                worksheet.column_dimensions[column[0].column_letter].width = width
                for cell in column[1:]:
                    if isinstance(cell.value, float):
                        cell.number_format = "0.000000000"


def _update_runtime_cell(output_file: Path, runtime_s: float) -> None:
    paths.ensure_output_path(output_file)
    workbook = load_workbook(output_file)
    worksheet = workbook["Summary"]
    headers = {cell.value: cell.column for cell in worksheet[1]}
    worksheet.cell(row=2, column=headers["Runtime_s"], value=runtime_s)
    worksheet.cell(row=2, column=headers["Runtime_s"]).number_format = "0.000000000"
    workbook.save(output_file)


def run_baseline(export: bool = True) -> BaselineResult:
    started = perf_counter()
    relay.validate_fixed_node_flight_altitude_extension()
    relay_parameters = load_relay_parameters()
    communication_parameters = communication.load_communication_parameters()
    gateway = communication.load_g01_endpoint(communication_parameters)
    dem = communication.load_dem()
    positions = load_relay_positions(relay_parameters, gateway)
    position_map = {row.position_id: row for row in positions}
    demands, coverage_map = load_demands_and_coverage(positions)

    for position in positions:
        print(
            f"{position.position_id}/{position.candidate_id}: "
            f"out={position.outbound_flight_time_s:.6f}s, "
            f"back={position.return_flight_time_s:.6f}s, "
            f"Eout={position.outbound_energy_kwh:.9f}kWh, "
            f"Eback={position.return_energy_kwh:.9f}kWh, "
            f"Tservice_max={position.maximum_service_time_s:.6f}s",
            flush=True,
        )

    priority, assignments, tasks, scheduled = choose_service_assignment(
        demands, coverage_map, positions, relay_parameters
    )
    print(
        f"Chosen position priority={priority}; tasks={len(tasks)}; "
        f"relay sorties={len(scheduled)}",
        flush=True,
    )
    for row in scheduled:
        print(
            f"{row.sortie_id}: {row.relay_uav_id}/{row.energy_component_id}, "
            f"{row.position_id}, SERVICE=[{row.service_start_s:.6f}, "
            f"{row.service_end_s:.6f}], E={row.total_energy_kwh:.9f}kWh, "
            f"SOC={row.soc_after:.9f}",
            flush=True,
        )

    trajectories = rebuild_transport_trajectories()
    communication_result = verify_continuous_communication(
        trajectories,
        demands,
        scheduled,
        position_map,
        communication_parameters,
        gateway,
        dem,
    )

    transport_table = pd.read_excel(paths.require_file(PROBLEM2_FILE), sheet_name="Sorties")
    p2_summary = pd.read_excel(paths.require_file(PROBLEM2_FILE), sheet_name="Summary").iloc[0]
    transport_z1 = float(p2_summary["Z1"])
    transport_makespan = float(transport_table["Return_Time_s"].max())
    transport_energy = float(transport_table["Energy_kWh"].sum())
    relay_makespan = max(row.return_end_s for row in scheduled)
    relay_energy = sum(row.total_energy_kwh for row in scheduled)
    objective = {
        "Z1_Weighted_Tardiness": transport_z1,
        "Z2_Joint_Makespan_s": max(transport_makespan, relay_makespan),
        "Z3_Total_Energy_kWh": transport_energy + relay_energy,
        "Z4_Total_Sorties": len(transport_table) + len(scheduled),
        "Transport_Makespan_s": transport_makespan,
        "Relay_Makespan_s": relay_makespan,
        "Transport_Energy_kWh": transport_energy,
        "Relay_Energy_kWh": relay_energy,
        "Transport_Sortie_Count": len(transport_table),
        "Relay_Sortie_Count": len(scheduled),
    }
    feasibility = build_feasibility_table(
        positions,
        tasks,
        scheduled,
        relay_parameters,
        communication_result,
        trajectories,
        transport_table,
    )
    result = BaselineResult(
        parameters=relay_parameters,
        positions=positions,
        assignments=assignments,
        service_tasks=tasks,
        relay_sorties=scheduled,
        communication=communication_result,
        objective=objective,
        feasibility=feasibility,
        chosen_position_priority=priority,
        runtime_s=perf_counter() - started,
    )
    if export:
        export_baseline(result, transport_table)
        result.runtime_s = perf_counter() - started
        _update_runtime_cell(OUTPUT_FILE, result.runtime_s)

    print(
        f"Continuous communication: uncovered="
        f"{communication_result.uncovered_duration_s:.9f}s, "
        f"uncovered_demands={len(communication_result.uncovered_demand_ids)}",
        flush=True,
    )
    print(
        "Objectives: "
        f"Z1={objective['Z1_Weighted_Tardiness']:.6f}, "
        f"Z2={objective['Z2_Joint_Makespan_s']:.6f}s, "
        f"Z3={objective['Z3_Total_Energy_kWh']:.9f}kWh, "
        f"Z4={objective['Z4_Total_Sorties']}",
        flush=True,
    )
    print(
        f"Feasibility={int(feasibility['Result'].eq('PASS').sum())}/"
        f"{len(feasibility)} PASS",
        flush=True,
    )
    print(f"Runtime={result.runtime_s:.6f}s", flush=True)
    if export:
        print(f"Result={OUTPUT_FILE}", flush=True)
    return result


if __name__ == "__main__":
    run_baseline(export=True)
