"""Layered comparison of existing Pareto solutions; no optimization is run."""

from __future__ import annotations

import hashlib
from pathlib import Path
import sys
from typing import Dict, Iterable, List, Mapping, Sequence, Tuple

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from project_paths import bootstrap, input_file, output_dir, require_file

bootstrap()

MODULE_DIR = Path(__file__).resolve().parent
RESULT_DIR = output_dir("problem4")
OUTPUT = RESULT_DIR / "problem4_final_selection_analysis.xlsx"

SOURCE_PATHS = {
    "K2_Pareto": input_file("problem4", "problem4_K2_pareto.xlsx"),
    "K3_Pareto": input_file("problem4", "problem4_K3_pareto.xlsx"),
    "候选比较": input_file("problem4", "problem4_final_candidate_comparison.xlsx"),
    "能耗完工核验": input_file("problem4", "problem4_energy_makespan_verification.xlsx"),
}

UAV_GAP_COLUMNS = ["D_U_A", "D_U_B", "D_U_C"]
BATTERY_GAP_COLUMNS = ["D_B_A", "D_B_B", "D_B_C"]
ALL_GAP_COLUMNS = [
    *UAV_GAP_COLUMNS,
    *BATTERY_GAP_COLUMNS,
    "D_R",
    "D_E",
]
RESOURCE_COLUMNS = [
    "R_U_A",
    "R_U_B",
    "R_U_C",
    "R_U_total",
    "R_B_A",
    "R_B_B",
    "R_B_C",
    "R_B_total",
    "R_R",
    "R_E",
]
BALANCE_SIX = [
    "B_Q",
    "B_M",
    "B_V",
    "B_N_sortie",
    "B_T_transport",
    "B_T_comm",
]
BALANCE_EIGHT = [*BALANCE_SIX, "B_max", "B_mean"]
EXECUTION_COLUMNS = [
    "relay_occurrence_extra_total",
    "E_relay_kwh",
    "Delta_E_total_kwh",
]
DISPLAY_COLUMNS = [
    "partition_id",
    "group1_components",
    "group2_components",
    "group3_components",
    *ALL_GAP_COLUMNS,
    *RESOURCE_COLUMNS,
    "RS001_group_count",
    "RS002_group_count",
    "RS003_group_count",
    "relay_occurrence_extra_total",
    "E_relay_kwh",
    "Delta_E_total_kwh",
    "Cmax_s",
    *BALANCE_EIGHT,
    "resource_nondominated",
    "execution_cost_nondominated",
    "balance_nondominated_in_resource_front",
    "balance_best_Bmax_in_min_gap",
    "layered_role",
]


def _hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_sources() -> Dict[str, object]:
    k2 = pd.read_excel(
        require_file(SOURCE_PATHS["K2_Pareto"]), sheet_name="Pareto_K2"
    )
    k3 = pd.read_excel(
        require_file(SOURCE_PATHS["K3_Pareto"]), sheet_name="Pareto_K3"
    )
    candidate = pd.read_excel(
        require_file(SOURCE_PATHS["候选比较"]), sheet_name="candidate_comparison"
    ).set_index("方案", drop=False)
    energy_summary = pd.read_excel(
        require_file(SOURCE_PATHS["能耗完工核验"]), sheet_name="solution_summary"
    ).set_index("方案", drop=False)
    baseline = pd.read_excel(
        require_file(SOURCE_PATHS["能耗完工核验"]), sheet_name="problem3_baseline"
    ).set_index("指标")
    return {
        "pareto": {2: k2, 3: k3},
        "candidate": candidate,
        "energy_summary": energy_summary,
        "baseline": baseline,
    }


def _nondominated_flags(
    table: pd.DataFrame,
    columns: Sequence[str],
    tolerance: float = 1e-12,
) -> np.ndarray:
    values = table[list(columns)].to_numpy(dtype=float)
    flags = np.ones(len(values), dtype=bool)
    for index, candidate in enumerate(values):
        no_worse = np.all(values <= candidate + tolerance, axis=1)
        strictly_better = np.any(values < candidate - tolerance, axis=1)
        if np.any(no_worse & strictly_better):
            flags[index] = False
    return flags


def _baseline_value(baseline: pd.DataFrame, label: str) -> float:
    return float(baseline.loc[label, "数值"])


def augment_pareto(
    table: pd.DataFrame,
    baseline: pd.DataFrame,
) -> pd.DataFrame:
    result = table.copy()
    result["B_max"] = result[BALANCE_SIX].max(axis=1)
    result["B_mean"] = result[BALANCE_SIX].mean(axis=1)
    relay_energies = {
        relay_id: _baseline_value(baseline, f"{relay_id}完整架次能耗")
        for relay_id in ("RS001", "RS002", "RS003")
    }
    result["E_relay_kwh"] = sum(
        result[f"{relay_id}_group_count"] * relay_energies[relay_id]
        for relay_id in ("RS001", "RS002", "RS003")
    )
    problem3_relay = _baseline_value(baseline, "中继总能耗")
    problem3_cmax = _baseline_value(baseline, "联合完工时间")
    result["Delta_E_total_kwh"] = result["E_relay_kwh"] - problem3_relay
    result["Cmax_s"] = problem3_cmax
    if "group3_components" not in result.columns:
        result["group3_components"] = ""
    return result


def layered_filter(
    augmented: pd.DataFrame,
) -> Tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    minimum_gap = int(augmented["D_R"].min())
    stage1 = augmented.loc[augmented["D_R"].eq(minimum_gap)].copy()
    stage1["resource_nondominated"] = _nondominated_flags(
        stage1, RESOURCE_COLUMNS
    )
    resource_front = stage1.loc[stage1["resource_nondominated"]].copy()
    resource_front["execution_cost_nondominated"] = _nondominated_flags(
        resource_front, EXECUTION_COLUMNS
    )
    execution_front = resource_front.loc[
        resource_front["execution_cost_nondominated"]
    ].copy()
    resource_front["balance_nondominated_in_resource_front"] = (
        _nondominated_flags(resource_front, BALANCE_EIGHT)
    )
    minimum_bmax = float(stage1["B_max"].min())
    stage1["balance_best_Bmax_in_min_gap"] = np.isclose(
        stage1["B_max"], minimum_bmax, atol=1e-12, rtol=0.0
    )

    resource_flags = resource_front.set_index("partition_id")[
        ["execution_cost_nondominated", "balance_nondominated_in_resource_front"]
    ]
    for column in (
        "execution_cost_nondominated",
        "balance_nondominated_in_resource_front",
    ):
        stage1[column] = stage1["partition_id"].map(resource_flags[column]).eq(True)
    roles = []
    for row in stage1.itertuples(index=False):
        role_parts = ["A：最小中继UAV缺口"]
        if row.resource_nondominated:
            role_parts.append("B：资源向量非支配")
        if row.execution_cost_nondominated:
            role_parts.append("C：资源前沿内额外代价非支配")
        if row.balance_nondominated_in_resource_front:
            role_parts.append("D：资源前沿内均衡非支配")
        if row.balance_best_Bmax_in_min_gap:
            role_parts.append("最小缺口集合中B_max最小")
        roles.append("；".join(role_parts))
    stage1["layered_role"] = roles
    resource_front = stage1.loc[stage1["resource_nondominated"]].copy()
    execution_front = stage1.loc[
        stage1["resource_nondominated"]
        & stage1["execution_cost_nondominated"]
    ].copy()
    return stage1, resource_front, execution_front


def _ids(table: pd.DataFrame) -> str:
    return "、".join(table["partition_id"].astype(str))


def build_category_summary(
    k: int,
    pareto_total: int,
    stage1: pd.DataFrame,
    resource_front: pd.DataFrame,
    execution_front: pd.DataFrame,
) -> pd.DataFrame:
    min_gap = int(stage1["D_R"].iloc[0])
    min_energy = float(stage1["Delta_E_total_kwh"].min())
    minimum_energy_schemes = stage1.loc[
        np.isclose(
            stage1["Delta_E_total_kwh"],
            min_energy,
            atol=1e-12,
            rtol=0.0,
        )
    ]
    min_bmax = float(stage1["B_max"].min())
    balance_best = stage1.loc[
        np.isclose(stage1["B_max"], min_bmax, atol=1e-12, rtol=0.0)
    ]
    resource_balance_best = resource_front.sort_values(
        ["B_max", "B_mean", "partition_id"]
    ).head(1)
    return pd.DataFrame(
        [
            {
                "K": k,
                "类别": "Pareto解总数",
                "方案": "—",
                "数值或说明": pareto_total,
            },
            {
                "K": k,
                "类别": "A：最小中继UAV库存缺口集合",
                "方案": _ids(stage1),
                "数值或说明": f"D_R_min={min_gap}；共{len(stage1)}个方案",
            },
            {
                "K": k,
                "类别": "B：最小缺口层资源向量非支配集合",
                "方案": _ids(resource_front),
                "数值或说明": f"共{len(resource_front)}个方案；未设置资源权重",
            },
            {
                "K": k,
                "类别": "C：最小新增能耗方案",
                "方案": _ids(minimum_energy_schemes),
                "数值或说明": f"Delta_E_total={min_energy:.6f} kWh",
            },
            {
                "K": k,
                "类别": "C：资源前沿内额外代价非支配集合",
                "方案": _ids(execution_front),
                "数值或说明": "比较中继复制数、中继能耗和新增能耗",
            },
            {
                "K": k,
                "类别": "D：最小缺口集合中最小B_max",
                "方案": _ids(balance_best),
                "数值或说明": (
                    f"B_max={float(balance_best['B_max'].iloc[0]):.6f}；"
                    f"B_mean={float(balance_best['B_mean'].iloc[0]):.6f}"
                ),
            },
            {
                "K": k,
                "类别": "D：资源前沿内均衡较优方案",
                "方案": _ids(resource_balance_best),
                "数值或说明": (
                    f"B_max={float(resource_balance_best['B_max'].iloc[0]):.6f}；"
                    f"B_mean={float(resource_balance_best['B_mean'].iloc[0]):.6f}"
                ),
            },
        ]
    )


def build_assessment(
    k2_stage1: pd.DataFrame,
    k3_stage1: pd.DataFrame,
) -> pd.DataFrame:
    rows = []
    cases = (
        (
            "K2-051",
            k2_stage1,
            "K2-032、K2-060",
            "属于最小缺口层且为该层最均衡方案；但资源向量分别被K2-032和K2-060支配。"
            "K2-032新增能耗更低，K2-060资源更少但均衡性明显较弱。",
        ),
        (
            "K3-147",
            k3_stage1,
            "K3-002、K3-007",
            "属于最小缺口层且为该层最均衡方案；但资源向量被K3-007支配，"
            "新增能耗也高于K3-002和K3-007。其优势是工作量均衡显著更好。",
        ),
    )
    for partition_id, stage1, alternatives, conclusion in cases:
        row = stage1.set_index("partition_id").loc[partition_id]
        rows.append(
            {
                "方案": partition_id,
                "属于最小D_R集合": True,
                "资源向量非支配": bool(row["resource_nondominated"]),
                "额外代价非支配": bool(row["execution_cost_nondominated"]),
                "最小缺口层B_max最小": bool(
                    row["balance_best_Bmax_in_min_gap"]
                ),
                "D_R": int(row["D_R"]),
                "新增能耗/kWh": float(row["Delta_E_total_kwh"]),
                "B_max": float(row["B_max"]),
                "B_mean": float(row["B_mean"]),
                "同缺口资源前沿替代方案": alternatives,
                "结论": conclusion,
            }
        )
    return pd.DataFrame(rows)


def build_compromise_note() -> pd.DataFrame:
    return pd.DataFrame(
        [
            {
                "K": 2,
                "是否存在无歧义的综合折中候选": False,
                "说明": (
                    "K2-032同时满足最小缺口、资源非支配和最小新增能耗，但B_max较高；"
                    "K2-060位于资源前沿且均衡优于K2-032，但新增能耗较高；"
                    "K2-051均衡最佳但资源向量被支配。因此不在无权重条件下强行指定单一方案。"
                ),
            },
            {
                "K": 3,
                "是否存在无歧义的综合折中候选": False,
                "说明": (
                    "K3-007同时位于资源前沿和最小新增能耗集合，并优于K3-002的均衡性，"
                    "但其B_max仍明显高；K3-147均衡最佳但资源和新增能耗均不占优。"
                    "因此保留资源—能耗与均衡之间的人工选择。"
                ),
            },
        ]
    )


def validation_table(
    sources: Mapping[str, object],
    augmented: Mapping[int, pd.DataFrame],
    stage1_tables: Mapping[int, pd.DataFrame],
    hashes_before: Mapping[str, str],
) -> pd.DataFrame:
    candidate: pd.DataFrame = sources["candidate"]  # type: ignore[assignment]
    energy_summary: pd.DataFrame = sources["energy_summary"]  # type: ignore[assignment]
    checks = []
    checks.append(("K2 Pareto count is 27", len(augmented[2]) == 27))
    checks.append(("K3 Pareto count is 58", len(augmented[3]) == 58))
    checks.append(("K2 minimum D_R is 1", int(stage1_tables[2]["D_R"].min()) == 1))
    checks.append(("K3 minimum D_R is 2", int(stage1_tables[3]["D_R"].min()) == 2))
    checks.append(("K2-051 is in K2 minimum-gap set", "K2-051" in set(stage1_tables[2]["partition_id"])))
    checks.append(("K3-147 is in K3 minimum-gap set", "K3-147" in set(stage1_tables[3]["partition_id"])))
    known_consistent = True
    for k, partition_id in ((2, "K2-051"), (3, "K3-147"), (2, "K2-041"), (3, "K3-301")):
        source = augmented[k].set_index("partition_id").loc[partition_id]
        known_consistent &= np.isclose(
            float(source["B_max"]),
            float(candidate.loc[partition_id, "B_max"]),
            atol=1e-12,
            rtol=0.0,
        )
        known_consistent &= np.isclose(
            float(source["B_mean"]),
            float(candidate.loc[partition_id, "B_mean"]),
            atol=1e-12,
            rtol=0.0,
        )
        known_consistent &= np.isclose(
            float(source["E_relay_kwh"]),
            float(energy_summary.loc[partition_id, "relay_energy_kwh"]),
            atol=1e-12,
            rtol=0.0,
        )
        known_consistent &= np.isclose(
            float(source["Delta_E_total_kwh"]),
            float(energy_summary.loc[partition_id, "delta_total_energy_kwh"]),
            atol=1e-12,
            rtol=0.0,
        )
    checks.append(("Known candidate balance and energy values match existing workbooks", known_consistent))
    hashes_after = {name: _hash(path) for name, path in SOURCE_PATHS.items()}
    checks.append(("All four source workbooks remain unchanged", hashes_before == hashes_after))
    return pd.DataFrame(
        {
            "check": [name for name, _ in checks],
            "result": ["PASS" if passed else "FAIL" for _, passed in checks],
        }
    )


def format_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    thin = Side(style="thin", color="B7B7B7")
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.sheet_view.showGridLines = False
        for cell in worksheet[1]:
            cell.font = Font(name="宋体", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4472C4")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="宋体", size=10)
                cell.alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                if isinstance(cell.value, float):
                    cell.number_format = "0.000000"
        for column_cells in worksheet.columns:
            contents = [
                str(cell.value) if cell.value is not None else ""
                for cell in column_cells
            ]
            width = min(55, max(10, max(len(value) for value in contents) + 2))
            worksheet.column_dimensions[column_cells[0].column_letter].width = width
        worksheet.page_setup.orientation = "landscape"
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0
        worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    workbook.save(path)


def main() -> None:
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    hashes_before = {name: _hash(path) for name, path in SOURCE_PATHS.items()}
    sources = read_sources()
    baseline: pd.DataFrame = sources["baseline"]  # type: ignore[assignment]
    pareto: Mapping[int, pd.DataFrame] = sources["pareto"]  # type: ignore[assignment]
    augmented = {
        k: augment_pareto(table, baseline) for k, table in pareto.items()
    }
    layered = {k: layered_filter(augmented[k]) for k in (2, 3)}
    stage1 = {k: layered[k][0] for k in (2, 3)}
    resource_front = {k: layered[k][1] for k in (2, 3)}
    execution_front = {k: layered[k][2] for k in (2, 3)}
    category = pd.concat(
        [
            build_category_summary(
                k,
                len(augmented[k]),
                stage1[k],
                resource_front[k],
                execution_front[k],
            )
            for k in (2, 3)
        ],
        ignore_index=True,
    )
    assessment = build_assessment(stage1[2], stage1[3])
    compromise = build_compromise_note()
    validation = validation_table(
        sources, augmented, stage1, hashes_before
    )
    if not validation["result"].eq("PASS").all():
        failed = validation.loc[validation["result"].ne("PASS"), "check"].tolist()
        raise AssertionError(f"Selection analysis validation failed: {failed}")

    methodology = pd.DataFrame(
        [
            {"层级": 1, "名称": "库存缺口", "规则": "分别保留Pareto集中D_R达到最小值的全部方案；同时逐类型报告其他库存缺口。"},
            {"层级": 2, "名称": "资源配置规模", "规则": "在第一层集合中，对10维资源需求向量做无权重非支配过滤。"},
            {"层级": 3, "名称": "额外执行代价", "规则": "比较额外中继任务实例数、中继总能耗和新增总能耗；Cmax只核验。"},
            {"层级": 4, "名称": "工作量均衡", "规则": "比较六个原始均衡指标及B_max、B_mean，不将均衡性设为最高优先级。"},
            {"层级": "结论", "名称": "代表方案措辞", "规则": "仅标记分层候选和权衡关系，不称为全局最优方案，不修改Pareto集合。"},
        ]
    )
    source_audit = pd.DataFrame(
        [
            {
                "source": name,
                "path": str(path),
                "sha256_before": hashes_before[name],
                "sha256_after": _hash(path),
                "unchanged": hashes_before[name] == _hash(path),
            }
            for name, path in SOURCE_PATHS.items()
        ]
    )
    with pd.ExcelWriter(OUTPUT, engine="openpyxl") as writer:
        category.to_excel(writer, sheet_name="layer_summary", index=False)
        stage1[2].reindex(columns=DISPLAY_COLUMNS).to_excel(
            writer, sheet_name="K2_min_gap_Pareto", index=False
        )
        stage1[3].reindex(columns=DISPLAY_COLUMNS).to_excel(
            writer, sheet_name="K3_min_gap_Pareto", index=False
        )
        resource_front[2].reindex(columns=DISPLAY_COLUMNS).to_excel(
            writer, sheet_name="K2_resource_front", index=False
        )
        resource_front[3].reindex(columns=DISPLAY_COLUMNS).to_excel(
            writer, sheet_name="K3_resource_front", index=False
        )
        assessment.to_excel(writer, sheet_name="named_scheme_check", index=False)
        compromise.to_excel(writer, sheet_name="compromise_assessment", index=False)
        methodology.to_excel(writer, sheet_name="methodology", index=False)
        validation.to_excel(writer, sheet_name="validation", index=False)
        source_audit.to_excel(writer, sheet_name="source_audit", index=False)
    format_workbook(OUTPUT)

    for k in (2, 3):
        print("\n" + "=" * 88)
        print(
            f"K={k}: Pareto={len(augmented[k])}, "
            f"D_R_min={int(stage1[k]['D_R'].iloc[0])}, "
            f"minimum-gap solutions={len(stage1[k])}"
        )
        print("=" * 88)
        print(
            stage1[k][
                [
                    "partition_id",
                    *ALL_GAP_COLUMNS,
                    *RESOURCE_COLUMNS,
                    "relay_occurrence_extra_total",
                    "E_relay_kwh",
                    "Delta_E_total_kwh",
                    *BALANCE_SIX,
                    "B_max",
                    "B_mean",
                    "resource_nondominated",
                ]
            ].to_string(index=False)
        )
        print(f"Resource nondominated: {_ids(resource_front[k])}")
        print(f"Execution-cost nondominated within resource front: {_ids(execution_front[k])}")
    print("\n" + category.to_string(index=False))
    print("\n" + assessment.to_string(index=False))
    print("\nValidation:")
    print(validation.to_string(index=False))
    print(f"\noutput={OUTPUT}")


if __name__ == "__main__":
    main()
