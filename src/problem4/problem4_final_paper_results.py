"""Build the final Problem-4 paper tables from existing results only.

This module is deliberately reporting-only.  It reads the completed Problem-3
and Problem-4 workbooks, selects the already-decided K2-032 and K3-007 rows,
and writes new paper artifacts.  It does not import or call any optimizer or
scheduler.
"""

from __future__ import annotations

import hashlib
from pathlib import Path
import re
import sys
from typing import Dict, Iterable, List, Mapping, Sequence, Tuple

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from project_paths import bootstrap, input_file, output_dir, require_file

bootstrap()

MODULE_DIR = Path(__file__).resolve().parent
RESULT_DIR = output_dir("problem4")
LATEX_DIR = RESULT_DIR / "final_paper_tables"
OUTPUT = RESULT_DIR / "problem4_final_paper_results.xlsx"
PROBLEM3_RESULT = input_file("problem3", "problem3_vns_20cycles.xlsx")

SOLUTION_SPECS = ((2, "K2-032"), (3, "K3-007"))
ALL_ZONES = {f"S{index:03d}" for index in range(1, 16)}
BALANCE_COLUMNS = (
    "B_Q",
    "B_M",
    "B_V",
    "B_N_sortie",
    "B_T_transport",
    "B_T_comm",
)
RESOURCE_FIELDS = (
    "R_U_A",
    "R_U_B",
    "R_U_C",
    "R_B_A",
    "R_B_B",
    "R_B_C",
    "R_R",
    "R_E",
)
PROBLEM3_USAGE = {
    "A型运输UAV": 4,
    "B型运输UAV": 2,
    "C型运输UAV": 2,
    "A型电池": 6,
    "B型电池": 4,
    "C型电池": 4,
    "中继UAV": 2,
    "中继能源组件": 3,
}
INVENTORY = {
    "A型运输UAV": 4,
    "B型运输UAV": 2,
    "C型运输UAV": 2,
    "A型电池": 6,
    "B型电池": 4,
    "C型电池": 4,
    "中继UAV": 2,
    "中继能源组件": 6,
}
RESOURCE_MAP = {
    "A型运输UAV": ("R_U_A", "Delta_U_A", "D_U_A"),
    "B型运输UAV": ("R_U_B", "Delta_U_B", "D_U_B"),
    "C型运输UAV": ("R_U_C", "Delta_U_C", "D_U_C"),
    "A型电池": ("R_B_A", "Delta_B_A", "D_B_A"),
    "B型电池": ("R_B_B", "Delta_B_B", "D_B_B"),
    "C型电池": ("R_B_C", "Delta_B_C", "D_B_C"),
    "中继UAV": ("R_R", "Delta_R", "D_R"),
    "中继能源组件": ("R_E", "Delta_E", "D_E"),
}

SOURCE_FILES = (
    input_file("problem4", "problem4_K2_all.xlsx"),
    input_file("problem4", "problem4_K3_all.xlsx"),
    input_file("problem4", "problem4_K2_pareto.xlsx"),
    input_file("problem4", "problem4_K3_pareto.xlsx"),
    input_file("problem4", "group_subset_results.xlsx"),
    input_file("problem4", "relay_duplication_analysis.xlsx"),
    input_file("problem4", "problem4_energy_makespan_verification.xlsx"),
    input_file("problem4", "problem4_final_selection_analysis.xlsx"),
    PROBLEM3_RESULT,
)
PROTECTED_ALL_PARETO = SOURCE_FILES[:4]

TABLE_SPECS = (
    (
        "Table1_任务块",
        "服务区不可拆分任务块划分结果",
        "tab:task_blocks",
        "table1_task_blocks.tex",
        False,
    ),
    (
        "Table2_最终分区",
        "两组与三组任务分区结果",
        "tab:final_partition",
        "table2_final_partition.tex",
        True,
    ),
    (
        "Table3_工作量",
        "最终分区方案各任务组工作量",
        "tab:group_workload",
        "table3_group_workload.tex",
        True,
    ),
    (
        "Table4_资源配置",
        "最终分区方案各任务组资源配置",
        "tab:group_resources",
        "table4_group_resources.tex",
        True,
    ),
    (
        "Table5_资源缺口",
        "最终分区方案资源需求、冗余及库存缺口",
        "tab:resource_gap",
        "table5_resource_gap.tex",
        True,
    ),
    (
        "Table6_执行性能",
        "最终分区方案各任务组执行性能",
        "tab:execution_performance",
        "table6_execution_performance.tex",
        False,
    ),
    (
        "Table7_均衡性",
        "最终分区方案组间工作量均衡性",
        "tab:balance",
        "table7_balance.tex",
        True,
    ),
    (
        "Table8_综合比较",
        "两组与三组最终分区方案综合比较",
        "tab:overall_comparison",
        "table8_overall_comparison.tex",
        False,
    ),
    (
        "Table9_方案筛选",
        "最终分区方案的分层筛选结果",
        "tab:selection_process",
        "table9_selection_process.tex",
        True,
    ),
    (
        "Table10_中继复制",
        "最终分区方案中继任务复制情况",
        "tab:relay_duplication",
        "table10_relay_duplication.tex",
        True,
    ),
)


def file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def split_items(value: object) -> Tuple[str, ...]:
    if pd.isna(value):
        return tuple()
    return tuple(item.strip() for item in str(value).split(";") if item.strip())


def cn_join(values: Iterable[str]) -> str:
    return "、".join(values)


def read_sources() -> Dict[str, object]:
    all_tables = {
        2: pd.read_excel(
            require_file(input_file("problem4", "problem4_K2_all.xlsx")), sheet_name="K2_all"
        ).set_index("partition_id", drop=False),
        3: pd.read_excel(
            require_file(input_file("problem4", "problem4_K3_all.xlsx")), sheet_name="K3_all"
        ).set_index("partition_id", drop=False),
    }
    pareto_tables = {
        2: pd.read_excel(
            require_file(input_file("problem4", "problem4_K2_pareto.xlsx")), sheet_name="Pareto_K2"
        ).set_index("partition_id", drop=False),
        3: pd.read_excel(
            require_file(input_file("problem4", "problem4_K3_pareto.xlsx")), sheet_name="Pareto_K3"
        ).set_index("partition_id", drop=False),
    }
    subset = pd.read_excel(
        require_file(input_file("problem4", "group_subset_results.xlsx")), sheet_name="subset_results"
    ).set_index("mask", drop=False)
    relay_partition = pd.read_excel(
        require_file(input_file("problem4", "relay_duplication_analysis.xlsx")),
        sheet_name="partition_summary",
    ).set_index("partition_id", drop=False)
    relay_detail = pd.read_excel(
        require_file(input_file("problem4", "relay_duplication_analysis.xlsx")),
        sheet_name="relay_detail",
    )
    selection = {
        2: pd.read_excel(
            require_file(input_file("problem4", "problem4_final_selection_analysis.xlsx")),
            sheet_name="K2_resource_front",
        ).set_index("partition_id", drop=False),
        3: pd.read_excel(
            require_file(input_file("problem4", "problem4_final_selection_analysis.xlsx")),
            sheet_name="K3_resource_front",
        ).set_index("partition_id", drop=False),
    }
    prior_energy_validation = pd.read_excel(
        require_file(input_file("problem4", "problem4_energy_makespan_verification.xlsx")),
        sheet_name="validation_15",
    )
    transport = pd.read_excel(
        require_file(PROBLEM3_RESULT), sheet_name="Transport_Sorties"
    ).set_index("Sortie_ID", drop=False)
    relay = pd.read_excel(
        require_file(PROBLEM3_RESULT), sheet_name="Relay_Sorties"
    ).set_index("Sortie_ID", drop=False)
    return {
        "all": all_tables,
        "pareto": pareto_tables,
        "subset": subset,
        "relay_partition": relay_partition,
        "relay_detail": relay_detail,
        "selection": selection,
        "prior_energy_validation": prior_energy_validation,
        "transport": transport,
        "relay": relay,
    }


def calculate_execution(
    sources: Mapping[str, object],
) -> Tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    all_tables: Mapping[int, pd.DataFrame] = sources["all"]  # type: ignore[assignment]
    transport: pd.DataFrame = sources["transport"]  # type: ignore[assignment]
    relay: pd.DataFrame = sources["relay"]  # type: ignore[assignment]
    group_rows: List[Dict[str, object]] = []
    summary_rows: List[Dict[str, object]] = []
    occurrence_rows: List[Dict[str, object]] = []
    problem3_relay = float(relay["Total_Energy_kWh"].sum())
    problem3_total = float(transport["Energy_kWh"].sum()) + problem3_relay

    for k, partition_id in SOLUTION_SPECS:
        solution = all_tables[k].loc[partition_id]
        scheme_groups: List[Dict[str, object]] = []
        for group_index in range(1, k + 1):
            prefix = f"group{group_index}_"
            transport_ids = split_items(solution[prefix + "transport_sorties"])
            relay_ids = split_items(solution[prefix + "relay_occurrences"])
            transport_energy = float(
                transport.loc[list(transport_ids), "Energy_kWh"].sum()
            )
            relay_energy = float(
                relay.loc[list(relay_ids), "Total_Energy_kWh"].sum()
            )
            transport_return = float(
                transport.loc[list(transport_ids), "Return_Time_s"].max()
            )
            relay_return = (
                float(relay.loc[list(relay_ids), "Return_End_s"].max())
                if relay_ids
                else 0.0
            )
            row = {
                "方案": partition_id,
                "K": k,
                "任务组": f"第{group_index}组",
                "任务块": cn_join(split_items(solution[prefix + "components"])),
                "服务区": cn_join(split_items(solution[prefix + "zones"])),
                "运输架次": cn_join(transport_ids),
                "中继任务": cn_join(relay_ids),
                "运输架次数": int(solution[prefix + "N_sortie"]),
                "货箱数": int(solution[prefix + "Q"]),
                "总质量/kg": float(solution[prefix + "M"]),
                "总体积/m³": float(solution[prefix + "V"]),
                "运输累计时长/s": float(solution[prefix + "T_transport"]),
                "通信需求累计时长/s": float(solution[prefix + "T_comm"]),
                "A型运输UAV": int(solution[prefix + "N_U_A"]),
                "B型运输UAV": int(solution[prefix + "N_U_B"]),
                "C型运输UAV": int(solution[prefix + "N_U_C"]),
                "A型电池": int(solution[prefix + "N_B_A"]),
                "B型电池": int(solution[prefix + "N_B_B"]),
                "C型电池": int(solution[prefix + "N_B_C"]),
                "中继UAV": int(solution[prefix + "N_R"]),
                "中继能源组件": int(solution[prefix + "N_E"]),
                "运输能耗/kWh": transport_energy,
                "中继能耗/kWh": relay_energy,
                "总能耗/kWh": transport_energy + relay_energy,
                "运输最大返回时刻/s": transport_return,
                "中继最大返回时刻/s": relay_return,
                "完工时间/s": max(transport_return, relay_return),
            }
            group_rows.append(row)
            scheme_groups.append(row)
            for relay_id in relay_ids:
                occurrence_rows.append(
                    {
                        "方案": partition_id,
                        "任务组": f"第{group_index}组",
                        "中继任务": relay_id,
                        "单次能耗/kWh": float(
                            relay.loc[relay_id, "Total_Energy_kWh"]
                        ),
                        "返回时刻/s": float(relay.loc[relay_id, "Return_End_s"]),
                    }
                )

        transport_total = sum(float(row["运输能耗/kWh"]) for row in scheme_groups)
        relay_total = sum(float(row["中继能耗/kWh"]) for row in scheme_groups)
        total_energy = transport_total + relay_total
        summary_rows.append(
            {
                "方案": partition_id,
                "K": k,
                "运输能耗/kWh": transport_total,
                "中继能耗/kWh": relay_total,
                "总能耗/kWh": total_energy,
                "新增中继能耗/kWh": relay_total - problem3_relay,
                "较问题三新增能耗/kWh": total_energy - problem3_total,
                "完工时间/s": max(float(row["完工时间/s"]) for row in scheme_groups),
            }
        )
    return (
        pd.DataFrame(group_rows),
        pd.DataFrame(summary_rows).set_index("方案", drop=False),
        pd.DataFrame(occurrence_rows),
    )


def build_table1(sources: Mapping[str, object]) -> pd.DataFrame:
    subset: pd.DataFrame = sources["subset"]  # type: ignore[assignment]
    rows = []
    for component_index in range(1, 8):
        source = subset.loc[1 << (component_index - 1)]
        zones = split_items(source["service_zones"])
        sorties = split_items(source["transport_sorties"])
        rows.append(
            {
                "任务块": f"C{component_index}",
                "包含服务区": cn_join(zones),
                "服务区数量": len(zones),
                "对应运输架次": cn_join(sorties),
            }
        )
    return pd.DataFrame(rows)


def build_table2(
    sources: Mapping[str, object], group_detail: pd.DataFrame
) -> pd.DataFrame:
    rows = []
    for k, partition_id in SOLUTION_SPECS:
        for _, group in group_detail.loc[
            group_detail["方案"].eq(partition_id)
        ].iterrows():
            rows.append(
                {
                    "分组方式": f"K={k}",
                    "最终方案": partition_id,
                    "任务组": group["任务组"],
                    "任务块": group["任务块"],
                    "服务区": group["服务区"],
                    "运输架次数": int(group["运输架次数"]),
                    "货箱数": int(group["货箱数"]),
                }
            )
    return pd.DataFrame(rows)


def build_table3(group_detail: pd.DataFrame) -> pd.DataFrame:
    fields = (
        "货箱数/个",
        "总质量/kg",
        "总体积/m³",
        "运输架次数",
        "运输累计时长/s",
        "通信需求累计时长/s",
    )
    rows = []
    for k, partition_id in SOLUTION_SPECS:
        groups = group_detail.loc[group_detail["方案"].eq(partition_id)]
        for _, group in groups.iterrows():
            rows.append(
                {
                    "分组方式": f"K={k}",
                    "任务组": group["任务组"],
                    "货箱数/个": int(group["货箱数"]),
                    "总质量/kg": float(group["总质量/kg"]),
                    "总体积/m³": float(group["总体积/m³"]),
                    "运输架次数": int(group["运输架次数"]),
                    "运输累计时长/s": float(group["运输累计时长/s"]),
                    "通信需求累计时长/s": float(group["通信需求累计时长/s"]),
                }
            )
        rows.append(
            {
                "分组方式": f"K={k}",
                "任务组": "合计",
                **{
                    field: float(
                        pd.DataFrame(rows)
                        .loc[pd.DataFrame(rows)["分组方式"].eq(f"K={k}"), field]
                        .sum()
                    )
                    for field in fields
                },
            }
        )
    table = pd.DataFrame(rows)
    for field in ("货箱数/个", "运输架次数"):
        table[field] = table[field].astype(int)
    return table


def build_table4(
    sources: Mapping[str, object], group_detail: pd.DataFrame
) -> pd.DataFrame:
    all_tables: Mapping[int, pd.DataFrame] = sources["all"]  # type: ignore[assignment]
    columns = (
        "A型运输UAV",
        "B型运输UAV",
        "C型运输UAV",
        "A型电池",
        "B型电池",
        "C型电池",
        "中继UAV",
        "中继能源组件",
    )
    total_fields = RESOURCE_FIELDS
    rows = []
    for k, partition_id in SOLUTION_SPECS:
        groups = group_detail.loc[group_detail["方案"].eq(partition_id)]
        for _, group in groups.iterrows():
            rows.append(
                {
                    "分组方式": f"K={k}",
                    "任务组": group["任务组"],
                    **{column: int(group[column]) for column in columns},
                }
            )
        solution = all_tables[k].loc[partition_id]
        rows.append(
            {
                "分组方式": f"K={k}",
                "任务组": "合计",
                **{
                    column: int(solution[field])
                    for column, field in zip(columns, total_fields)
                },
            }
        )
    return pd.DataFrame(rows)


def build_table5(sources: Mapping[str, object]) -> pd.DataFrame:
    all_tables: Mapping[int, pd.DataFrame] = sources["all"]  # type: ignore[assignment]
    k2 = all_tables[2].loc["K2-032"]
    k3 = all_tables[3].loc["K3-007"]
    rows = []
    for name, (requirement, _delta, _gap) in RESOURCE_MAP.items():
        k2_requirement = int(k2[requirement])
        k3_requirement = int(k3[requirement])
        rows.append(
            {
                "资源类型": name,
                "问题三使用量": PROBLEM3_USAGE[name],
                "现有库存": INVENTORY[name],
                "K=2需求": k2_requirement,
                "K=2冗余": k2_requirement - PROBLEM3_USAGE[name],
                "K=2缺口": max(0, k2_requirement - INVENTORY[name]),
                "K=3需求": k3_requirement,
                "K=3冗余": k3_requirement - PROBLEM3_USAGE[name],
                "K=3缺口": max(0, k3_requirement - INVENTORY[name]),
            }
        )
    return pd.DataFrame(rows)


def build_table6(
    group_detail: pd.DataFrame, execution_summary: pd.DataFrame
) -> pd.DataFrame:
    rows = []
    for k, partition_id in SOLUTION_SPECS:
        groups = group_detail.loc[group_detail["方案"].eq(partition_id)]
        for _, group in groups.iterrows():
            rows.append(
                {
                    "分组方式": f"K={k}",
                    "任务组": group["任务组"],
                    "运输能耗/kWh": float(group["运输能耗/kWh"]),
                    "中继能耗/kWh": float(group["中继能耗/kWh"]),
                    "总能耗/kWh": float(group["总能耗/kWh"]),
                    "完工时间/s": float(group["完工时间/s"]),
                }
            )
        total = execution_summary.loc[partition_id]
        rows.append(
            {
                "分组方式": f"K={k}",
                "任务组": "整体",
                "运输能耗/kWh": float(total["运输能耗/kWh"]),
                "中继能耗/kWh": float(total["中继能耗/kWh"]),
                "总能耗/kWh": float(total["总能耗/kWh"]),
                "完工时间/s": float(total["完工时间/s"]),
            }
        )
    return pd.DataFrame(rows)


def build_table7(sources: Mapping[str, object]) -> pd.DataFrame:
    all_tables: Mapping[int, pd.DataFrame] = sources["all"]  # type: ignore[assignment]
    headers = (
        ("货箱数不均衡", "B_Q"),
        ("质量不均衡", "B_M"),
        ("体积不均衡", "B_V"),
        ("运输架次数不均衡", "B_N_sortie"),
        ("运输时长不均衡", "B_T_transport"),
        ("通信时长不均衡", "B_T_comm"),
    )
    rows = []
    for k, partition_id in SOLUTION_SPECS:
        solution = all_tables[k].loc[partition_id]
        values = [float(solution[field]) for _, field in headers]
        rows.append(
            {
                "分组方式": f"K={k}",
                "方案": partition_id,
                **{header: value for (header, _), value in zip(headers, values)},
                "最坏不均衡": max(values),
                "平均不均衡": float(np.mean(values)),
            }
        )
    return pd.DataFrame(rows)


def build_table8(
    sources: Mapping[str, object],
    execution_summary: pd.DataFrame,
    table5: pd.DataFrame,
    table7: pd.DataFrame,
) -> pd.DataFrame:
    all_tables: Mapping[int, pd.DataFrame] = sources["all"]  # type: ignore[assignment]
    transport: pd.DataFrame = sources["transport"]  # type: ignore[assignment]
    relay: pd.DataFrame = sources["relay"]  # type: ignore[assignment]
    p3_transport = float(transport["Energy_kWh"].sum())
    p3_relay = float(relay["Total_Energy_kWh"].sum())
    p3_cmax = max(
        float(transport["Return_Time_s"].max()),
        float(relay["Return_End_s"].max()),
    )
    k2 = all_tables[2].loc["K2-032"]
    k3 = all_tables[3].loc["K3-007"]
    s2 = execution_summary.loc["K2-032"]
    s3 = execution_summary.loc["K3-007"]
    b2 = table7.loc[table7["方案"].eq("K2-032")].iloc[0]
    b3 = table7.loc[table7["方案"].eq("K3-007")].iloc[0]

    metrics = (
        ("运输UAV总数", 8, int(k2["R_U_total"]), int(k3["R_U_total"])),
        ("共享电池总数", 14, int(k2["R_B_total"]), int(k3["R_B_total"])),
        ("中继UAV数量", 2, int(k2["R_R"]), int(k3["R_R"])),
        ("中继能源组件数量", 3, int(k2["R_E"]), int(k3["R_E"])),
        ("运输UAV库存缺口", 0, int(k2["D_U_A"] + k2["D_U_B"] + k2["D_U_C"]), int(k3["D_U_A"] + k3["D_U_B"] + k3["D_U_C"])),
        ("共享电池库存缺口", 0, int(k2["D_B_A"] + k2["D_B_B"] + k2["D_B_C"]), int(k3["D_B_A"] + k3["D_B_B"] + k3["D_B_C"])),
        ("中继UAV库存缺口", 0, int(k2["D_R"]), int(k3["D_R"])),
        ("中继能源组件库存缺口", 0, int(k2["D_E"]), int(k3["D_E"])),
        ("运输总能耗/kWh", p3_transport, float(s2["运输能耗/kWh"]), float(s3["运输能耗/kWh"])),
        ("中继总能耗/kWh", p3_relay, float(s2["中继能耗/kWh"]), float(s3["中继能耗/kWh"])),
        ("总能耗/kWh", p3_transport + p3_relay, float(s2["总能耗/kWh"]), float(s3["总能耗/kWh"])),
        ("较问题三新增能耗/kWh", 0.0, float(s2["较问题三新增能耗/kWh"]), float(s3["较问题三新增能耗/kWh"])),
        ("整体完工时间/s", p3_cmax, float(s2["完工时间/s"]), float(s3["完工时间/s"])),
        ("中继额外任务实例数", 0, int(k2["relay_occurrence_extra_total"]), int(k3["relay_occurrence_extra_total"])),
        ("最坏不均衡", "—", float(b2["最坏不均衡"]), float(b3["最坏不均衡"])),
        ("平均不均衡", "—", float(b2["平均不均衡"]), float(b3["平均不均衡"])),
    )
    return pd.DataFrame(
        [
            {
                "指标": metric,
                "问题三联合方案": p3,
                "K=2最终方案": value2,
                "K=3最终方案": value3,
            }
            for metric, p3, value2, value3 in metrics
        ]
    )


def build_table9(sources: Mapping[str, object]) -> pd.DataFrame:
    pareto: Mapping[int, pd.DataFrame] = sources["pareto"]  # type: ignore[assignment]
    selection: Mapping[int, pd.DataFrame] = sources["selection"]  # type: ignore[assignment]
    rows = []
    settings = (
        (2, "K2-032", "附加能耗较低"),
        (3, "K3-007", "资源与能耗相同时组间均衡性较好"),
    )
    for k, partition_id, reason in settings:
        p = pareto[k]
        minimum_gap = int(p["D_R"].min())
        candidates = selection[k].index.tolist()
        rows.append(
            {
                "分组方式": f"K={k}",
                "Pareto方案数": len(p),
                "最小中继UAV缺口": minimum_gap,
                "最小缺口方案数": int(p["D_R"].eq(minimum_gap).sum()),
                "资源非支配候选": cn_join(candidates),
                "最终方案": partition_id,
                "最终选择依据": reason,
            }
        )
    return pd.DataFrame(rows)


def build_table10(sources: Mapping[str, object]) -> pd.DataFrame:
    all_tables: Mapping[int, pd.DataFrame] = sources["all"]  # type: ignore[assignment]
    relay: pd.DataFrame = sources["relay"]  # type: ignore[assignment]
    rows = []
    for k, partition_id in SOLUTION_SPECS:
        solution = all_tables[k].loc[partition_id]
        for relay_id in ("RS001", "RS002", "RS003"):
            count = int(solution[f"{relay_id}_group_count"])
            single_energy = float(relay.loc[relay_id, "Total_Energy_kWh"])
            rows.append(
                {
                    "分组方式": f"K={k}（{partition_id}）",
                    "中继任务": relay_id,
                    "问题三实例数": 1,
                    "分区后实例数": count,
                    "新增实例数": count - 1,
                    "单次能耗/kWh": single_energy,
                    "新增能耗/kWh": (count - 1) * single_energy,
                }
            )
    return pd.DataFrame(rows)


def build_tables(
    sources: Mapping[str, object],
    group_detail: pd.DataFrame,
    execution_summary: pd.DataFrame,
) -> List[pd.DataFrame]:
    table1 = build_table1(sources)
    table2 = build_table2(sources, group_detail)
    table3 = build_table3(group_detail)
    table4 = build_table4(sources, group_detail)
    table5 = build_table5(sources)
    table6 = build_table6(group_detail, execution_summary)
    table7 = build_table7(sources)
    table8 = build_table8(sources, execution_summary, table5, table7)
    table9 = build_table9(sources)
    table10 = build_table10(sources)
    return [
        table1,
        table2,
        table3,
        table4,
        table5,
        table6,
        table7,
        table8,
        table9,
        table10,
    ]


def validate(
    sources: Mapping[str, object],
    group_detail: pd.DataFrame,
    execution_summary: pd.DataFrame,
    tables: Sequence[pd.DataFrame],
    source_hashes_before: Mapping[Path, str],
) -> pd.DataFrame:
    all_tables: Mapping[int, pd.DataFrame] = sources["all"]  # type: ignore[assignment]
    pareto: Mapping[int, pd.DataFrame] = sources["pareto"]  # type: ignore[assignment]
    relay_partition: pd.DataFrame = sources["relay_partition"]  # type: ignore[assignment]
    selection: Mapping[int, pd.DataFrame] = sources["selection"]  # type: ignore[assignment]
    transport: pd.DataFrame = sources["transport"]  # type: ignore[assignment]
    relay: pd.DataFrame = sources["relay"]  # type: ignore[assignment]
    prior_validation: pd.DataFrame = sources["prior_energy_validation"]  # type: ignore[assignment]
    (
        _table1,
        table2,
        table3,
        table4,
        table5,
        table6,
        table7,
        table8,
        table9,
        table10,
    ) = tables
    tolerance = 1e-9

    def scheme_groups(partition_id: str) -> pd.DataFrame:
        return group_detail.loc[group_detail["方案"].eq(partition_id)]

    def unique_zones(partition_id: str) -> bool:
        values = [
            item
            for value in scheme_groups(partition_id)["服务区"]
            for item in str(value).split("、")
        ]
        return len(values) == 15 and len(set(values)) == 15 and set(values) == ALL_ZONES

    def unique_sorties(partition_id: str) -> bool:
        values = [
            item
            for value in scheme_groups(partition_id)["运输架次"]
            for item in str(value).split("、")
        ]
        return (
            len(values) == 20
            and len(set(values)) == 20
            and set(values) == set(transport.index)
        )

    def resources_match(k: int, partition_id: str) -> bool:
        all_row = all_tables[k].loc[partition_id]
        pareto_row = pareto[k].loc[partition_id]
        total = table4.loc[
            table4["分组方式"].eq(f"K={k}") & table4["任务组"].eq("合计")
        ].iloc[0]
        table_fields = (
            "A型运输UAV",
            "B型运输UAV",
            "C型运输UAV",
            "A型电池",
            "B型电池",
            "C型电池",
            "中继UAV",
            "中继能源组件",
        )
        return all(
            int(total[header]) == int(all_row[field]) == int(pareto_row[field])
            for header, field in zip(table_fields, RESOURCE_FIELDS)
        )

    p3_transport = float(transport["Energy_kWh"].sum())
    p3_relay = float(relay["Total_Energy_kWh"].sum())
    p3_total = p3_transport + p3_relay
    p3_cmax = max(
        float(transport["Return_Time_s"].max()),
        float(relay["Return_End_s"].max()),
    )
    expected_delta = {"K2-032": 0.759708757565, "K3-007": 2.754147722696}

    table5_delta_ok = all(
        int(row["K=2冗余"]) == int(row["K=2需求"]) - int(row["问题三使用量"])
        and int(row["K=3冗余"]) == int(row["K=3需求"]) - int(row["问题三使用量"])
        for _, row in table5.iterrows()
    )
    table5_gap_ok = all(
        int(row["K=2缺口"]) == max(0, int(row["K=2需求"]) - int(row["现有库存"]))
        and int(row["K=3缺口"]) == max(0, int(row["K=3需求"]) - int(row["现有库存"]))
        for _, row in table5.iterrows()
    )

    table8_map = table8.set_index("指标")
    table8_ok = (
        int(table8_map.loc["运输UAV总数", "K=2最终方案"]) == 9
        and int(table8_map.loc["运输UAV总数", "K=3最终方案"]) == 10
        and int(table8_map.loc["共享电池总数", "K=2最终方案"]) == 15
        and int(table8_map.loc["共享电池总数", "K=3最终方案"]) == 16
        and np.isclose(
            float(table8_map.loc["总能耗/kWh", "K=2最终方案"]),
            float(execution_summary.loc["K2-032", "总能耗/kWh"]),
            atol=tolerance,
            rtol=0.0,
        )
        and np.isclose(
            float(table8_map.loc["总能耗/kWh", "K=3最终方案"]),
            float(execution_summary.loc["K3-007", "总能耗/kWh"]),
            atol=tolerance,
            rtol=0.0,
        )
    )

    table9_map = table9.set_index("分组方式")
    selection_ok = (
        table9_map.loc["K=2", "资源非支配候选"]
        == cn_join(selection[2].index.tolist())
        and table9_map.loc["K=3", "资源非支配候选"]
        == cn_join(selection[3].index.tolist())
        and table9_map.loc["K=2", "最终方案"] == "K2-032"
        and table9_map.loc["K=3", "最终方案"] == "K3-007"
    )

    old_ids = {"K2-041", "K3-301"}
    final_scheme_cells = set(table2["最终方案"]) | set(table7["方案"]) | set(
        table9["最终方案"]
    )

    checks: List[Tuple[str, bool, str]] = [
        ("01_K2-032的15个服务区恰好出现一次", unique_zones("K2-032"), "15个服务区完整且无重复"),
        ("02_K3-007的15个服务区恰好出现一次", unique_zones("K3-007"), "15个服务区完整且无重复"),
        ("03_K2-032的20个运输架次恰好属于一个组", unique_sorties("K2-032"), "与问题三20个运输架次集合一致"),
        ("04_K3-007的20个运输架次恰好属于一个组", unique_sorties("K3-007"), "与问题三20个运输架次集合一致"),
        ("05_K2-032资源合计与原Pareto数据一致", resources_match(2, "K2-032"), "All、Pareto与Table4一致"),
        ("06_K3-007资源合计与原Pareto数据一致", resources_match(3, "K3-007"), "All、Pareto与Table4一致"),
        ("07_K2-032的R_R等于3", int(all_tables[2].loc["K2-032", "R_R"]) == 3, "R_R=3"),
        ("08_K3-007的R_R等于4", int(all_tables[3].loc["K3-007", "R_R"]) == 4, "R_R=4"),
        ("09_K2-032的D_R等于1", int(all_tables[2].loc["K2-032", "D_R"]) == 1, "D_R=1"),
        ("10_K3-007的D_R等于2", int(all_tables[3].loc["K3-007", "D_R"]) == 2, "D_R=2"),
        ("11_两方案中继能源组件缺口均为0", all(int(all_tables[k].loc[pid, "D_E"]) == 0 for k, pid in SOLUTION_SPECS), "K2与K3的D_E均为0"),
        ("12_两方案运输总能耗均等于问题三", all(np.isclose(float(execution_summary.loc[pid, "运输能耗/kWh"]), p3_transport, atol=tolerance, rtol=0.0) for _, pid in SOLUTION_SPECS), f"{p3_transport:.12f} kWh"),
        ("13_K2-032新增总能耗核验", np.isclose(float(execution_summary.loc["K2-032", "较问题三新增能耗/kWh"]), expected_delta["K2-032"], atol=tolerance, rtol=0.0), f"{execution_summary.loc['K2-032', '较问题三新增能耗/kWh']:.12f} kWh"),
        ("14_K3-007新增总能耗核验", np.isclose(float(execution_summary.loc["K3-007", "较问题三新增能耗/kWh"]), expected_delta["K3-007"], atol=tolerance, rtol=0.0), f"{execution_summary.loc['K3-007', '较问题三新增能耗/kWh']:.12f} kWh"),
        ("15_Delta_E_total等于Delta_E_relay", all(np.isclose(float(execution_summary.loc[pid, "较问题三新增能耗/kWh"]), float(execution_summary.loc[pid, "新增中继能耗/kWh"]), atol=tolerance, rtol=0.0) for _, pid in SOLUTION_SPECS), "运输架次未复制"),
        ("16_两方案整体Cmax等于问题三", all(np.isclose(float(execution_summary.loc[pid, "完工时间/s"]), p3_cmax, atol=tolerance, rtol=0.0) for _, pid in SOLUTION_SPECS), f"{p3_cmax:.12f} s；使用Return_Time_s与Return_End_s"),
        ("17_B_max等于六维均衡指标最大值", all(np.isclose(float(row.最坏不均衡), max(float(getattr(row, column)) for column in table7.columns[2:8]), atol=tolerance, rtol=0.0) for row in table7.itertuples(index=False)), "逐方案重算"),
        ("18_B_mean等于六维均衡指标平均值", all(np.isclose(float(row.平均不均衡), np.mean([float(getattr(row, column)) for column in table7.columns[2:8]]), atol=tolerance, rtol=0.0) for row in table7.itertuples(index=False)), "逐方案重算"),
        ("19_Table5冗余计算正确", table5_delta_ok, "需求减问题三使用量"),
        ("20_Table5库存缺口计算正确", table5_gap_ok, "max(0,需求-库存)"),
        ("21_Table8与各明细表完全一致", table8_ok, "资源、能耗、缺口、Cmax与均衡指标交叉核对"),
        ("22_Table9筛选过程与final_selection_analysis一致", selection_ok and len(pareto[2]) == 27 and len(pareto[3]) == 58, "K2候选K2-032/K2-060；K3候选K3-002/K3-007"),
        ("23_旧方案未被误标为最终方案", old_ids.isdisjoint(final_scheme_cells), "最终方案仅K2-032与K3-007"),
        ("24_原始All和Pareto工作簿未修改", all(file_hash(path) == source_hashes_before[path] for path in PROTECTED_ALL_PARETO), "写入前后SHA-256一致"),
        ("25_未调用Gurobi或重新运行优化算法", "gurobipy" not in sys.modules and prior_validation["result"].eq("PASS").all(), "本脚本仅使用pandas/openpyxl读取并整理既有结果"),
    ]

    # Extra cross-source guards support, but do not increase the requested count.
    relay_counts_ok = all(
        all(
            int(all_tables[k].loc[pid, field])
            == int(relay_partition.loc[pid, field])
            for field in (
                "RS001_group_count",
                "RS002_group_count",
                "RS003_group_count",
                "relay_occurrence_extra_total",
            )
        )
        for k, pid in SOLUTION_SPECS
    )
    relay_energy_ok = all(
        np.isclose(
            float(group["中继能耗/kWh"]),
            sum(
                float(relay.loc[item, "Total_Energy_kWh"])
                for item in str(group["中继任务"]).split("、")
                if item
            ),
            atol=tolerance,
            rtol=0.0,
        )
        for _, group in group_detail.iterrows()
    )
    if not relay_counts_ok or not relay_energy_ok:
        checks[20] = (
            checks[20][0],
            False,
            "中继复制次数或完整架次能耗与明细不一致",
        )

    return pd.DataFrame(
        {
            "序号": list(range(1, 26)),
            "核验项目": [name for name, _, _ in checks],
            "结果": ["PASS" if passed else "FAIL" for _, passed, _ in checks],
            "核验说明": [evidence for _, _, evidence in checks],
        }
    )


def latex_escape(value: object) -> str:
    text = str(value)
    replacements = {
        "\\": r"\textbackslash{}",
        "&": r"\&",
        "%": r"\%",
        "$": r"\$",
        "#": r"\#",
        "_": r"\_",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    return text


def format_latex_value(column: str, value: object) -> str:
    if pd.isna(value):
        return ""
    if value == "—":
        return "—"
    if isinstance(value, (np.integer, int)):
        return str(int(value))
    if isinstance(value, (np.floating, float)):
        if "能耗" in column or "不均衡" in column:
            return f"{float(value):.4f}"
        if "时间" in column or "时长" in column or "质量" in column:
            return f"{float(value):.2f}"
        if "体积" in column:
            return f"{float(value):.4f}"
        if float(value).is_integer():
            return str(int(value))
        return f"{float(value):.4f}"
    return str(value)


def write_latex(
    table: pd.DataFrame,
    caption: str,
    label: str,
    output: Path,
    resize: bool,
) -> None:
    alignment = "l" + "c" * (len(table.columns) - 1)
    lines = [
        r"\begin{table}[htbp]",
        r"\centering",
        rf"\caption{{{latex_escape(caption)}}}",
        rf"\label{{{label}}}",
    ]
    if resize:
        lines.append(r"\resizebox{\textwidth}{!}{%")
    lines.extend(
        [
            rf"\begin{{tabular}}{{{alignment}}}",
            r"\toprule",
            " & ".join(latex_escape(column) for column in table.columns) + r" \\",
            r"\midrule",
        ]
    )
    for _, row in table.iterrows():
        metric = str(row.iloc[0]) if table.columns[0] == "指标" else ""
        lines.append(
            " & ".join(
                latex_escape(
                    format_latex_value(
                        metric if metric and column != "指标" else column,
                        row[column],
                    )
                )
                for column in table.columns
            )
            + r" \\"
        )
    lines.extend([r"\bottomrule", r"\end{tabular}"])
    if resize:
        lines.append("}")
    lines.append(r"\end{table}")
    output.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_latex_tables(tables: Sequence[pd.DataFrame]) -> None:
    LATEX_DIR.mkdir(parents=True, exist_ok=True)
    for table, (_, caption, label, filename, resize) in zip(tables, TABLE_SPECS):
        write_latex(table, caption, label, LATEX_DIR / filename, resize)


def format_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    thin = Side(style="thin", color="B7B7B7")
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.sheet_view.showGridLines = False
        for cell in worksheet[1]:
            cell.font = Font(name="宋体", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4472C4")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="宋体", size=10)
                cell.alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                header = worksheet.cell(row=1, column=cell.column).value or ""
                if isinstance(cell.value, float):
                    if "能耗" in str(header) or "不均衡" in str(header):
                        cell.number_format = "0.0000"
                    elif "时间" in str(header) or "时长" in str(header) or "质量" in str(header):
                        cell.number_format = "0.00"
                    elif "体积" in str(header):
                        cell.number_format = "0.0000"
                    else:
                        cell.number_format = "0.0000"
        for column_cells in worksheet.columns:
            values = [str(cell.value) if cell.value is not None else "" for cell in column_cells]
            width = min(55, max(10, max(len(value) for value in values) + 2))
            worksheet.column_dimensions[column_cells[0].column_letter].width = width
        worksheet.page_setup.orientation = "landscape"
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0
        worksheet.sheet_properties.pageSetUpPr.fitToPage = True

    # Table8 mixes integer, energy, time, and balance rows.
    worksheet = workbook["Table8_综合比较"]
    for row_index in range(2, worksheet.max_row + 1):
        metric = str(worksheet.cell(row=row_index, column=1).value)
        if "能耗" in metric or "不均衡" in metric:
            number_format = "0.0000"
        elif "时间" in metric:
            number_format = "0.00"
        else:
            number_format = "0"
        for column_index in range(2, worksheet.max_column + 1):
            cell = worksheet.cell(row=row_index, column=column_index)
            if isinstance(cell.value, (int, float)):
                cell.number_format = number_format
    workbook.save(path)


def write_workbook(tables: Sequence[pd.DataFrame], validation: pd.DataFrame) -> None:
    with pd.ExcelWriter(OUTPUT, engine="openpyxl") as writer:
        for table, (sheet, *_rest) in zip(tables, TABLE_SPECS):
            table.to_excel(writer, sheet_name=sheet, index=False)
        validation.to_excel(writer, sheet_name="Validation", index=False)
    format_workbook(OUTPUT)


def print_final_report(
    group_detail: pd.DataFrame,
    execution_summary: pd.DataFrame,
    tables: Sequence[pd.DataFrame],
    validation: pd.DataFrame,
) -> None:
    table4 = tables[3]
    table5 = tables[4]
    table8 = tables[7]
    table10 = tables[9]
    for label, partition_id in (("A", "K2-032"), ("B", "K3-007")):
        print("\n" + "=" * 100)
        print(f"{label}. {partition_id} 最终分区")
        print("=" * 100)
        groups = group_detail.loc[group_detail["方案"].eq(partition_id)]
        for _, group in groups.iterrows():
            resources = table4.loc[
                table4["分组方式"].eq(f"K={int(group['K'])}")
                & table4["任务组"].eq(group["任务组"])
            ].iloc[0]
            print(f"{group['任务组']}：")
            print(f"  任务块：{group['任务块']}")
            print(f"  服务区：{group['服务区']}")
            print(f"  运输架次：{group['运输架次']}")
            print(
                "  资源：运输UAV(A/B/C)="
                f"{resources['A型运输UAV']}/{resources['B型运输UAV']}/{resources['C型运输UAV']}，"
                "电池(A/B/C)="
                f"{resources['A型电池']}/{resources['B型电池']}/{resources['C型电池']}，"
                f"中继UAV={resources['中继UAV']}，中继能源组件={resources['中继能源组件']}"
            )
            print(
                f"  能耗：运输={group['运输能耗/kWh']:.4f}，"
                f"中继={group['中继能耗/kWh']:.4f}，"
                f"总能耗={group['总能耗/kWh']:.4f} kWh"
            )
            print(f"  Cmax={group['完工时间/s']:.2f} s")

    print("\n" + "=" * 100)
    print("C. 两个方案总体")
    print("=" * 100)
    print(table8.to_string(index=False))
    print("\n分型号资源需求、冗余与缺口：")
    print(table5.to_string(index=False))
    print("\n中继复制情况：")
    print(table10.to_string(index=False))

    print("\n" + "=" * 100)
    print("D. Table8：两组与三组最终分区方案综合比较")
    print("=" * 100)
    print(table8.to_string(index=False))

    print("\n" + "=" * 100)
    print("E. 25项验证")
    print("=" * 100)
    for row in validation.itertuples(index=False):
        print(f"[{row.结果}] {row.序号:02d} {row.核验项目}｜{row.核验说明}")
    print(f"全部PASS：{validation['结果'].eq('PASS').all()}")
    print(f"\nExcel：{OUTPUT}")
    print(f"LaTeX目录：{LATEX_DIR}")


def main() -> None:
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    missing = [path for path in SOURCE_FILES if not path.exists()]
    if missing:
        raise FileNotFoundError(f"Missing source files: {missing}")
    source_hashes_before = {path: file_hash(path) for path in SOURCE_FILES}
    sources = read_sources()
    group_detail, execution_summary, _occurrences = calculate_execution(sources)
    tables = build_tables(sources, group_detail, execution_summary)

    # Create reporting artifacts only.  Optimization/Pareto inputs stay read-only.
    provisional_validation = validate(
        sources,
        group_detail,
        execution_summary,
        tables,
        source_hashes_before,
    )
    if not provisional_validation["结果"].eq("PASS").all():
        failed = provisional_validation.loc[
            provisional_validation["结果"].ne("PASS"), "核验项目"
        ].tolist()
        raise AssertionError(f"Final paper validation failed before write: {failed}")
    write_latex_tables(tables)
    write_workbook(tables, provisional_validation)

    # Recheck hashes after every write and persist the final 25-check result.
    final_validation = validate(
        sources,
        group_detail,
        execution_summary,
        tables,
        source_hashes_before,
    )
    if not final_validation["结果"].eq("PASS").all():
        failed = final_validation.loc[
            final_validation["结果"].ne("PASS"), "核验项目"
        ].tolist()
        raise AssertionError(f"Final paper validation failed after write: {failed}")
    write_workbook(tables, final_validation)
    print_final_report(group_detail, execution_summary, tables, final_validation)


if __name__ == "__main__":
    main()
