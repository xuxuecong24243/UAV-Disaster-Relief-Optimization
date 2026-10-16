"""Build a read-only comparison report from existing Problem-4 workbooks."""

from __future__ import annotations

from pathlib import Path
import sys
from typing import Dict, Iterable, List

import numpy as np
import pandas as pd
from openpyxl import load_workbook


sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from project_paths import bootstrap, input_file, output_dir, require_file

bootstrap()

MODULE_DIR = Path(__file__).resolve().parent
RESULT_DIR = output_dir("problem4")
OUTPUT = RESULT_DIR / "problem4_final_candidate_comparison.xlsx"

K2_IDS = ("K2-041", "K2-051", "K2-060")
K3_IDS = ("K3-147", "K3-301")
CANDIDATE_IDS = (*K2_IDS, *K3_IDS)

BALANCE_COLUMNS = (
    "B_Q",
    "B_M",
    "B_V",
    "B_N_sortie",
    "B_T_transport",
    "B_T_comm",
)
RESOURCE_COLUMNS = (
    "R_U_A",
    "R_U_B",
    "R_U_C",
    "R_U_total",
    "R_B_A",
    "R_B_B",
    "R_B_C",
    "R_B_total",
    "R_R",
    "R_E",
)
GAP_COLUMNS = (
    "D_U_A",
    "D_U_B",
    "D_U_C",
    "D_B_A",
    "D_B_B",
    "D_B_C",
    "D_R",
    "D_E",
)
DELTA_COLUMNS = (
    "Delta_U_A",
    "Delta_U_B",
    "Delta_U_C",
    "Delta_B_A",
    "Delta_B_B",
    "Delta_B_C",
    "Delta_R",
    "Delta_E",
)
RELAY_COLUMNS = (
    "RS001_group_count",
    "RS002_group_count",
    "RS003_group_count",
    "relay_occurrence_extra_total",
)


def _read_sources() -> Dict[str, pd.DataFrame]:
    sources = {
        "K2_all": pd.read_excel(
            require_file(input_file("problem4", "problem4_K2_all.xlsx")), sheet_name="K2_all"
        ),
        "K3_all": pd.read_excel(
            require_file(input_file("problem4", "problem4_K3_all.xlsx")), sheet_name="K3_all"
        ),
        "K2_pareto": pd.read_excel(
            require_file(input_file("problem4", "problem4_K2_pareto.xlsx")), sheet_name="Pareto_K2"
        ),
        "K3_pareto": pd.read_excel(
            require_file(input_file("problem4", "problem4_K3_pareto.xlsx")), sheet_name="Pareto_K3"
        ),
        "K2_representative": pd.read_excel(
            require_file(input_file("problem4", "problem4_representative.xlsx")),
            sheet_name="K2_candidates",
        ),
        "K3_representative": pd.read_excel(
            require_file(input_file("problem4", "problem4_representative.xlsx")),
            sheet_name="K3_candidates",
        ),
    }
    return {
        name: table.set_index("partition_id", drop=False)
        for name, table in sources.items()
    }


def _equal_value(left: object, right: object) -> bool:
    if pd.isna(left) and pd.isna(right):
        return True
    if isinstance(left, (int, float, np.number)) and isinstance(
        right, (int, float, np.number)
    ):
        return bool(np.isclose(float(left), float(right), rtol=0.0, atol=1e-10))
    return str(left) == str(right)


def _verify_sources(sources: Dict[str, pd.DataFrame]) -> pd.DataFrame:
    rows: List[Dict[str, object]] = []
    for candidate_id in CANDIDATE_IDS:
        k = 2 if candidate_id.startswith("K2") else 3
        names = (f"K{k}_all", f"K{k}_pareto", f"K{k}_representative")
        reference = sources[names[0]]
        present = all(candidate_id in sources[name].index for name in names)
        consistent = present
        if present:
            for name in names[1:]:
                comparison = sources[name]
                common = [
                    column
                    for column in reference.columns
                    if column in comparison.columns
                ]
                consistent &= all(
                    _equal_value(
                        reference.loc[candidate_id, column],
                        comparison.loc[candidate_id, column],
                    )
                    for column in common
                )
        rows.append(
            {
                "方案": candidate_id,
                "All文件存在": present,
                "Pareto文件存在": present,
                "representative文件存在": present,
                "三类来源公共字段一致": consistent,
                "核验结果": "PASS" if present and consistent else "FAIL",
            }
        )
    verification = pd.DataFrame(rows)
    if not verification["核验结果"].eq("PASS").all():
        raise AssertionError("Candidate source workbook verification failed.")
    return verification


def _partition_text(row: pd.Series, k: int) -> str:
    return " | ".join(
        f"G{group_index}={{{row[f'group{group_index}_components']}}}"
        for group_index in range(1, k + 1)
    )


def _build_tables(sources: Dict[str, pd.DataFrame]):
    selected_rows = []
    group_rows = []
    for candidate_id in CANDIDATE_IDS:
        k = 2 if candidate_id.startswith("K2") else 3
        source = sources[f"K{k}_all"]
        row = source.loc[candidate_id].copy()
        row["方案"] = candidate_id
        row["K"] = k
        row["任务分区"] = _partition_text(row, k)
        row["B_max"] = float(row[list(BALANCE_COLUMNS)].max())
        row["B_mean"] = float(row[list(BALANCE_COLUMNS)].mean())
        selected_rows.append(row)
        for group_index in range(1, k + 1):
            prefix = f"group{group_index}_"
            group_rows.append(
                {
                    "方案": candidate_id,
                    "K": k,
                    "任务组": f"Group{group_index}",
                    "Components": row[prefix + "components"],
                    "服务区": row[prefix + "zones"],
                    "运输架次数": row[prefix + "N_sortie"],
                    "货箱数量": row[prefix + "Q"],
                    "总质量": row[prefix + "M"],
                    "总体积": row[prefix + "V"],
                    "运输累计时长": row[prefix + "T_transport"],
                    "通信需求累计时长": row[prefix + "T_comm"],
                }
            )

    selected = pd.DataFrame(selected_rows).set_index("方案", drop=False)
    group_columns: List[str] = []
    for group_index in range(1, 4):
        prefix = f"group{group_index}_"
        group_columns.extend(
            [
                prefix + "components",
                prefix + "zones",
                prefix + "N_sortie",
                prefix + "Q",
                prefix + "M",
                prefix + "V",
                prefix + "T_transport",
                prefix + "T_comm",
            ]
        )
    detailed_columns = [
        "方案",
        "K",
        "任务分区",
        *group_columns,
        *RESOURCE_COLUMNS,
        *GAP_COLUMNS,
        *DELTA_COLUMNS,
        *BALANCE_COLUMNS,
        *RELAY_COLUMNS,
        "B_max",
        "B_mean",
    ]
    detailed = selected.reindex(columns=detailed_columns).reset_index(drop=True)
    compact = selected[
        [
            "方案",
            "任务分区",
            "R_U_total",
            "R_B_total",
            "R_R",
            "R_E",
            "D_R",
            "relay_occurrence_extra_total",
            "B_max",
            "B_mean",
        ]
    ].reset_index(drop=True)
    return selected, detailed, pd.DataFrame(group_rows), compact


def _difference_rows(selected: pd.DataFrame) -> pd.DataFrame:
    columns = [
        *RESOURCE_COLUMNS,
        *GAP_COLUMNS,
        *DELTA_COLUMNS,
        *BALANCE_COLUMNS,
        *RELAY_COLUMNS,
        "B_max",
        "B_mean",
    ]
    pairs = (
        ("K2-041", "K2-051"),
        ("K2-051", "K2-041"),
        ("K3-147", "K3-301"),
        ("K3-301", "K3-147"),
    )
    rows = []
    for left, right in pairs:
        values = {
            column: float(selected.loc[left, column])
            - float(selected.loc[right, column])
            for column in columns
        }
        rows.append(
            {
                "比较": f"{left} - {right}",
                "正值含义": f"{left}数值更高",
                "负值含义": f"{left}数值更低",
                **values,
            }
        )
    return pd.DataFrame(rows)


def _format_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for column_cells in worksheet.columns:
            values = [str(cell.value) if cell.value is not None else "" for cell in column_cells]
            width = min(60, max(10, max(len(value) for value in values) + 2))
            worksheet.column_dimensions[column_cells[0].column_letter].width = width
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.000000"
    workbook.save(path)


def main() -> None:
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    sources = _read_sources()
    verification = _verify_sources(sources)
    selected, detailed, group_details, compact = _build_tables(sources)
    differences = _difference_rows(selected)
    with pd.ExcelWriter(OUTPUT, engine="openpyxl") as writer:
        detailed.to_excel(
            writer, sheet_name="candidate_comparison", index=False
        )
        group_details.to_excel(writer, sheet_name="group_details", index=False)
        compact.to_excel(writer, sheet_name="compact_summary", index=False)
        differences.to_excel(
            writer, sheet_name="pairwise_differences", index=False
        )
        verification.to_excel(
            writer, sheet_name="source_verification", index=False
        )
    _format_workbook(OUTPUT)
    print(
        compact.to_string(
            index=False,
            formatters={
                "B_max": lambda value: f"{value:.6f}",
                "B_mean": lambda value: f"{value:.6f}",
            },
        )
    )
    print(f"\noutput={OUTPUT}")


if __name__ == "__main__":
    main()
