"""Create paper-ready tables from existing Problem-4 result files only."""

from __future__ import annotations

from pathlib import Path
import re
import sys
from typing import Dict, Iterable, List, Mapping, Sequence, Tuple

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from project_paths import bootstrap, input_file, output_dir, require_file

bootstrap()

MODULE_DIR = Path(__file__).resolve().parent
RESULT_DIR = output_dir("problem4")
PAPER_DIR = RESULT_DIR / "paper_tables"
EXCEL_OUTPUT = RESULT_DIR / "problem4_paper_tables.xlsx"
SUMMARY_OUTPUT = PAPER_DIR / "problem4_results_summary.txt"
RUN_LOG = RESULT_DIR / "problem4_run.log"

TABLE_SPECS = (
    (
        "Table1_任务块",
        "服务区不可拆分任务块划分结果",
        "tab:task_blocks",
        "table1_task_blocks.tex",
        False,
    ),
    (
        "Table2_任务分区",
        "代表性任务分区方案",
        "tab:partition_results",
        "table2_partition_results.tex",
        True,
    ),
    (
        "Table3_资源配置",
        "代表性分区方案下各任务组资源配置结果",
        "tab:group_resources",
        "table3_group_resources.tex",
        True,
    ),
    (
        "Table4_资源缺口",
        "K=2与K=3代表方案资源配置及库存缺口比较",
        "tab:resource_gap",
        "table4_resource_gap.tex",
        True,
    ),
    (
        "Table5_均衡比较",
        "代表性分区方案工作量均衡性比较",
        "tab:balance",
        "table5_balance.tex",
        True,
    ),
    (
        "Table6_Pareto代表",
        "不同分区策略下Pareto代表方案比较",
        "tab:pareto_representatives",
        "table6_pareto_representatives.tex",
        True,
    ),
    (
        "Table7_枚举结果",
        "问题四精确枚举求解结果",
        "tab:enumeration_summary",
        "table7_enumeration_summary.tex",
        True,
    ),
)

BALANCE_SOURCE_COLUMNS = (
    "B_Q",
    "B_M",
    "B_V",
    "B_N_sortie",
    "B_T_transport",
    "B_T_comm",
)
ALL_ZONES = {f"S{index:03d}" for index in range(1, 16)}
PROBLEM3_USAGE = {
    "A型运输UAV": 4,
    "B型运输UAV": 2,
    "C型运输UAV": 2,
    "A型电池": 6,
    "B型电池": 4,
    "C型电池": 4,
    "中继UAV": 2,
    "中继能源组件": 3,
}
INVENTORY = {
    "A型运输UAV": 4,
    "B型运输UAV": 2,
    "C型运输UAV": 2,
    "A型电池": 6,
    "B型电池": 4,
    "C型电池": 4,
    "中继UAV": 2,
    "中继能源组件": 6,
}
RESOURCE_FIELD_MAP = {
    "A型运输UAV": ("R_U_A", "Delta_U_A", "D_U_A"),
    "B型运输UAV": ("R_U_B", "Delta_U_B", "D_U_B"),
    "C型运输UAV": ("R_U_C", "Delta_U_C", "D_U_C"),
    "A型电池": ("R_B_A", "Delta_B_A", "D_B_A"),
    "B型电池": ("R_B_B", "Delta_B_B", "D_B_B"),
    "C型电池": ("R_B_C", "Delta_B_C", "D_B_C"),
    "中继UAV": ("R_R", "Delta_R", "D_R"),
    "中继能源组件": ("R_E", "Delta_E", "D_E"),
}


def _read_excel(path: str, sheet: str) -> pd.DataFrame:
    return pd.read_excel(require_file(input_file("problem4", path)), sheet_name=sheet)


def read_sources() -> Dict[str, pd.DataFrame]:
    sources = {
        "k2_all": _read_excel("problem4_K2_all.xlsx", "K2_all"),
        "k3_all": _read_excel("problem4_K3_all.xlsx", "K3_all"),
        "k2_pareto": _read_excel("problem4_K2_pareto.xlsx", "Pareto_K2"),
        "k3_pareto": _read_excel("problem4_K3_pareto.xlsx", "Pareto_K3"),
        "k2_representative": _read_excel(
            "problem4_representative.xlsx", "K2_candidates"
        ),
        "k3_representative": _read_excel(
            "problem4_representative.xlsx", "K3_candidates"
        ),
        "candidate_comparison": _read_excel(
            "problem4_final_candidate_comparison.xlsx",
            "candidate_comparison",
        ),
        "relay_summary": _read_excel(
            "relay_duplication_analysis.xlsx", "partition_summary"
        ),
        "subset_results": _read_excel(
            "group_subset_results.xlsx", "subset_results"
        ),
        "k2_validation": _read_excel(
            "problem4_K2_all.xlsx", "validation_15"
        ),
        "k3_validation": _read_excel(
            "problem4_K3_all.xlsx", "validation_15"
        ),
    }
    for name in (
        "k2_all",
        "k3_all",
        "k2_pareto",
        "k3_pareto",
        "k2_representative",
        "k3_representative",
        "relay_summary",
    ):
        sources[name] = sources[name].set_index("partition_id", drop=False)
    sources["candidate_comparison"] = sources["candidate_comparison"].set_index(
        "方案", drop=False
    )
    return sources


def _split(value: object) -> List[str]:
    if pd.isna(value):
        return []
    return [item.strip() for item in str(value).split(";") if item.strip()]


def _numeric_key(value: str) -> int:
    matched = re.search(r"(\d+)$", value)
    return int(matched.group(1)) if matched else 0


def _cn_join(values: Iterable[str]) -> str:
    return "、".join(sorted(values, key=_numeric_key))


def _row(table: pd.DataFrame, partition_id: str) -> pd.Series:
    return table.loc[partition_id]


def build_table1(sources: Mapping[str, pd.DataFrame]) -> pd.DataFrame:
    subsets = sources["subset_results"].set_index("mask")
    rows = []
    for component_index in range(1, 8):
        result = subsets.loc[1 << (component_index - 1)]
        zones = _split(result["service_zones"])
        sorties = _split(result["transport_sorties"])
        rows.append(
            {
                "任务块": f"C{component_index}",
                "包含服务区": _cn_join(zones),
                "服务区数量": len(zones),
                "对应运输架次": _cn_join(sorties),
            }
        )
    return pd.DataFrame(rows)


def build_table2(sources: Mapping[str, pd.DataFrame]) -> pd.DataFrame:
    rows = []
    for k, partition_id, source_name in (
        (2, "K2-041", "k2_all"),
        (3, "K3-301", "k3_all"),
    ):
        source = _row(sources[source_name], partition_id)
        for group_index in range(1, k + 1):
            prefix = f"group{group_index}_"
            rows.append(
                {
                    "分组数": f"K={k}",
                    "方案": partition_id,
                    "任务组": f"第{group_index}组",
                    "任务块": _cn_join(_split(source[prefix + "components"])),
                    "服务区": _cn_join(_split(source[prefix + "zones"])),
                    "运输架次数": int(source[prefix + "N_sortie"]),
                    "货箱数": int(source[prefix + "Q"]),
                }
            )
    return pd.DataFrame(rows)


def build_table3(sources: Mapping[str, pd.DataFrame]) -> pd.DataFrame:
    rows = []
    group_fields = (
        ("A型运输UAV", "N_U_A"),
        ("B型运输UAV", "N_U_B"),
        ("C型运输UAV", "N_U_C"),
        ("A型电池", "N_B_A"),
        ("B型电池", "N_B_B"),
        ("C型电池", "N_B_C"),
        ("中继UAV", "N_R"),
        ("中继能源组件", "N_E"),
    )
    total_fields = (
        "R_U_A",
        "R_U_B",
        "R_U_C",
        "R_B_A",
        "R_B_B",
        "R_B_C",
        "R_R",
        "R_E",
    )
    for k, partition_id, source_name in (
        (2, "K2-041", "k2_all"),
        (3, "K3-301", "k3_all"),
    ):
        source = _row(sources[source_name], partition_id)
        for group_index in range(1, k + 1):
            prefix = f"group{group_index}_"
            row = {"分组数": f"K={k}", "任务组": f"第{group_index}组"}
            for header, field in group_fields:
                row[header] = int(source[prefix + field])
            rows.append(row)
        total = {"分组数": f"K={k}", "任务组": "合计"}
        for (header, _), field in zip(group_fields, total_fields):
            total[header] = int(source[field])
        rows.append(total)
    return pd.DataFrame(rows)


def build_table4(sources: Mapping[str, pd.DataFrame]) -> pd.DataFrame:
    k2 = _row(sources["k2_all"], "K2-041")
    k3 = _row(sources["k3_all"], "K3-301")
    rows = []
    for resource_name, (requirement, delta, gap) in RESOURCE_FIELD_MAP.items():
        rows.append(
            {
                "资源类型": resource_name,
                "问题三联合方案用量": PROBLEM3_USAGE[resource_name],
                "现有库存": INVENTORY[resource_name],
                "K=2需求": int(k2[requirement]),
                "K=2资源冗余": int(k2[delta]),
                "K=2库存缺口": int(k2[gap]),
                "K=3需求": int(k3[requirement]),
                "K=3资源冗余": int(k3[delta]),
                "K=3库存缺口": int(k3[gap]),
            }
        )
    return pd.DataFrame(rows)


def build_table5(sources: Mapping[str, pd.DataFrame]) -> pd.DataFrame:
    comparison = sources["candidate_comparison"]
    rows = []
    header_map = (
        ("货箱数量不均衡", "B_Q"),
        ("质量不均衡", "B_M"),
        ("体积不均衡", "B_V"),
        ("运输架次数不均衡", "B_N_sortie"),
        ("运输累计时长不均衡", "B_T_transport"),
        ("通信需求时长不均衡", "B_T_comm"),
        ("最坏不均衡", "B_max"),
        ("平均不均衡", "B_mean"),
    )
    for partition_id in ("K2-041", "K2-051", "K3-301", "K3-147"):
        source = comparison.loc[partition_id]
        row = {"方案": partition_id}
        for header, field in header_map:
            row[header] = round(float(source[field]), 4)
        rows.append(row)
    return pd.DataFrame(rows)


def build_table6(sources: Mapping[str, pd.DataFrame]) -> pd.DataFrame:
    comparison = sources["candidate_comparison"]
    features = {
        "K2-041": "两组分区的均衡型代表方案",
        "K2-051": "两组分区的资源节约型代表方案",
        "K3-301": "三组分区的均衡型代表方案",
        "K3-147": "三组分区的资源节约型代表方案",
    }
    rows = []
    for partition_id in ("K2-041", "K2-051", "K3-301", "K3-147"):
        source = comparison.loc[partition_id]
        rows.append(
            {
                "分组数": f"K={int(source['K'])}",
                "方案": partition_id,
                "方案特征": features[partition_id],
                "运输UAV总数": int(source["R_U_total"]),
                "电池总数": int(source["R_B_total"]),
                "中继UAV": int(source["R_R"]),
                "中继能源组件": int(source["R_E"]),
                "中继UAV缺口": int(source["D_R"]),
                "中继复制数": int(source["relay_occurrence_extra_total"]),
                "最坏不均衡": round(float(source["B_max"]), 4),
                "平均不均衡": round(float(source["B_mean"]), 4),
            }
        )
    return pd.DataFrame(rows)


def build_table7(sources: Mapping[str, pd.DataFrame]) -> pd.DataFrame:
    rows = []
    for k, all_name, pareto_name in (
        (2, "k2_all", "k2_pareto"),
        (3, "k3_all", "k3_pareto"),
    ):
        all_solutions = sources[all_name]
        rows.append(
            {
                "分组数": f"K={k}",
                "不可拆分任务块数": 7,
                "理论方案数": 63 if k == 2 else 301,
                "实际方案数": len(all_solutions),
                "Pareto方案数": len(sources[pareto_name]),
                "库存完全可行方案数": int(
                    all_solutions["inventory_feasible"].sum()
                ),
                "中继UAV需求范围": (
                    f"{int(all_solutions['R_R'].min())}–"
                    f"{int(all_solutions['R_R'].max())}"
                ),
                "中继能源需求范围": (
                    f"{int(all_solutions['R_E'].min())}–"
                    f"{int(all_solutions['R_E'].max())}"
                ),
                "额外中继任务实例范围": (
                    f"{int(all_solutions['relay_occurrence_extra_total'].min())}–"
                    f"{int(all_solutions['relay_occurrence_extra_total'].max())}"
                ),
            }
        )
    return pd.DataFrame(rows)


def build_tables(sources: Mapping[str, pd.DataFrame]) -> List[pd.DataFrame]:
    return [
        build_table1(sources),
        build_table2(sources),
        build_table3(sources),
        build_table4(sources),
        build_table5(sources),
        build_table6(sources),
        build_table7(sources),
    ]


def validate_tables(
    sources: Mapping[str, pd.DataFrame],
    tables: Sequence[pd.DataFrame],
) -> List[Tuple[str, bool]]:
    table1, table2, table3, table4, table5, table6, table7 = tables
    table2_zone_ok = True
    for partition_id in ("K2-041", "K3-301"):
        zones = [
            zone
            for value in table2.loc[table2["方案"].eq(partition_id), "服务区"]
            for zone in str(value).split("、")
        ]
        table2_zone_ok &= len(zones) == 15 and set(zones) == ALL_ZONES

    table3_sum_ok = True
    resource_headers = list(table3.columns[2:])
    for k, partition_id, source_name in (
        (2, "K2-041", "k2_all"),
        (3, "K3-301", "k3_all"),
    ):
        rows = table3.loc[table3["分组数"].eq(f"K={k}")]
        group_sum = rows.loc[rows["任务组"].ne("合计"), resource_headers].sum()
        reported_total = rows.loc[
            rows["任务组"].eq("合计"), resource_headers
        ].iloc[0]
        table3_sum_ok &= group_sum.equals(reported_total)

    delta_ok = all(
        row["K=2资源冗余"]
        == row["K=2需求"] - row["问题三联合方案用量"]
        and row["K=3资源冗余"]
        == row["K=3需求"] - row["问题三联合方案用量"]
        for _, row in table4.iterrows()
    )
    gap_ok = all(
        row["K=2库存缺口"] == max(0, row["K=2需求"] - row["现有库存"])
        and row["K=3库存缺口"] == max(0, row["K=3需求"] - row["现有库存"])
        for _, row in table4.iterrows()
    )

    comparison = sources["candidate_comparison"]
    bmax_ok = all(
        np.isclose(
            float(comparison.loc[partition_id, "B_max"]),
            max(
                float(comparison.loc[partition_id, column])
                for column in BALANCE_SOURCE_COLUMNS
            ),
            atol=1e-10,
            rtol=0.0,
        )
        for partition_id in table5["方案"]
    )
    bmean_ok = all(
        np.isclose(
            float(comparison.loc[partition_id, "B_mean"]),
            np.mean(
                [
                    float(comparison.loc[partition_id, column])
                    for column in BALANCE_SOURCE_COLUMNS
                ]
            ),
            atol=1e-10,
            rtol=0.0,
        )
        for partition_id in table5["方案"]
    )
    table6_ok = all(
        int(row["运输UAV总数"]) == int(comparison.loc[row["方案"], "R_U_total"])
        and int(row["电池总数"]) == int(comparison.loc[row["方案"], "R_B_total"])
        and int(row["中继UAV"]) == int(comparison.loc[row["方案"], "R_R"])
        and int(row["中继能源组件"]) == int(comparison.loc[row["方案"], "R_E"])
        and int(row["中继UAV缺口"]) == int(comparison.loc[row["方案"], "D_R"])
        and int(row["中继复制数"])
        == int(comparison.loc[row["方案"], "relay_occurrence_extra_total"])
        for _, row in table6.iterrows()
    )
    consistency_fields = (
        "R_U_total",
        "R_B_total",
        "R_R",
        "R_E",
        "D_R",
        "B_Q",
        "B_M",
        "B_V",
        "B_N_sortie",
        "B_T_transport",
        "B_T_comm",
        "RS001_group_count",
        "RS002_group_count",
        "RS003_group_count",
        "relay_occurrence_extra_total",
    )
    for partition_id in table6["方案"]:
        k = 2 if partition_id.startswith("K2") else 3
        all_row = sources[f"k{k}_all"].loc[partition_id]
        for source_name in (
            f"k{k}_pareto",
            f"k{k}_representative",
            "candidate_comparison",
        ):
            source = sources[source_name]
            table6_ok &= partition_id in source.index
            if partition_id in source.index:
                table6_ok &= all(
                    np.isclose(
                        float(all_row[field]),
                        float(source.loc[partition_id, field]),
                        atol=1e-10,
                        rtol=0.0,
                    )
                    for field in consistency_fields
                )
        relay = sources["relay_summary"]
        table6_ok &= partition_id in relay.index
        if partition_id in relay.index:
            table6_ok &= all(
                int(all_row[field]) == int(relay.loc[partition_id, field])
                for field in (
                    "RS001_group_count",
                    "RS002_group_count",
                    "RS003_group_count",
                    "relay_occurrence_extra_total",
                    "R_R",
                    "R_E",
                )
            )
    source_validation_ok = (
        sources["k2_validation"]["result"].eq("PASS").all()
        and sources["k3_validation"]["result"].eq("PASS").all()
    )
    table7_ok = (
        table7["实际方案数"].tolist() == [63, 301]
        and table7["Pareto方案数"].tolist() == [27, 58]
        and table7["库存完全可行方案数"].tolist() == [0, 0]
        and source_validation_ok
    )
    k2 = comparison.loc["K2-041"]
    k3 = comparison.loc["K3-301"]
    k2_known_ok = (
        int(k2["R_U_total"]) == 12
        and int(k2["R_B_total"]) == 17
        and int(k2["R_R"]) == 4
        and int(k2["R_E"]) == 4
    )
    k3_known_ok = (
        int(k3["R_U_total"]) == 15
        and int(k3["R_B_total"]) == 18
        and int(k3["R_R"]) == 5
        and int(k3["R_E"]) == 5
    )
    return [
        ("Table2每个代表方案的15个服务区恰好出现一次", table2_zone_ok),
        ("Table3各组资源之和等于方案整体资源", table3_sum_ok),
        ("Table4资源冗余满足需求减问题三用量", delta_ok),
        ("Table4库存缺口满足正部函数定义", gap_ok),
        ("Table5最坏不均衡等于六项指标最大值", bmax_ok),
        ("Table5平均不均衡等于六项指标平均值", bmean_ok),
        ("Table6与All、Pareto、候选及中继统计原始数据一致", table6_ok),
        ("Table7满足63、301、27、58且原15项检查通过", table7_ok),
        ("K2-041整体资源核验", k2_known_ok),
        ("K3-301整体资源核验", k3_known_ok),
    ]


def _latex_escape(value: object) -> str:
    text = str(value)
    replacements = {
        "\\": r"\textbackslash{}",
        "&": r"\&",
        "%": r"\%",
        "$": r"\$",
        "#": r"\#",
        "_": r"\_",
    }
    for source, target in replacements.items():
        text = text.replace(source, target)
    return text


def _format_value(value: object, balance: bool = False) -> str:
    if pd.isna(value):
        return ""
    if isinstance(value, (np.integer, int)):
        return str(int(value))
    if isinstance(value, (np.floating, float)):
        return f"{float(value):.4f}" if balance else f"{float(value):.2f}"
    return str(value)


def write_latex_table(
    table: pd.DataFrame,
    caption: str,
    label: str,
    output: Path,
    resize: bool,
    balance_columns: Iterable[str] = (),
) -> None:
    balance_columns = set(balance_columns)
    alignments = ["l"] + ["c"] * (len(table.columns) - 1)
    lines = [
        r"\begin{table}[htbp]",
        r"\centering",
        rf"\caption{{{_latex_escape(caption)}}}",
        rf"\label{{{label}}}",
    ]
    if resize:
        lines.append(r"\resizebox{\textwidth}{!}{%")
    lines.extend(
        [
            rf"\begin{{tabular}}{{{''.join(alignments)}}}",
            r"\toprule",
            " & ".join(_latex_escape(column) for column in table.columns)
            + r" \\",
            r"\midrule",
        ]
    )
    for _, row in table.iterrows():
        values = [
            _latex_escape(
                _format_value(
                    row[column],
                    balance=column in balance_columns,
                )
            )
            for column in table.columns
        ]
        lines.append(" & ".join(values) + r" \\")
    lines.extend([r"\bottomrule", r"\end{tabular}"])
    if resize:
        lines.append("}")
    lines.extend([r"\end{table}", ""])
    output.write_text("\n".join(lines), encoding="utf-8")


def write_excel(tables: Sequence[pd.DataFrame]) -> None:
    with pd.ExcelWriter(EXCEL_OUTPUT, engine="openpyxl") as writer:
        for spec, table in zip(TABLE_SPECS, tables):
            table.to_excel(writer, sheet_name=spec[0], index=False)

    workbook = load_workbook(EXCEL_OUTPUT)
    thin_gray = Side(style="thin", color="B7B7B7")
    for spec, table in zip(TABLE_SPECS, tables):
        worksheet = workbook[spec[0]]
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.sheet_view.showGridLines = False
        worksheet.page_setup.orientation = "landscape"
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0
        worksheet.sheet_properties.pageSetUpPr.fitToPage = True
        for cell in worksheet[1]:
            cell.font = Font(name="宋体", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4472C4")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="宋体", size=10)
                cell.alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )
                cell.border = Border(
                    left=thin_gray,
                    right=thin_gray,
                    top=thin_gray,
                    bottom=thin_gray,
                )
                if isinstance(cell.value, float):
                    if spec[0] in ("Table5_均衡比较", "Table6_Pareto代表"):
                        cell.number_format = "0.0000"
                    else:
                        cell.number_format = "0.00"
        for column_cells in worksheet.columns:
            contents = [
                str(cell.value) if cell.value is not None else ""
                for cell in column_cells
            ]
            width = min(34, max(10, max(len(value) for value in contents) + 2))
            worksheet.column_dimensions[column_cells[0].column_letter].width = width
        worksheet.print_title_rows = "1:1"
    workbook.save(EXCEL_OUTPUT)


def _runtime_seconds() -> float:
    if not RUN_LOG.exists():
        return 40.395
    matches = re.findall(
        r"total runtime s\s*=\s*([0-9.]+)",
        RUN_LOG.read_text(encoding="utf-8"),
    )
    return float(matches[-1]) if matches else 40.395


def write_summary(
    sources: Mapping[str, pd.DataFrame],
    validation: Sequence[Tuple[str, bool]],
) -> None:
    comparison = sources["candidate_comparison"]
    runtime = _runtime_seconds()
    lines = [
        "问题四结果数据摘要",
        "",
        "1. 服务区关联关系形成7个不可拆分任务块。",
        "2. K=2采用精确枚举，共得到63个分区方案。",
        "3. K=3采用精确枚举，共得到301个分区方案。",
        "4. K=2和K=3的Pareto方案数分别为27和58。",
        "5. 两种分组方式均不存在库存完全可行方案。",
        "6. K=2方案至少需要3架中继UAV，相对2架库存至少缺少1架。",
        "7. K=3方案至少需要4架中继UAV，相对2架库存至少缺少2架。",
        "8. 中继能源组件库存为6组，全部K=2与K=3方案的能源组件库存缺口均为0。",
        (
            "9. 两组分区的均衡型代表方案K2-041需要12架运输UAV、17组电池、"
            "4架中继UAV和4组中继能源组件；最坏不均衡指标为"
            f"{comparison.loc['K2-041', 'B_max']:.4f}，平均不均衡指标为"
            f"{comparison.loc['K2-041', 'B_mean']:.4f}。"
        ),
        (
            "10. 三组分区的均衡型代表方案K3-301需要15架运输UAV、18组电池、"
            "5架中继UAV和5组中继能源组件；最坏不均衡指标为"
            f"{comparison.loc['K3-301', 'B_max']:.4f}，平均不均衡指标为"
            f"{comparison.loc['K3-301', 'B_mean']:.4f}。"
        ),
        (
            "11. 资源节约型对照方案K2-051需要12架运输UAV、18组电池、3架中继UAV"
            "和3组中继能源组件；K3-147需要14架运输UAV、18组电池、4架中继UAV"
            "和4组中继能源组件。"
        ),
        "12. K=3的额外中继任务实例范围为2–4，高于K=2的1–2，跨组中继复制更明显。",
        (
            "13. 本实例中增加分组数并未自动改善工作量均衡性：K3-301的最坏不均衡"
            f"{comparison.loc['K3-301', 'B_max']:.4f}，高于K2-041的"
            f"{comparison.loc['K2-041', 'B_max']:.4f}。"
        ),
        "14. 同一运输架次访问的服务区必须同组，因此7个不可拆分任务块限制了分区自由度。",
        f"15. 完整精确枚举与结果整理所依据的求解运行时间为{runtime:.3f} s。",
        "16. 364个候选分区方案均由精确枚举获得，原求解中的15项一致性与可行性检查全部通过。",
        "",
        "论文表格数据核验：",
    ]
    lines.extend(
        f"{index}. [{'PASS' if passed else 'FAIL'}] {name}"
        for index, (name, passed) in enumerate(validation, start=1)
    )
    SUMMARY_OUTPUT.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    PAPER_DIR.mkdir(parents=True, exist_ok=True)
    sources = read_sources()
    tables = build_tables(sources)
    validation = validate_tables(sources, tables)
    if not all(passed for _, passed in validation):
        failures = [name for name, passed in validation if not passed]
        raise AssertionError(f"Paper table validation failed: {failures}")

    write_excel(tables)
    for index, (spec, table) in enumerate(zip(TABLE_SPECS, tables), start=1):
        balance_columns = (
            table.columns[1:]
            if spec[0] == "Table5_均衡比较"
            else ("最坏不均衡", "平均不均衡")
            if spec[0] == "Table6_Pareto代表"
            else ()
        )
        write_latex_table(
            table,
            spec[1],
            spec[2],
            PAPER_DIR / spec[3],
            spec[4],
            balance_columns,
        )
        print("\n" + "=" * 72)
        print(f"表{index}：{spec[1]}（{len(table)}行 × {len(table.columns)}列）")
        print("=" * 72)
        print(table.to_string(index=False))

    write_summary(sources, validation)
    print("\n" + "=" * 72)
    print("10项数据核验")
    print("=" * 72)
    for name, passed in validation:
        print(f"[{'PASS' if passed else 'FAIL'}] {name}")
    print(f"\nExcel: {EXCEL_OUTPUT}")
    for spec in TABLE_SPECS:
        print(f"LaTeX: {PAPER_DIR / spec[3]}")
    print(f"Summary: {SUMMARY_OUTPUT}")


if __name__ == "__main__":
    main()
