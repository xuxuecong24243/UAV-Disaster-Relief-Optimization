"""问题1：单点往返条件下的货箱组批与机型选择字典序优化。"""

from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
import sys

# Keep historical module names while resolving all locations from this repository.
_RELEASE_ROOT = Path(__file__).resolve().parents[2]
if str(_RELEASE_ROOT) not in sys.path:
    sys.path.insert(0, str(_RELEASE_ROOT))
from project_paths import ROOT, bootstrap, raw_file, dem_file, input_file, output_dir, output_file, require_file, assert_writable_output

bootstrap()
from time import perf_counter
from typing import Dict, List, Tuple

import gurobipy as gp
from gurobipy import GRB
import numpy as np
from openpyxl import load_workbook
import pandas as pd

from problem1_max_safe_payload import (
    ENERGY_TOLERANCE_KWH,
    UAVModel,
    load_uav_parameters as load_energy_models,
    round_trip_energy,
)


PROJECT_DIR = ROOT
DATA_DIR = raw_file("运输无人机数据.xlsx").parent
BOX_FILE = raw_file('物资需求与配送时限.xlsx')
UAV_FILE = raw_file('运输无人机数据.xlsx')
RESULT_DIR = output_dir("problem1")
SEGMENT_FILE = input_file("problem1", 'segment_preprocessing.xlsx')
MAX_PAYLOAD_FILE = input_file("problem1", 'problem1_max_safe_payload.xlsx')
CANDIDATE_FILE = RESULT_DIR / "problem1_batch_candidates.xlsx"
OPTIMIZATION_FILE = RESULT_DIR / "problem1_batch_optimization.xlsx"

DEPOT = "O01"
SERVICE_ZONES = [f"S{i:03d}" for i in range(1, 16)]
SEGMENT_NODES = [DEPOT] + [f"S{i:02d}" for i in range(1, 16)]
UAV_TYPES = ["A", "B", "C"]
CAPACITY_TOLERANCE = 1e-9
OBJECTIVE_TOLERANCE_KWH = 1e-7
TIME_TOLERANCE_S = 1e-6


@dataclass(frozen=True)
class UAVParameters:
    """组批阶段使用的机型容量、时间、速度及能耗参数。"""

    uav_type: str
    model_name: str
    max_volume_m3: float
    cruise_speed_m_s: float
    prep_time_s: float
    load_time_per_box_s: float
    handover_base_time_s: float
    handover_time_per_box_s: float
    climb_speed_m_s: float
    descent_speed_m_s: float
    descent_efficiency: float
    energy_model: UAVModel


@dataclass(frozen=True)
class OptimizationResult:
    """三阶段字典序求解结果。"""

    selected_indices: Tuple[int, ...]
    minimum_sorties: int
    minimum_energy_kwh: float
    minimum_operation_time_s: float
    stage1_runtime_s: float
    stage2_runtime_s: float
    stage3_runtime_s: float
    solver_status: str


def to_segment_zone(service_zone: str) -> str:
    """将附件的 S001-S015 映射到航段文件的 S01-S15。"""
    return f"S{int(str(service_zone).strip()[1:]):02d}"


def to_box_zone(segment_zone: str) -> str:
    """将航段文件的 S01-S15 映射到附件的 S001-S015。"""
    return f"S{int(str(segment_zone).strip()[1:]):03d}"


def load_box_data(box_file: Path) -> pd.DataFrame:
    """读取附件的逐箱货箱清单并统一为求解字段。"""
    sheet_name = "逐箱货箱清单"
    raw = pd.read_excel(require_file(box_file), sheet_name=sheet_name)
    column_map = {
        "货箱编号": "box_id",
        "服务区编号": "service_zone",
        "物资类型": "material_type",
        "单箱质量（kg）": "weight_kg",
        "单箱体积（m³）": "volume_m3",
        "是否首批保障": "is_first_batch",
        "首批截止时间（s）": "first_batch_deadline_s",
        "期望送达时间（s）": "expected_arrival_time_s",
        "应急优先系数": "emergency_priority",
    }
    missing = [column for column in column_map if column not in raw.columns]
    if missing:
        raise ValueError("逐箱货箱清单缺少字段: " + ", ".join(missing))

    boxes = raw[list(column_map)].rename(columns=column_map).copy()
    boxes["box_id"] = boxes["box_id"].astype(str).str.strip()
    boxes["service_zone"] = boxes["service_zone"].astype(str).str.strip().str.upper()
    boxes["material_type"] = boxes["material_type"].astype(str).str.strip()
    boxes["is_first_batch"] = boxes["is_first_batch"].map({"是": True, "否": False})
    numeric_columns = [
        "weight_kg",
        "volume_m3",
        "first_batch_deadline_s",
        "expected_arrival_time_s",
        "emergency_priority",
    ]
    boxes[numeric_columns] = boxes[numeric_columns].apply(
        pd.to_numeric, errors="coerce"
    )

    if len(boxes) != 80 or boxes["box_id"].nunique() != 80:
        raise ValueError("逐箱货箱清单应包含 80 个不重复货箱")
    if not boxes["service_zone"].isin(SERVICE_ZONES).all():
        raise ValueError("逐箱货箱清单中存在 S001-S015 之外的服务区")
    if boxes["is_first_batch"].isna().any():
        raise ValueError("是否首批保障字段存在无法识别的取值")
    if boxes[["weight_kg", "volume_m3", "expected_arrival_time_s"]].isna().any().any():
        raise ValueError("货箱质量、体积或期望送达时间存在空值")

    boxes["is_first_batch"] = boxes["is_first_batch"].astype(bool)
    zone_order = {zone: index for index, zone in enumerate(SERVICE_ZONES)}
    boxes["_zone_order"] = boxes["service_zone"].map(zone_order)
    boxes = boxes.sort_values(["_zone_order"], kind="stable").drop(
        columns="_zone_order"
    )
    return boxes.reset_index(drop=True)


def load_uav_parameters(uav_file: Path) -> Dict[str, UAVParameters]:
    """读取附件三类机型参数，并复用最大安全载荷程序的能耗参数。"""
    energy_models = load_energy_models(uav_file)
    raw = pd.read_excel(require_file(uav_file), sheet_name="数据", header=1)
    column_map = {
        "机型编号": "uav_type",
        "机型名称": "model_name",
        "可用装载体积（m³）": "max_volume_m3",
        "计划巡航速度（m/s）": "cruise_speed_m_s",
        "工位固定准备时间（s）": "prep_time_s",
        "每箱装载时间（s）": "load_time_per_box_s",
        "接收点基础交接时间（s）": "handover_base_time_s",
        "每箱增加交接时间（s）": "handover_time_per_box_s",
        "最大爬升速度（m/s）": "climb_speed_m_s",
        "最大下降速度（m/s）": "descent_speed_m_s",
        "下降能耗效率": "descent_efficiency",
    }
    missing = [column for column in column_map if column not in raw.columns]
    if missing:
        raise ValueError("运输无人机数据缺少字段: " + ", ".join(missing))

    table = raw[list(column_map)].rename(columns=column_map)
    table["uav_type"] = table["uav_type"].astype(str).str.strip().str.upper()
    numeric_columns = [column for column in column_map.values() if column not in {
        "uav_type", "model_name"
    }]
    table = table.loc[
        table["uav_type"].isin(UAV_TYPES)
        & table["model_name"].notna()
        & table[numeric_columns].notna().all(axis=1)
    ].copy()
    table[numeric_columns] = table[numeric_columns].apply(
        pd.to_numeric, errors="raise"
    )
    table = table.set_index("uav_type").loc[UAV_TYPES]

    parameters: Dict[str, UAVParameters] = {}
    for uav_type in UAV_TYPES:
        row = table.loc[uav_type]
        parameters[uav_type] = UAVParameters(
            uav_type=uav_type,
            model_name=str(row["model_name"]).strip(),
            max_volume_m3=float(row["max_volume_m3"]),
            cruise_speed_m_s=float(row["cruise_speed_m_s"]),
            prep_time_s=float(row["prep_time_s"]),
            load_time_per_box_s=float(row["load_time_per_box_s"]),
            handover_base_time_s=float(row["handover_base_time_s"]),
            handover_time_per_box_s=float(row["handover_time_per_box_s"]),
            climb_speed_m_s=float(row["climb_speed_m_s"]),
            descent_speed_m_s=float(row["descent_speed_m_s"]),
            descent_efficiency=float(row["descent_efficiency"]),
            energy_model=energy_models[uav_type],
        )
    return parameters


def load_segment_data(segment_file: Path) -> Dict[str, pd.DataFrame]:
    """读取距离、爬升和下降矩阵，不重复计算 DEM。"""
    matrices = {}
    for sheet_name in (
        "distance_matrix",
        "climb_height_matrix",
        "descent_height_matrix",
    ):
        matrix = pd.read_excel(require_file(segment_file), sheet_name=sheet_name, index_col=0)
        matrix.index = matrix.index.astype(str).str.strip()
        matrix.columns = matrix.columns.astype(str).str.strip()
        if (
            matrix.index.tolist() != SEGMENT_NODES
            or matrix.columns.tolist() != SEGMENT_NODES
        ):
            raise ValueError(f"{sheet_name} 的节点顺序或名称不正确")
        matrices[sheet_name] = matrix.astype(float)
    return matrices


def load_max_safe_payload(max_payload_file: Path) -> Dict[Tuple[str, str], float]:
    """从已有结果读取 q_gi^max，不在本程序中重新求解。"""
    table = pd.read_excel(require_file(max_payload_file), sheet_name="max_safe_payload")
    required = {"uav_type", "service_zone", "max_payload_kg", "feasible"}
    missing = required.difference(table.columns)
    if missing:
        raise ValueError("最大安全载荷结果缺少字段: " + ", ".join(sorted(missing)))

    table = table.loc[table["feasible"]].copy()
    payloads = {
        (str(row.uav_type).strip().upper(), to_box_zone(row.service_zone)): float(
            row.max_payload_kg
        )
        for row in table.itertuples(index=False)
    }
    expected_keys = {
        (uav_type, service_zone)
        for uav_type in UAV_TYPES
        for service_zone in SERVICE_ZONES
    }
    if set(payloads) != expected_keys:
        raise ValueError("最大安全载荷结果未包含全部 3×15 个可行组合")
    return payloads


def segment_flight_time(
    distance_m: float,
    climb_height_m: float,
    descent_height_m: float,
    uav: UAVParameters,
) -> float:
    """计算一个有向航段的爬升、巡航和下降总时间。"""
    return (
        climb_height_m / uav.climb_speed_m_s
        + distance_m / uav.cruise_speed_m_s
        + descent_height_m / uav.descent_speed_m_s
    )


def round_trip_flight_time(
    service_zone: str,
    uav: UAVParameters,
    matrices: Dict[str, pd.DataFrame],
) -> float:
    """计算 O01-S_i-O01 的往返飞行时间。"""
    destination = to_segment_zone(service_zone)
    distance = matrices["distance_matrix"]
    climb = matrices["climb_height_matrix"]
    descent = matrices["descent_height_matrix"]
    outbound = segment_flight_time(
        float(distance.loc[DEPOT, destination]),
        float(climb.loc[DEPOT, destination]),
        float(descent.loc[DEPOT, destination]),
        uav,
    )
    return_trip = segment_flight_time(
        float(distance.loc[destination, DEPOT]),
        float(climb.loc[destination, DEPOT]),
        float(descent.loc[destination, DEPOT]),
        uav,
    )
    return outbound + return_trip


def batch_operation_time(
    box_count: int, flight_time_s: float, uav: UAVParameters
) -> Tuple[float, float, float, float]:
    """返回准备、装载、交接及总作业时间。"""
    prep_time = uav.prep_time_s
    loading_time = box_count * uav.load_time_per_box_s
    handover_time = (
        uav.handover_base_time_s
        + box_count * uav.handover_time_per_box_s
    )
    total = prep_time + loading_time + flight_time_s + handover_time
    return prep_time, loading_time, handover_time, total


def candidate_energy(
    service_zone: str,
    payload_kg: float,
    uav: UAVParameters,
    matrices: Dict[str, pd.DataFrame],
):
    """复用已有程序公式，计算实际载荷去程和空载返程能耗。"""
    destination = to_segment_zone(service_zone)
    distance = matrices["distance_matrix"]
    climb = matrices["climb_height_matrix"]
    return round_trip_energy(
        payload_kg,
        float(distance.loc[DEPOT, destination]),
        float(distance.loc[destination, DEPOT]),
        float(climb.loc[DEPOT, destination]),
        float(climb.loc[destination, DEPOT]),
        uav.energy_model,
    )


def generate_candidates_for_zone(
    service_zone: str,
    zone_boxes: pd.DataFrame,
    uavs: Dict[str, UAVParameters],
    max_payloads: Dict[Tuple[str, str], float],
    matrices: Dict[str, pd.DataFrame],
) -> List[dict]:
    """使用 DFS/回溯及质量、体积上界剪枝生成一个服务区的候选方案。"""
    records = zone_boxes.to_dict("records")
    max_weight = max(max_payloads[(uav_type, service_zone)] for uav_type in UAV_TYPES)
    max_volume = max(uavs[uav_type].max_volume_m3 for uav_type in UAV_TYPES)
    flight_times = {
        uav_type: round_trip_flight_time(service_zone, uavs[uav_type], matrices)
        for uav_type in UAV_TYPES
    }
    candidates: List[dict] = []
    selected_indices: List[int] = []

    def dfs(start: int, total_weight: float, total_volume: float) -> None:
        if (
            total_weight > max_weight + CAPACITY_TOLERANCE
            or total_volume > max_volume + CAPACITY_TOLERANCE
        ):
            return

        if selected_indices:
            box_ids = tuple(records[index]["box_id"] for index in selected_indices)
            box_count = len(box_ids)
            for uav_type in UAV_TYPES:
                uav = uavs[uav_type]
                safe_payload = max_payloads[(uav_type, service_zone)]
                if (
                    total_weight <= safe_payload + CAPACITY_TOLERANCE
                    and total_volume <= uav.max_volume_m3 + CAPACITY_TOLERANCE
                ):
                    energy = candidate_energy(
                        service_zone, total_weight, uav, matrices
                    )
                    prep, loading, handover, operation = batch_operation_time(
                        box_count, flight_times[uav_type], uav
                    )
                    candidates.append(
                        {
                            "service_zone": service_zone,
                            "uav_type": uav_type,
                            "box_count": box_count,
                            "box_ids": ";".join(box_ids),
                            "box_ids_tuple": box_ids,
                            "total_weight_kg": total_weight,
                            "total_volume_m3": total_volume,
                            "max_safe_payload_kg": safe_payload,
                            "max_volume_m3": uav.max_volume_m3,
                            "weight_remaining_kg": safe_payload - total_weight,
                            "volume_remaining_m3": uav.max_volume_m3 - total_volume,
                            "outbound_energy_kwh": energy.outbound.total_energy_kwh,
                            "return_energy_kwh": energy.return_trip.total_energy_kwh,
                            "energy_kwh": energy.total_energy_kwh,
                            "flight_time_s": flight_times[uav_type],
                            "prep_time_s": prep,
                            "loading_time_s": loading,
                            "handover_time_s": handover,
                            "operation_time_s": operation,
                        }
                    )

        for index in range(start, len(records)):
            new_weight = total_weight + float(records[index]["weight_kg"])
            new_volume = total_volume + float(records[index]["volume_m3"])
            if (
                new_weight > max_weight + CAPACITY_TOLERANCE
                or new_volume > max_volume + CAPACITY_TOLERANCE
            ):
                continue
            selected_indices.append(index)
            dfs(index + 1, new_weight, new_volume)
            selected_indices.pop()

    dfs(0, 0.0, 0.0)
    return candidates


def generate_all_candidates(
    boxes: pd.DataFrame,
    uavs: Dict[str, UAVParameters],
    max_payloads: Dict[Tuple[str, str], float],
    matrices: Dict[str, pd.DataFrame],
) -> pd.DataFrame:
    """按 S001-S015 顺序生成全部候选方案和全局唯一编号。"""
    records: List[dict] = []
    for service_zone in SERVICE_ZONES:
        zone_boxes = boxes.loc[boxes["service_zone"] == service_zone]
        zone_candidates = generate_candidates_for_zone(
            service_zone, zone_boxes, uavs, max_payloads, matrices
        )
        records.extend(zone_candidates)
        print(
            f"候选生成: {service_zone}, 货箱 {len(zone_boxes)} 个, "
            f"候选 {len(zone_candidates)} 个"
        )

    candidates = pd.DataFrame(records)
    candidates.insert(
        0,
        "candidate_id",
        [f"R{index:06d}" for index in range(1, len(candidates) + 1)],
    )
    return candidates


def build_candidate_summary(
    boxes: pd.DataFrame, candidates: pd.DataFrame
) -> pd.DataFrame:
    """统计服务区×机型候选数、服务区总数和全局总数。"""
    rows = []
    total_candidates = len(candidates)
    for service_zone in SERVICE_ZONES:
        box_count = int((boxes["service_zone"] == service_zone).sum())
        zone_candidates = candidates.loc[
            candidates["service_zone"] == service_zone
        ]
        zone_total = len(zone_candidates)
        for uav_type in UAV_TYPES:
            rows.append(
                {
                    "service_zone": service_zone,
                    "uav_type": uav_type,
                    "box_count": box_count,
                    "candidate_count": int(
                        (zone_candidates["uav_type"] == uav_type).sum()
                    ),
                    "service_zone_candidate_count": zone_total,
                    "all_service_zones_candidate_count": total_candidates,
                }
            )
    rows.append(
        {
            "service_zone": "TOTAL",
            "uav_type": "ALL",
            "box_count": len(boxes),
            "candidate_count": total_candidates,
            "service_zone_candidate_count": total_candidates,
            "all_service_zones_candidate_count": total_candidates,
        }
    )
    return pd.DataFrame(rows)


def build_set_partitioning_model(
    boxes: pd.DataFrame, candidates: pd.DataFrame
):
    """建立仅含货箱唯一覆盖约束的 0-1 集合划分模型。"""
    model = gp.Model("problem1_batch_set_partitioning")
    model.Params.MIPGap = 0.0
    model.Params.FeasibilityTol = 1e-9
    model.Params.IntFeasTol = 1e-9
    model.Params.OptimalityTol = 1e-9

    indices = range(len(candidates))
    variables = model.addVars(indices, vtype=GRB.BINARY, name="z")
    box_to_candidates = defaultdict(list)
    for index, box_ids in enumerate(candidates["box_ids_tuple"]):
        for box_id in box_ids:
            box_to_candidates[box_id].append(index)

    for box_id in boxes["box_id"]:
        model.addConstr(
            gp.quicksum(variables[index] for index in box_to_candidates[box_id])
            == 1,
            name=f"cover_{box_id}",
        )

    sortie_expression = gp.quicksum(variables[index] for index in indices)
    energy_values = candidates["energy_kwh"].to_numpy(dtype=float)
    time_values = candidates["operation_time_s"].to_numpy(dtype=float)
    energy_expression = gp.quicksum(
        energy_values[index] * variables[index] for index in indices
    )
    time_expression = gp.quicksum(
        time_values[index] * variables[index] for index in indices
    )
    return (
        model,
        variables,
        sortie_expression,
        energy_expression,
        time_expression,
    )


def optimize_stage(model: gp.Model, objective, name: str) -> float:
    """设置一个阶段的目标并求到全局最优，返回该阶段墙钟时间。"""
    model.setObjective(objective, GRB.MINIMIZE)
    start = perf_counter()
    model.optimize()
    runtime = perf_counter() - start
    if model.Status != GRB.OPTIMAL:
        raise RuntimeError(f"{name} 未得到最优解，Gurobi status={model.Status}")
    return runtime


def solve_lexicographic(
    model: gp.Model,
    variables,
    sortie_expression,
    energy_expression,
    time_expression,
    candidate_count: int,
) -> OptimizationResult:
    """依次最小化架次数、能耗和累计作业时间。"""
    stage1_runtime = optimize_stage(model, sortie_expression, "Stage 1")
    minimum_sorties = int(round(model.ObjVal))
    print(f"Stage 1 optimal sorties = {minimum_sorties}")

    model.addConstr(sortie_expression == minimum_sorties, name="fix_stage1_sorties")
    stage2_runtime = optimize_stage(model, energy_expression, "Stage 2")
    minimum_energy = float(model.ObjVal)
    print(f"Stage 2 optimal energy = {minimum_energy:.12f} kWh")

    model.addConstr(
        energy_expression <= minimum_energy + OBJECTIVE_TOLERANCE_KWH,
        name="preserve_stage2_energy",
    )
    stage3_runtime = optimize_stage(model, time_expression, "Stage 3")
    minimum_time = float(model.ObjVal)
    print(f"Stage 3 optimal operation time = {minimum_time:.9f} s")

    selected = tuple(
        index for index in range(candidate_count) if variables[index].X > 0.5
    )
    return OptimizationResult(
        selected_indices=selected,
        minimum_sorties=minimum_sorties,
        minimum_energy_kwh=minimum_energy,
        minimum_operation_time_s=minimum_time,
        stage1_runtime_s=stage1_runtime,
        stage2_runtime_s=stage2_runtime,
        stage3_runtime_s=stage3_runtime,
        solver_status="OPTIMAL",
    )


def extract_solution(
    boxes: pd.DataFrame,
    candidates: pd.DataFrame,
    result: OptimizationResult,
) -> Tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """构造选中批次、逐箱分配及服务区汇总表。"""
    selected = candidates.iloc[list(result.selected_indices)].copy()
    selected = selected.sort_values(
        ["service_zone", "candidate_id"], kind="stable"
    ).reset_index(drop=True)
    selected.insert(
        0, "sortie_id", [f"P{index:03d}" for index in range(1, len(selected) + 1)]
    )

    assignment_records = []
    box_lookup = boxes.set_index("box_id")
    for row in selected.itertuples(index=False):
        for box_id in row.box_ids_tuple:
            box = box_lookup.loc[box_id]
            assignment_records.append(
                {
                    "box_id": box_id,
                    "service_zone": box["service_zone"],
                    "material_type": box["material_type"],
                    "weight_kg": box["weight_kg"],
                    "volume_m3": box["volume_m3"],
                    "is_first_batch": box["is_first_batch"],
                    "first_batch_deadline_s": box["first_batch_deadline_s"],
                    "expected_arrival_time_s": box["expected_arrival_time_s"],
                    "emergency_priority": box["emergency_priority"],
                    "sortie_id": row.sortie_id,
                    "candidate_id": row.candidate_id,
                    "uav_type": row.uav_type,
                }
            )
    assignment = pd.DataFrame(assignment_records)
    assignment = assignment.set_index("box_id").loc[boxes["box_id"]].reset_index()

    summary_rows = []
    for service_zone in SERVICE_ZONES:
        zone_boxes = boxes.loc[boxes["service_zone"] == service_zone]
        zone_selected = selected.loc[selected["service_zone"] == service_zone]
        summary_rows.append(
            {
                "service_zone": service_zone,
                "box_count": len(zone_boxes),
                "sortie_count": len(zone_selected),
                "A_sorties": int((zone_selected["uav_type"] == "A").sum()),
                "B_sorties": int((zone_selected["uav_type"] == "B").sum()),
                "C_sorties": int((zone_selected["uav_type"] == "C").sum()),
                "total_weight_kg": float(zone_boxes["weight_kg"].sum()),
                "total_volume_m3": float(zone_boxes["volume_m3"].sum()),
                "total_energy_kwh": float(zone_selected["energy_kwh"].sum()),
                "total_operation_time_s": float(
                    zone_selected["operation_time_s"].sum()
                ),
            }
        )
    service_summary = pd.DataFrame(summary_rows)
    return selected, assignment, service_summary


def validate_solution(
    boxes: pd.DataFrame,
    selected: pd.DataFrame,
    assignment: pd.DataFrame,
    uavs: Dict[str, UAVParameters],
    matrices: Dict[str, pd.DataFrame],
    result: OptimizationResult,
) -> None:
    """自动检查唯一覆盖、同区、容量、能量安全和字典序目标。"""
    counts = Counter(
        box_id
        for box_ids in selected["box_ids_tuple"]
        for box_id in box_ids
    )
    unique_coverage = (
        len(counts) == len(boxes)
        and set(counts) == set(boxes["box_id"])
        and all(count == 1 for count in counts.values())
        and len(assignment) == len(boxes)
    )

    box_zones = boxes.set_index("box_id")["service_zone"].to_dict()
    zone_consistency = all(
        all(box_zones[box_id] == row.service_zone for box_id in row.box_ids_tuple)
        for row in selected.itertuples(index=False)
    )
    payload_feasible = bool(
        (
            selected["total_weight_kg"]
            <= selected["max_safe_payload_kg"] + CAPACITY_TOLERANCE
        ).all()
    )
    volume_feasible = bool(
        (
            selected["total_volume_m3"]
            <= selected["max_volume_m3"] + CAPACITY_TOLERANCE
        ).all()
    )

    energy_safe = True
    energy_consistent = True
    for row in selected.itertuples(index=False):
        uav = uavs[row.uav_type]
        recalculated = candidate_energy(
            row.service_zone, row.total_weight_kg, uav, matrices
        ).total_energy_kwh
        limit = (1.0 - uav.energy_model.reserve_ratio) * (
            uav.energy_model.usable_energy_kwh
        )
        energy_safe &= recalculated <= limit + ENERGY_TOLERANCE_KWH
        energy_consistent &= np.isclose(
            recalculated, row.energy_kwh, atol=1e-12, rtol=1e-12
        )

    final_sorties = len(selected)
    final_energy = float(selected["energy_kwh"].sum())
    final_time = float(selected["operation_time_s"].sum())
    objectives_valid = (
        final_sorties == result.minimum_sorties
        and final_energy
        <= result.minimum_energy_kwh + OBJECTIVE_TOLERANCE_KWH + 1e-9
        and final_energy >= result.minimum_energy_kwh - 1e-8
        and np.isclose(
            final_time,
            result.minimum_operation_time_s,
            atol=TIME_TOLERANCE_S,
            rtol=1e-12,
        )
    )

    checks = [
        ("All 80 boxes are assigned exactly once.", unique_coverage),
        ("All selected batches contain boxes from one service zone.", zone_consistency),
        ("All selected batches satisfy payload limits.", payload_feasible),
        ("All selected batches satisfy volume limits.", volume_feasible),
        ("All selected batches satisfy energy safety constraints.", energy_safe),
        ("Selected-batch energies match the reused energy formula.", energy_consistent),
        ("Lexicographic objectives verified.", objectives_valid),
    ]
    print("\n最终方案自动验证:")
    for description, passed in checks:
        print(f"[{'PASS' if passed else 'FAIL'}] {description}")
    failures = [description for description, passed in checks if not passed]
    if failures:
        raise RuntimeError("最终方案验证失败: " + "; ".join(failures))


def format_workbook(writer, matrix_sheets=()) -> None:
    """设置结果工作簿的冻结窗格、列宽和数值格式。"""
    for worksheet in writer.book.worksheets:
        worksheet.freeze_panes = "B2" if worksheet.title in matrix_sheets else "A2"
        for column_cells in worksheet.iter_cols():
            values = ["" if cell.value is None else str(cell.value) for cell in column_cells]
            width = min(max(max(map(len, values)) + 2, 12), 48)
            worksheet.column_dimensions[column_cells[0].column_letter].width = width
            for cell in column_cells[1:]:
                if isinstance(cell.value, float):
                    cell.number_format = "0.000000000"


def save_candidates(
    output_file: Path, candidates: pd.DataFrame, summary: pd.DataFrame
) -> None:
    """保存全部候选方案及候选统计。"""
    output_file.parent.mkdir(parents=True, exist_ok=True)
    exported = candidates.drop(columns="box_ids_tuple")
    with pd.ExcelWriter(assert_writable_output(output_file), engine="openpyxl") as writer:
        exported.to_excel(writer, sheet_name="candidate_solutions", index=False)
        summary.to_excel(writer, sheet_name="candidate_summary", index=False)
        format_workbook(writer)


def save_optimization_results(
    output_file: Path,
    selected: pd.DataFrame,
    assignment: pd.DataFrame,
    service_summary: pd.DataFrame,
    result: OptimizationResult,
    candidate_count: int,
    box_count: int,
    runtime_s: float,
) -> None:
    """保存最终批次、逐箱分配、分区汇总和三阶段求解摘要。"""
    selected_export = selected.drop(columns="box_ids_tuple")
    optimization_summary = pd.DataFrame(
        [
            {
                "minimum_sorties": result.minimum_sorties,
                "minimum_energy_kwh": result.minimum_energy_kwh,
                "minimum_operation_time_s": result.minimum_operation_time_s,
                "candidate_count": candidate_count,
                "selected_candidate_count": len(selected),
                "box_count": box_count,
                "solver_status": result.solver_status,
                "runtime_s": runtime_s,
                "stage1_runtime_s": result.stage1_runtime_s,
                "stage2_runtime_s": result.stage2_runtime_s,
                "stage3_runtime_s": result.stage3_runtime_s,
                "energy_tolerance_kwh": OBJECTIVE_TOLERANCE_KWH,
            }
        ]
    )
    with pd.ExcelWriter(assert_writable_output(output_file), engine="openpyxl") as writer:
        selected_export.to_excel(writer, sheet_name="selected_batches", index=False)
        assignment.to_excel(writer, sheet_name="box_assignment", index=False)
        service_summary.to_excel(
            writer, sheet_name="service_zone_summary", index=False
        )
        optimization_summary.to_excel(
            writer, sheet_name="optimization_summary", index=False
        )
        format_workbook(writer)


def update_recorded_total_runtime(output_file: Path, runtime_s: float) -> None:
    """在结果文件完成首次写出后记录完整的程序墙钟时间。"""
    workbook = load_workbook(output_file)
    worksheet = workbook["optimization_summary"]
    headers = {
        cell.value: cell.column for cell in worksheet[1] if cell.value is not None
    }
    worksheet.cell(row=2, column=headers["runtime_s"], value=runtime_s)
    workbook.save(assert_writable_output(output_file))


def print_input_summary(
    boxes: pd.DataFrame,
    uavs: Dict[str, UAVParameters],
    max_payloads: Dict[Tuple[str, str], float],
) -> None:
    """打印实际读取文件、sheet、货箱数和机型参数。"""
    print(f"赛题原文: {PROJECT_DIR / '山区洪涝灾害下无人机运输与通信协同优化.docx'}")
    print(f"货箱文件: {BOX_FILE} [逐箱货箱清单]")
    print(
        f"航段文件: {SEGMENT_FILE} "
        "[distance_matrix, climb_height_matrix, descent_height_matrix]"
    )
    print(f"无人机参数文件: {UAV_FILE} [数据]")
    print(f"最大安全载荷文件: {MAX_PAYLOAD_FILE} [max_safe_payload]")
    print("问题一不建立具体无人机/电池时序，首批与期望送达字段仅保留输出。")
    print("\n各服务区货箱数量:")
    print(boxes.groupby("service_zone").size().reindex(SERVICE_ZONES).to_string())
    print("\n实际机型容量与作业参数:")
    for uav_type in UAV_TYPES:
        uav = uavs[uav_type]
        payload_range = [max_payloads[(uav_type, zone)] for zone in SERVICE_ZONES]
        print(
            f"{uav_type}: V={uav.max_volume_m3} m^3, "
            f"q_max范围=[{min(payload_range):.6f}, {max(payload_range):.6f}] kg, "
            f"速度(up/cruise/down)=({uav.climb_speed_m_s}/"
            f"{uav.cruise_speed_m_s}/{uav.descent_speed_m_s}) m/s"
        )


def main() -> None:
    program_start = perf_counter()
    boxes = load_box_data(BOX_FILE)
    uavs = load_uav_parameters(UAV_FILE)
    matrices = load_segment_data(SEGMENT_FILE)
    max_payloads = load_max_safe_payload(MAX_PAYLOAD_FILE)
    print_input_summary(boxes, uavs, max_payloads)

    candidate_start = perf_counter()
    candidates = generate_all_candidates(boxes, uavs, max_payloads, matrices)
    candidate_generation_runtime = perf_counter() - candidate_start
    candidate_summary = build_candidate_summary(boxes, candidates)
    save_candidates(CANDIDATE_FILE, candidates, candidate_summary)
    print(f"\n全部候选运输方案总数: {len(candidates)}")
    print(f"候选生成时间: {candidate_generation_runtime:.6f} s")
    print(f"候选方案已保存: {CANDIDATE_FILE}")

    (
        model,
        variables,
        sortie_expression,
        energy_expression,
        time_expression,
    ) = build_set_partitioning_model(boxes, candidates)
    result = solve_lexicographic(
        model,
        variables,
        sortie_expression,
        energy_expression,
        time_expression,
        len(candidates),
    )
    selected, assignment, service_summary = extract_solution(
        boxes, candidates, result
    )
    validate_solution(
        boxes, selected, assignment, uavs, matrices, result
    )

    runtime_before_final_save = perf_counter() - program_start
    save_optimization_results(
        OPTIMIZATION_FILE,
        selected,
        assignment,
        service_summary,
        result,
        len(candidates),
        len(boxes),
        runtime_before_final_save,
    )
    total_runtime = perf_counter() - program_start
    update_recorded_total_runtime(OPTIMIZATION_FILE, total_runtime)

    type_counts = selected["uav_type"].value_counts().reindex(UAV_TYPES, fill_value=0)
    print("\n最终机型架次:")
    print(type_counts.to_string())
    print(f"Stage 1 求解时间: {result.stage1_runtime_s:.6f} s")
    print(f"Stage 2 求解时间: {result.stage2_runtime_s:.6f} s")
    print(f"Stage 3 求解时间: {result.stage3_runtime_s:.6f} s")
    print(f"总运行时间: {total_runtime:.6f} s")
    print(f"最终优化结果已保存: {OPTIMIZATION_FILE}")


if __name__ == "__main__":
    main()
