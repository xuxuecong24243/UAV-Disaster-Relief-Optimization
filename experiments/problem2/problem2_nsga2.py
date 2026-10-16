"""问题二第二种方法：Constraint-handled Random-Key NSGA-II。"""

from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path
import sys

# Keep historical module names while resolving all locations from this repository.
_RELEASE_ROOT = Path(__file__).resolve().parents[2]
if str(_RELEASE_ROOT) not in sys.path:
    sys.path.insert(0, str(_RELEASE_ROOT))
from project_paths import ROOT, bootstrap, raw_file, dem_file, input_file, output_dir, output_file, require_file, assert_writable_output

bootstrap()
from time import perf_counter
from typing import Dict, Iterable, List, Optional, Sequence, Tuple
import math
import random

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
import numpy as np
from openpyxl import load_workbook
import pandas as pd

# 直接复用 ALNS 中经过完整验证的物理模型、数据接口和资源调度解码器。
import problem2_alns as core


RESULT_DIR = output_dir("problem2")
OUTPUT_FILE = RESULT_DIR / "problem2_nsga2_result.xlsx"
LOG_FILE = RESULT_DIR / "problem2_nsga2_run.log"

CONVERGENCE_FIGURE = RESULT_DIR / "问题2_NSGA2收敛图.png"
PARETO_FIGURE = RESULT_DIR / "问题2_NSGA2_Pareto前沿投影.png"
PARALLEL_FIGURE = RESULT_DIR / "问题2_NSGA2四目标平行坐标图.png"
ROUTE_FIGURE = RESULT_DIR / "问题2_NSGA2代表方案运输路线图.png"
UAV_GANTT_FIGURE = RESULT_DIR / "问题2_NSGA2代表方案无人机甘特图.png"
BATTERY_GANTT_FIGURE = RESULT_DIR / "问题2_NSGA2代表方案电池甘特图.png"
DELIVERY_FIGURE = RESULT_DIR / "问题2_NSGA2代表方案送达时间对比图.png"

TOLERANCE = 1e-9


@dataclass(frozen=True)
class NSGA2Config:
    """正式实验参数；mutation_probability=None 表示 1/(3n)。"""

    population_size: int = 24
    max_generations: int = 20
    crossover_probability: float = 0.90
    mutation_probability: Optional[float] = None
    eta_c: float = 15.0
    eta_m: float = 20.0
    beta: float = 1.25
    random_seed: int = 20260924
    log_interval: int = 1


@dataclass
class Individual:
    genes: np.ndarray
    solution: List[core.Sortie] = field(default_factory=list)
    evaluation: Optional[core.Evaluation] = None
    objectives: Tuple[float, float, float, float] = (
        math.inf,
        math.inf,
        math.inf,
        math.inf,
    )
    cv: float = math.inf
    rank: int = 10**9
    crowding_distance: float = 0.0
    signature: Tuple[Tuple, ...] = field(default_factory=tuple)


@dataclass(frozen=True)
class InsertionCandidate:
    target_index: int
    sortie: core.Sortie
    is_new_sortie: bool
    delta_duration_s: float
    delta_energy_kwh: float
    timeliness_proxy: float
    minimum_hard_slack_s: float


@dataclass
class SearchContext:
    data: core.ProblemData
    config: NSGA2Config
    rng: random.Random
    box_ids: Tuple[str, ...]
    box_index: Dict[str, int]
    evaluation_cache: Dict[
        Tuple[Tuple, ...],
        Tuple[core.Evaluation, Tuple[float, float, float, float], float],
    ] = field(default_factory=dict)


CONFIG = NSGA2Config()


def log(message: str, reset: bool = False) -> None:
    print(message)
    mode = "w" if reset else "a"
    with LOG_FILE.open(mode, encoding="utf-8") as stream:
        stream.write(message + "\n")


def load_data() -> core.ProblemData:
    """与 ALNS 完全共用附件、问题一结果和航段预处理读取逻辑。"""
    return core.load_data()


def preprocess_static_data(data: core.ProblemData) -> None:
    """预热三机型的节点间飞行时间缓存；不重复读取Excel或DEM。"""
    nodes = (core.DEPOT, *core.SERVICE_ZONES)
    for uav_type in core.UAV_TYPES:
        for start in nodes:
            for end in nodes:
                if start != end:
                    core.segment_flight_time(data, uav_type, start, end)


def evaluate_sortie(sortie: core.Sortie, data: core.ProblemData) -> core.SortieMetrics:
    return core.get_sortie_metrics(sortie, data)


def check_sortie_capacity(sortie: core.Sortie, data: core.ProblemData) -> bool:
    metrics = evaluate_sortie(sortie, data)
    parameters = data.uav_types[sortie.uav_type]
    return (
        metrics.total_weight_kg <= parameters.payload_limit_kg + TOLERANCE
        and metrics.total_volume_m3 <= parameters.volume_limit_m3 + TOLERANCE
    )


def check_sortie_energy(sortie: core.Sortie, data: core.ProblemData) -> bool:
    metrics = evaluate_sortie(sortie, data)
    parameters = data.uav_types[sortie.uav_type]
    return (
        metrics.energy_kwh <= parameters.energy_limit_kwh + TOLERANCE
        and metrics.return_soc >= parameters.reserve_ratio - TOLERANCE
    )


def repair_hard_deadline(sortie: core.Sortie, data: core.ProblemData) -> bool:
    """判断架次若从0开始，内部硬时限是否理论可行。"""
    delivery_relative = evaluate_sortie(sortie, data).delivery_map
    return all(
        delivery_relative[box_id] <= deadline + TOLERANCE
        for box_id in sortie.box_ids
        for deadline in data.boxes[box_id].hard_deadlines
    )


def gamma_type_priority(gamma_average: float) -> Tuple[str, str, str]:
    if gamma_average < 1.0 / 3.0:
        return ("A", "B", "C")
    if gamma_average < 2.0 / 3.0:
        return ("B", "C", "A")
    return ("C", "B", "A")


def repair_sortie(
    box_ids: Sequence[str],
    route: Sequence[str],
    gamma: np.ndarray,
    context: SearchContext,
) -> Optional[core.Sortie]:
    """按 Gamma_r 机型序尝试容量、能量与内部硬时限修复。"""
    indices = [context.box_index[box_id] for box_id in box_ids]
    gamma_average = float(np.mean(gamma[indices]))
    canonical_box_ids = tuple(sorted(box_ids))
    for uav_type in gamma_type_priority(gamma_average):
        sortie = core.Sortie(canonical_box_ids, tuple(route), uav_type)
        if not check_sortie_capacity(sortie, context.data):
            continue
        if not check_sortie_energy(sortie, context.data):
            continue
        if not repair_hard_deadline(sortie, context.data):
            continue
        return sortie
    return None


def approximate_type_start(
    solution: Sequence[core.Sortie], candidate: core.Sortie, context: SearchContext
) -> float:
    """构造阶段预计开始时刻：同机型已有工作量除以实体机数量。"""
    workload = sum(
        evaluate_sortie(sortie, context.data).duration_s
        for sortie in solution
        if sortie.uav_type == candidate.uav_type
    )
    resource_count = len(context.data.uav_ids_by_type[candidate.uav_type])
    return workload / max(resource_count, 1)


def insertion_timeliness_proxy(
    solution: Sequence[core.Sortie], sortie: core.Sortie, context: SearchContext
) -> Tuple[float, float]:
    """计算构造阶段预计加权延迟与最小硬时限裕量。"""
    metrics = evaluate_sortie(sortie, context.data)
    estimated_start = approximate_type_start(solution, sortie, context)
    delivery_relative = metrics.delivery_map
    soft_tardiness = 0.0
    hard_slacks: List[float] = []
    for box_id in sortie.box_ids:
        box = context.data.boxes[box_id]
        estimated_delivery = estimated_start + delivery_relative[box_id]
        if not box.hard_deadlines:
            soft_tardiness += box.emergency_priority * max(
                0.0, estimated_delivery - box.expected_time_s
            )
        for deadline in box.hard_deadlines:
            hard_slacks.append(deadline - estimated_delivery)
    # 无硬时限候选取0，避免被“无穷裕量”错误地排到硬时限候选之前。
    minimum_slack = min(hard_slacks) if hard_slacks else 0.0
    return soft_tardiness, minimum_slack


def generate_insertion_candidates(
    solution: Sequence[core.Sortie],
    box_id: str,
    gamma: np.ndarray,
    context: SearchContext,
) -> List[InsertionCandidate]:
    """枚举已有架次所有访问位置及新架次，删除局部不可行候选。"""
    candidates: List[InsertionCandidate] = []
    seen = set()
    target_zone = context.data.boxes[box_id].service_zone

    for target_index, old_sortie in enumerate(solution):
        new_box_ids = tuple(old_sortie.box_ids) + (box_id,)
        if target_zone in old_sortie.route:
            route_options = [old_sortie.route]
        else:
            route_options = [
                old_sortie.route[:position]
                + (target_zone,)
                + old_sortie.route[position:]
                for position in range(len(old_sortie.route) + 1)
            ]
        old_metrics = evaluate_sortie(old_sortie, context.data)
        for route in route_options:
            repaired = repair_sortie(new_box_ids, route, gamma, context)
            if repaired is None or repaired in seen:
                continue
            seen.add(repaired)
            new_metrics = evaluate_sortie(repaired, context.data)
            timeliness, minimum_slack = insertion_timeliness_proxy(
                solution, repaired, context
            )
            candidates.append(
                InsertionCandidate(
                    target_index=target_index,
                    sortie=repaired,
                    is_new_sortie=False,
                    delta_duration_s=new_metrics.duration_s - old_metrics.duration_s,
                    delta_energy_kwh=new_metrics.energy_kwh - old_metrics.energy_kwh,
                    timeliness_proxy=timeliness,
                    minimum_hard_slack_s=minimum_slack,
                )
            )

    repaired_new = repair_sortie((box_id,), (target_zone,), gamma, context)
    if repaired_new is None:
        raise RuntimeError(f"货箱 {box_id} 无法建立任何局部可行的新架次")
    metrics = evaluate_sortie(repaired_new, context.data)
    timeliness, minimum_slack = insertion_timeliness_proxy(
        solution, repaired_new, context
    )
    candidates.append(
        InsertionCandidate(
            target_index=-1,
            sortie=repaired_new,
            is_new_sortie=True,
            delta_duration_s=metrics.duration_s,
            delta_energy_kwh=metrics.energy_kwh,
            timeliness_proxy=timeliness,
            minimum_hard_slack_s=minimum_slack,
        )
    )
    return candidates


def rank_insertion_candidates(
    candidates: Sequence[InsertionCandidate], strategy_index: int
) -> List[InsertionCandidate]:
    """按四种互不加权的插入规则排序。"""
    if strategy_index == 0:
        # 配送及时性优先，同时优先保留更大的硬时限裕量。
        key = lambda item: (
            -item.minimum_hard_slack_s,
            item.timeliness_proxy,
            item.delta_duration_s,
            item.delta_energy_kwh,
        )
    elif strategy_index == 1:
        key = lambda item: (
            item.delta_duration_s,
            item.delta_energy_kwh,
            item.is_new_sortie,
        )
    elif strategy_index == 2:
        key = lambda item: (
            item.delta_energy_kwh,
            item.delta_duration_s,
            item.is_new_sortie,
        )
    else:
        # 架次数优先规则严格采用 (C4, Delta E) 字典序。
        key = lambda item: (
            1 if item.is_new_sortie else 0,
            item.delta_energy_kwh,
            item.delta_duration_s,
        )
    return sorted(
        candidates,
        key=lambda item: (key(item), item.target_index, item.sortie.route, item.sortie.uav_type),
    )


def biased_candidate_selection(
    ranked: Sequence[InsertionCandidate], theta: float, beta: float
) -> InsertionCandidate:
    """用区间内随机键 theta 对指数衰减排名分布做逆CDF选择。"""
    weights = np.exp(-beta * np.arange(len(ranked), dtype=float))
    cumulative = np.cumsum(weights / weights.sum())
    selected_index = int(np.searchsorted(cumulative, min(theta, 1.0 - 1e-15)))
    return ranked[min(selected_index, len(ranked) - 1)]


def decode_individual(
    genes: np.ndarray, context: SearchContext
) -> List[core.Sortie]:
    """将 X=(xi,eta,gamma) 解码为组批、路径和架次级机型方案。"""
    n = len(context.box_ids)
    xi = genes[:n]
    eta = genes[n : 2 * n]
    gamma = genes[2 * n :]
    processing_order = sorted(
        range(n), key=lambda index: (xi[index], context.box_ids[index])
    )
    solution: List[core.Sortie] = []

    for box_index in processing_order:
        box_id = context.box_ids[box_index]
        candidates = generate_insertion_candidates(
            solution, box_id, gamma, context
        )
        strategy = min(int(eta[box_index] * 4.0), 3)
        theta = eta[box_index] * 4.0 - strategy
        ranked = rank_insertion_candidates(candidates, strategy)
        selected = biased_candidate_selection(ranked, theta, context.config.beta)
        if selected.is_new_sortie:
            solution.append(selected.sortie)
        else:
            solution[selected.target_index] = selected.sortie
    return solution


def schedule_resources(
    solution: Sequence[core.Sortie], data: core.ProblemData
) -> List[core.ScheduledSortie]:
    return core.decode_resource_schedule(solution, data)


def calculate_charge_time(return_soc: float, full_charge_time_s: float) -> float:
    return core.charge_time(return_soc, full_charge_time_s)


def calculate_delivery_times(
    scheduled: Sequence[core.ScheduledSortie],
) -> Dict[str, float]:
    delivery_times: Dict[str, float] = {}
    for item in scheduled:
        for box_id, relative_time in item.metrics.delivery_relative_s:
            delivery_times[box_id] = item.start_s + relative_time
    return delivery_times


def evaluate_objectives(
    evaluation: core.Evaluation, data: core.ProblemData
) -> Tuple[float, float, float, float]:
    """直接返回四目标；Z1只累计没有硬时限的软要求货箱。"""
    z1 = sum(
        box.emergency_priority
        * max(0.0, evaluation.delivery_times[box.box_id] - box.expected_time_s)
        for box in data.boxes.values()
        if not box.hard_deadlines
    )
    return (
        z1,
        evaluation.z2_makespan_s,
        evaluation.z3_energy_kwh,
        float(evaluation.z4_sorties),
    )


def calculate_constraint_violation(
    evaluation: core.Evaluation, data: core.ProblemData
) -> float:
    """CV=sum max(0,T_j-d_j)/d_j，双硬时限取更严格者。"""
    if not evaluation.coverage_ok or len(evaluation.delivery_times) != len(data.boxes):
        return math.inf
    violation = 0.0
    for box in data.boxes.values():
        if box.hard_deadlines:
            effective_deadline = min(box.hard_deadlines)
            violation += max(
                0.0,
                evaluation.delivery_times[box.box_id] - effective_deadline,
            ) / effective_deadline
    return violation


def evaluate_individual(individual: Individual, context: SearchContext) -> None:
    # 架次组合缓存对亲缘个体很有效；设置上限只控制内存，不改变任何计算结果。
    if len(context.data.sortie_cache) > 120_000:
        context.data.sortie_cache.clear()
    solution = decode_individual(individual.genes, context)
    signature = core.solution_signature(solution)
    if signature in context.evaluation_cache:
        evaluation, objectives, cv = context.evaluation_cache[signature]
    else:
        # 权重在此只为满足公共函数接口；NSGA-II支配关系不使用该标量目标。
        evaluation = core.evaluate_solution(
            solution,
            context.data,
            (0.25, 0.25, 0.25, 0.25),
            require_complete=True,
        )
        objectives = evaluate_objectives(evaluation, context.data)
        cv = calculate_constraint_violation(evaluation, context.data)
        context.evaluation_cache[signature] = (evaluation, objectives, cv)
    individual.solution = solution
    individual.signature = signature
    individual.evaluation = evaluation
    individual.objectives = objectives
    individual.cv = cv


def pareto_dominates(
    objectives_a: Sequence[float], objectives_b: Sequence[float]
) -> bool:
    return all(a <= b + TOLERANCE for a, b in zip(objectives_a, objectives_b)) and any(
        a < b - TOLERANCE for a, b in zip(objectives_a, objectives_b)
    )


def constraint_dominates(a: Individual, b: Individual) -> bool:
    """Deb约束支配：可行优先；均不可行时CV小者优先。"""
    a_feasible = a.cv <= TOLERANCE
    b_feasible = b.cv <= TOLERANCE
    if a_feasible and not b_feasible:
        return True
    if not a_feasible and b_feasible:
        return False
    if not a_feasible and not b_feasible:
        if a.cv < b.cv - TOLERANCE:
            return True
        if a.cv > b.cv + TOLERANCE:
            return False
    return pareto_dominates(a.objectives, b.objectives)


def fast_non_dominated_sort(population: Sequence[Individual]) -> List[List[Individual]]:
    """标准O(N^2)快速非支配排序，支配关系包含CV。"""
    size = len(population)
    dominated_sets: List[List[int]] = [[] for _ in range(size)]
    domination_counts = [0] * size
    fronts_indices: List[List[int]] = [[]]

    for p_index in range(size):
        for q_index in range(size):
            if p_index == q_index:
                continue
            if constraint_dominates(population[p_index], population[q_index]):
                dominated_sets[p_index].append(q_index)
            elif constraint_dominates(population[q_index], population[p_index]):
                domination_counts[p_index] += 1
        if domination_counts[p_index] == 0:
            population[p_index].rank = 0
            fronts_indices[0].append(p_index)

    front_index = 0
    while front_index < len(fronts_indices) and fronts_indices[front_index]:
        next_front: List[int] = []
        for p_index in fronts_indices[front_index]:
            for q_index in dominated_sets[p_index]:
                domination_counts[q_index] -= 1
                if domination_counts[q_index] == 0:
                    population[q_index].rank = front_index + 1
                    next_front.append(q_index)
        if next_front:
            fronts_indices.append(next_front)
        front_index += 1
    return [[population[index] for index in front] for front in fronts_indices if front]


def calculate_crowding_distance(front: Sequence[Individual]) -> None:
    """在四个原始目标上计算拥挤距离，并处理零目标极差。"""
    if not front:
        return
    for individual in front:
        individual.crowding_distance = 0.0
    if len(front) <= 2:
        for individual in front:
            individual.crowding_distance = math.inf
        return

    for objective_index in range(4):
        ordered = sorted(
            front,
            key=lambda individual: individual.objectives[objective_index],
        )
        ordered[0].crowding_distance = math.inf
        ordered[-1].crowding_distance = math.inf
        minimum = ordered[0].objectives[objective_index]
        maximum = ordered[-1].objectives[objective_index]
        if abs(maximum - minimum) <= TOLERANCE:
            continue
        for index in range(1, len(ordered) - 1):
            if math.isinf(ordered[index].crowding_distance):
                continue
            previous_value = ordered[index - 1].objectives[objective_index]
            next_value = ordered[index + 1].objectives[objective_index]
            ordered[index].crowding_distance += (
                next_value - previous_value
            ) / (maximum - minimum)


def assign_rank_and_crowding(population: Sequence[Individual]) -> List[List[Individual]]:
    fronts = fast_non_dominated_sort(population)
    for front in fronts:
        calculate_crowding_distance(front)
    return fronts


def tournament_selection(
    population: Sequence[Individual], rng: random.Random
) -> Individual:
    """二元锦标赛依次比较约束支配、rank和拥挤距离。"""
    first, second = rng.sample(list(population), 2)
    if constraint_dominates(first, second):
        return first
    if constraint_dominates(second, first):
        return second
    if first.rank != second.rank:
        return first if first.rank < second.rank else second
    if abs(first.crowding_distance - second.crowding_distance) > TOLERANCE:
        return first if first.crowding_distance > second.crowding_distance else second
    return first if rng.random() < 0.5 else second


def sbx_crossover(
    parent_a: np.ndarray,
    parent_b: np.ndarray,
    probability: float,
    eta_c: float,
    rng: random.Random,
) -> Tuple[np.ndarray, np.ndarray]:
    """[0,1]有界模拟二进制交叉。"""
    child_a = parent_a.copy()
    child_b = parent_b.copy()
    if rng.random() > probability:
        return child_a, child_b

    for index in range(len(parent_a)):
        if rng.random() > 0.5 or abs(parent_a[index] - parent_b[index]) <= 1e-14:
            continue
        x1 = min(parent_a[index], parent_b[index])
        x2 = max(parent_a[index], parent_b[index])
        rand = rng.random()

        beta = 1.0 + 2.0 * x1 / (x2 - x1)
        alpha = 2.0 - beta ** (-(eta_c + 1.0))
        if rand <= 1.0 / alpha:
            beta_q = (rand * alpha) ** (1.0 / (eta_c + 1.0))
        else:
            beta_q = (1.0 / (2.0 - rand * alpha)) ** (1.0 / (eta_c + 1.0))
        value_a = 0.5 * ((x1 + x2) - beta_q * (x2 - x1))

        beta = 1.0 + 2.0 * (1.0 - x2) / (x2 - x1)
        alpha = 2.0 - beta ** (-(eta_c + 1.0))
        if rand <= 1.0 / alpha:
            beta_q = (rand * alpha) ** (1.0 / (eta_c + 1.0))
        else:
            beta_q = (1.0 / (2.0 - rand * alpha)) ** (1.0 / (eta_c + 1.0))
        value_b = 0.5 * ((x1 + x2) + beta_q * (x2 - x1))
        if rng.random() < 0.5:
            value_a, value_b = value_b, value_a
        child_a[index] = min(1.0, max(0.0, value_a))
        child_b[index] = min(1.0, max(0.0, value_b))
    return child_a, child_b


def polynomial_mutation(
    genes: np.ndarray,
    probability: float,
    eta_m: float,
    rng: random.Random,
) -> np.ndarray:
    """[0,1]有界Polynomial Mutation。"""
    mutated = genes.copy()
    mutation_power = 1.0 / (eta_m + 1.0)
    for index, value in enumerate(mutated):
        if rng.random() > probability:
            continue
        delta1 = value
        delta2 = 1.0 - value
        rand = rng.random()
        if rand < 0.5:
            xy = 1.0 - delta1
            value_term = 2.0 * rand + (1.0 - 2.0 * rand) * xy ** (eta_m + 1.0)
            delta_q = value_term ** mutation_power - 1.0
        else:
            xy = 1.0 - delta2
            value_term = (
                2.0 * (1.0 - rand)
                + 2.0 * (rand - 0.5) * xy ** (eta_m + 1.0)
            )
            delta_q = 1.0 - value_term ** mutation_power
        mutated[index] = min(1.0, max(0.0, value + delta_q))
    return mutated


def genes_from_order(
    order: Sequence[str],
    context: SearchContext,
    eta_base: float,
    gamma_by_box: Optional[Dict[str, float]] = None,
) -> np.ndarray:
    """把启发式货箱顺序转换为随机键种子。"""
    n = len(context.box_ids)
    xi = np.zeros(n, dtype=float)
    for rank, box_id in enumerate(order):
        xi[context.box_index[box_id]] = (rank + 0.5) / n
    eta = np.array(
        [
            min(1.0, max(0.0, eta_base + 0.02 * ((index % 5) - 2)))
            for index in range(n)
        ],
        dtype=float,
    )
    if gamma_by_box is None:
        gamma = np.array(
            [0.15 if index % 3 == 0 else 0.50 if index % 3 == 1 else 0.85 for index in range(n)],
            dtype=float,
        )
    else:
        gamma = np.array(
            [gamma_by_box.get(box_id, 0.50) for box_id in context.box_ids],
            dtype=float,
        )
    return np.concatenate((xi, eta, gamma))


def problem1_seed_genes(context: SearchContext) -> np.ndarray:
    order: List[str] = []
    gamma_by_box: Dict[str, float] = {}
    gamma_value = {"A": 0.15, "B": 0.50, "C": 0.85}
    for sortie in context.data.initial_solution:
        order.extend(sortie.box_ids)
        for box_id in sortie.box_ids:
            gamma_by_box[box_id] = gamma_value[sortie.uav_type]
    return genes_from_order(order, context, eta_base=0.875, gamma_by_box=gamma_by_box)


def build_problem1_feasible_reference(context: SearchContext) -> List[core.Sortie]:
    """仅拆出问题一种子中实际超硬时限货箱，建立两个A型紧急单点架次。"""
    seed_evaluation = core.evaluate_solution(
        context.data.initial_solution,
        context.data,
        (0.25, 0.25, 0.25, 0.25),
        require_complete=True,
    )
    violating = {
        box_id
        for box_id, delivery_time in seed_evaluation.delivery_times.items()
        if any(
            delivery_time > deadline + TOLERANCE
            for deadline in context.data.boxes[box_id].hard_deadlines
        )
    }
    reference: List[core.Sortie] = []
    for sortie in context.data.initial_solution:
        remaining = tuple(box_id for box_id in sortie.box_ids if box_id not in violating)
        if remaining:
            rebuilt = core.optimize_route(remaining, sortie.uav_type, context.data)
            if rebuilt is None:
                raise RuntimeError("问题一参考架次移除超时货箱后不可行")
            reference.append(rebuilt)
    for service_zone in sorted(
        {context.data.boxes[box_id].service_zone for box_id in violating}
    ):
        box_ids = tuple(
            sorted(
                box_id
                for box_id in violating
                if context.data.boxes[box_id].service_zone == service_zone
            )
        )
        emergency_sortie = core.Sortie(box_ids, (service_zone,), "A")
        if not core.check_sortie_feasibility(emergency_sortie, context.data):
            raise RuntimeError("问题一参考解的A型紧急修复架次不可行")
        reference.append(emergency_sortie)
    evaluation = core.evaluate_solution(
        reference,
        context.data,
        (0.25, 0.25, 0.25, 0.25),
        require_complete=True,
    )
    if not evaluation.feasible:
        raise RuntimeError("问题一种子做最小硬时限修复后仍不可行")
    return reference


def encode_reference_solution(
    reference: Sequence[core.Sortie], context: SearchContext
) -> np.ndarray:
    """反求一组随机键，使解码器确定性复现给定参考方案。"""
    n = len(context.box_ids)
    order = [box_id for sortie in reference for box_id in sortie.box_ids]
    xi = np.zeros(n, dtype=float)
    gamma = np.zeros(n, dtype=float)
    type_gene = {"A": 0.15, "B": 0.50, "C": 0.85}
    for rank, box_id in enumerate(order):
        xi[context.box_index[box_id]] = (rank + 0.5) / n
    for sortie in reference:
        for box_id in sortie.box_ids:
            gamma[context.box_index[box_id]] = type_gene[sortie.uav_type]

    eta = np.zeros(n, dtype=float)
    simulated: List[core.Sortie] = []
    reference_to_current: Dict[int, int] = {}
    for reference_index, reference_sortie in enumerate(reference):
        accumulated: List[str] = []
        for box_id in reference_sortie.box_ids:
            accumulated.append(box_id)
            candidates = generate_insertion_candidates(
                simulated, box_id, gamma, context
            )
            desired: Optional[InsertionCandidate] = None
            if reference_index not in reference_to_current:
                for candidate in candidates:
                    if (
                        candidate.is_new_sortie
                        and set(candidate.sortie.box_ids) == set(accumulated)
                        and candidate.sortie.uav_type == reference_sortie.uav_type
                    ):
                        desired = candidate
                        break
            else:
                current_index = reference_to_current[reference_index]
                accumulated_zones = {
                    context.data.boxes[item].service_zone for item in accumulated
                }
                partial_route = tuple(
                    zone
                    for zone in reference_sortie.route
                    if zone in accumulated_zones
                )
                for candidate in candidates:
                    if (
                        not candidate.is_new_sortie
                        and candidate.target_index == current_index
                        and set(candidate.sortie.box_ids) == set(accumulated)
                        and candidate.sortie.route == partial_route
                        and candidate.sortie.uav_type == reference_sortie.uav_type
                    ):
                        desired = candidate
                        break
            if desired is None:
                raise RuntimeError("随机键无法反向编码问题一可行参考方案")

            encodable_options = []
            for strategy_index in range(4):
                ranked_option = rank_insertion_candidates(
                    candidates, strategy_index=strategy_index
                )
                desired_rank = next(
                    index
                    for index, candidate in enumerate(ranked_option)
                    if candidate == desired
                )
                weights = np.exp(
                    -context.config.beta
                    * np.arange(len(ranked_option), dtype=float)
                )
                cumulative = np.cumsum(weights / weights.sum())
                lower = (
                    0.0
                    if desired_rank == 0
                    else float(cumulative[desired_rank - 1])
                )
                upper = float(cumulative[desired_rank])
                if upper > lower:
                    encodable_options.append(
                        (
                            desired_rank,
                            strategy_index,
                            (lower + upper) / 2.0,
                            ranked_option,
                        )
                    )
            if not encodable_options:
                raise RuntimeError("参考方案候选排名概率低于浮点分辨率")
            _, strategy_index, theta, ranked = min(
                encodable_options, key=lambda item: (item[0], item[1])
            )
            eta[context.box_index[box_id]] = (strategy_index + theta) / 4.0

            selected = biased_candidate_selection(
                ranked, theta, context.config.beta
            )
            if selected != desired:
                raise RuntimeError("参考方案随机键逆CDF编码失败")
            if selected.is_new_sortie:
                simulated.append(selected.sortie)
                reference_to_current[reference_index] = len(simulated) - 1
            else:
                simulated[selected.target_index] = selected.sortie

    if core.solution_signature(simulated) != core.solution_signature(reference):
        raise RuntimeError("参考方案随机键复现校验失败")
    return np.concatenate((xi, eta, gamma))


def build_feasible_reference_variants(
    reference: Sequence[core.Sortie], context: SearchContext, limit: int = 12
) -> List[List[core.Sortie]]:
    """从问题一修复参考解拆出不同软货箱，形成不重复的可行种子。"""
    variants: List[List[core.Sortie]] = [list(reference)]
    signatures = {core.solution_signature(reference)}
    # 按每个源架次的第1、第2...个软货箱轮换，避免种子集中在同一服务区。
    soft_by_source = {
        source_index: [
            box_id
            for box_id in source.box_ids
            if not context.data.boxes[box_id].hard_deadlines
        ]
        for source_index, source in enumerate(reference)
        if len(source.box_ids) > 1
    }
    maximum_soft_count = max(map(len, soft_by_source.values()), default=0)
    source_box_pairs = [
        (source_index, soft_by_source[source_index][soft_position])
        for soft_position in range(maximum_soft_count)
        for source_index in sorted(soft_by_source)
        if soft_position < len(soft_by_source[source_index])
    ]
    for source_index, box_id in source_box_pairs:
        source = reference[source_index]
        remaining = tuple(item for item in source.box_ids if item != box_id)
        rebuilt_source = core.optimize_route(
            remaining, source.uav_type, context.data
        )
        if rebuilt_source is None:
            continue
        single_options = [
            core.optimize_route((box_id,), uav_type, context.data)
            for uav_type in core.UAV_TYPES
        ]
        feasible_single = [item for item in single_options if item is not None]
        if not feasible_single:
            continue
        single = min(
            feasible_single,
            key=lambda item: (
                evaluate_sortie(item, context.data).energy_kwh,
                evaluate_sortie(item, context.data).duration_s,
                item.uav_type,
            ),
        )
        candidate = list(reference)
        candidate[source_index] = rebuilt_source
        candidate.append(single)
        signature = core.solution_signature(candidate)
        if signature in signatures:
            continue
        evaluation = core.evaluate_solution(
            candidate,
            context.data,
            (0.25, 0.25, 0.25, 0.25),
            require_complete=True,
        )
        if evaluation.feasible:
            variants.append(candidate)
            signatures.add(signature)
        if len(variants) >= limit:
            return variants
    return variants


def spatial_order(context: SearchContext) -> List[str]:
    """按服务区最近邻形成空间接近性启发式序列。"""
    unvisited = set(core.SERVICE_ZONES)
    zone_order: List[str] = []
    current = core.DEPOT
    while unvisited:
        next_zone = min(
            unvisited,
            key=lambda zone: (
                float(
                    context.data.distance.at[
                        core.to_segment_node(current), core.to_segment_node(zone)
                    ]
                ),
                zone,
            ),
        )
        zone_order.append(next_zone)
        unvisited.remove(next_zone)
        current = next_zone
    zone_rank = {zone: rank for rank, zone in enumerate(zone_order)}
    return sorted(
        context.box_ids,
        key=lambda box_id: (
            zone_rank[context.data.boxes[box_id].service_zone],
            min(context.data.boxes[box_id].hard_deadlines, default=math.inf),
            box_id,
        ),
    )


def heuristic_seed_genes(context: SearchContext) -> List[np.ndarray]:
    boxes = context.data.boxes
    orders = [
        sorted(
            context.box_ids,
            key=lambda box_id: (
                min(boxes[box_id].hard_deadlines, default=math.inf),
                boxes[box_id].expected_time_s,
                -boxes[box_id].emergency_priority,
                box_id,
            ),
        ),
        sorted(
            context.box_ids,
            key=lambda box_id: (
                0 if boxes[box_id].material_type == "医疗物资" else 1,
                min(boxes[box_id].hard_deadlines, default=math.inf),
                box_id,
            ),
        ),
        sorted(
            context.box_ids,
            key=lambda box_id: (
                boxes[box_id].expected_time_s,
                -boxes[box_id].emergency_priority,
                box_id,
            ),
        ),
        sorted(
            context.box_ids,
            key=lambda box_id: (
                -boxes[box_id].emergency_priority,
                boxes[box_id].expected_time_s,
                box_id,
            ),
        ),
        sorted(
            context.box_ids,
            key=lambda box_id: (
                -boxes[box_id].weight_kg,
                -boxes[box_id].volume_m3,
                box_id,
            ),
        ),
        spatial_order(context),
    ]
    eta_bases = [0.125, 0.125, 0.375, 0.125, 0.625, 0.875]
    return [
        genes_from_order(order, context, eta_base)
        for order, eta_base in zip(orders, eta_bases)
    ]


def build_makespan_reference(
    context: SearchContext, sortie_overhead_s: float = 0.0
) -> List[core.Sortie]:
    """构造资源完工时间优先的独立启发式种子。

    这里不调用 ALNS 算子。货箱按硬截止期、期望时刻和应急系数排序，
    每次完整枚举已有架次/新架次、访问位置和三种机型。候选先满足
    容量、逐段能量、SOC及架次内部硬时限，再用与最终评价完全相同的
    无人机--共享电池解码器计算部分方案 Cmax，选择 Cmax 最小者。
    """
    boxes = context.data.boxes
    processing_order = sorted(
        context.box_ids,
        key=lambda box_id: (
            min(boxes[box_id].hard_deadlines, default=math.inf),
            boxes[box_id].expected_time_s,
            -boxes[box_id].emergency_priority,
            box_id,
        ),
    )
    solution: List[core.Sortie] = []

    for box_id in processing_order:
        target_zone = boxes[box_id].service_zone
        candidate_solutions: List[List[core.Sortie]] = []
        seen_signatures = set()

        for target_index, old_sortie in enumerate(solution):
            new_box_ids = tuple(sorted((*old_sortie.box_ids, box_id)))
            if target_zone in old_sortie.route:
                route_options = [old_sortie.route]
            else:
                route_options = [
                    old_sortie.route[:position]
                    + (target_zone,)
                    + old_sortie.route[position:]
                    for position in range(len(old_sortie.route) + 1)
                ]
            for route in route_options:
                for uav_type in core.UAV_TYPES:
                    candidate_sortie = core.Sortie(
                        new_box_ids, tuple(route), uav_type
                    )
                    if not core.check_sortie_feasibility(
                        candidate_sortie, context.data
                    ) or not repair_hard_deadline(candidate_sortie, context.data):
                        continue
                    candidate = list(solution)
                    candidate[target_index] = candidate_sortie
                    signature = core.solution_signature(candidate)
                    if signature not in seen_signatures:
                        seen_signatures.add(signature)
                        candidate_solutions.append(candidate)

        for uav_type in core.UAV_TYPES:
            candidate_sortie = core.Sortie((box_id,), (target_zone,), uav_type)
            if not core.check_sortie_feasibility(
                candidate_sortie, context.data
            ) or not repair_hard_deadline(candidate_sortie, context.data):
                continue
            candidate = [*solution, candidate_sortie]
            signature = core.solution_signature(candidate)
            if signature not in seen_signatures:
                seen_signatures.add(signature)
                candidate_solutions.append(candidate)

        evaluated_candidates = []
        for candidate in candidate_solutions:
            evaluation = core.evaluate_solution(
                candidate,
                context.data,
                (0.25, 0.25, 0.25, 0.25),
                require_complete=False,
            )
            if not evaluation.feasible:
                continue
            # 固定架次开销只用于生成不同结构的种子，不进入NSGA-II目标。
            # overhead=0 是纯Cmax贪心；较大取值鼓励合并，供种群自行筛选。
            evaluated_candidates.append(
                (
                    evaluation.z2_makespan_s
                    + sortie_overhead_s * evaluation.z4_sorties,
                    evaluation.z2_makespan_s,
                    evaluation.z4_sorties,
                    evaluation.z3_energy_kwh,
                    core.solution_signature(candidate),
                    candidate,
                )
            )
        if not evaluated_candidates:
            raise RuntimeError(f"完工时间启发式无法可行插入货箱 {box_id}")
        solution = min(evaluated_candidates, key=lambda item: item[:-1])[-1]

    final_evaluation = core.evaluate_solution(
        solution,
        context.data,
        (0.25, 0.25, 0.25, 0.25),
        require_complete=True,
    )
    if not final_evaluation.feasible:
        raise RuntimeError("完工时间启发式种子未通过全局可行性检查")
    return solution


def merge_makespan_reference(
    reference: Sequence[core.Sortie],
    target_sorties: int,
    context: SearchContext,
) -> List[core.Sortie]:
    """以确定性两架次合并压缩种子，并用真实 Cmax 选择每一步。"""
    solution = list(reference)
    while len(solution) > target_sorties:
        merge_options = []
        for first_index in range(len(solution) - 1):
            for second_index in range(first_index + 1, len(solution)):
                merged_box_ids = tuple(
                    sorted(
                        (
                            *solution[first_index].box_ids,
                            *solution[second_index].box_ids,
                        )
                    )
                )
                for uav_type in core.UAV_TYPES:
                    merged = core.optimize_route(
                        merged_box_ids, uav_type, context.data
                    )
                    if merged is None or not repair_hard_deadline(
                        merged, context.data
                    ):
                        continue
                    candidate = [
                        sortie
                        for index, sortie in enumerate(solution)
                        if index not in (first_index, second_index)
                    ]
                    candidate.append(merged)
                    evaluation = core.evaluate_solution(
                        candidate,
                        context.data,
                        (0.25, 0.25, 0.25, 0.25),
                        require_complete=True,
                    )
                    if not evaluation.feasible:
                        continue
                    merge_options.append(
                        (
                            evaluation.z2_makespan_s,
                            evaluation.z3_energy_kwh,
                            core.solution_signature(candidate),
                            candidate,
                        )
                    )
        if not merge_options:
            break
        solution = min(merge_options, key=lambda item: item[:-1])[-1]
    return solution


def improve_makespan_reference(
    reference: Sequence[core.Sortie],
    context: SearchContext,
    max_iterations: int = 8,
) -> List[core.Sortie]:
    """用确定性单箱重定位/机型重选改进资源 Cmax 启发式种子。"""
    solution = list(reference)
    current = core.evaluate_solution(
        solution,
        context.data,
        (0.25, 0.25, 0.25, 0.25),
        require_complete=True,
    )
    for _ in range(max_iterations):
        options = []
        seen_signatures = set()

        # 同一组批重新选择机型和优化路径。
        for sortie_index, sortie in enumerate(solution):
            for uav_type in core.UAV_TYPES:
                rebuilt = core.optimize_route(
                    sortie.box_ids, uav_type, context.data
                )
                if rebuilt is None or not repair_hard_deadline(
                    rebuilt, context.data
                ):
                    continue
                candidate = list(solution)
                candidate[sortie_index] = rebuilt
                signature = core.solution_signature(candidate)
                if signature in seen_signatures:
                    continue
                seen_signatures.add(signature)
                evaluation = core.evaluate_solution(
                    candidate,
                    context.data,
                    (0.25, 0.25, 0.25, 0.25),
                    require_complete=True,
                )
                if evaluation.feasible:
                    options.append((evaluation, signature, candidate))

        # 把一个货箱从源架次移至另一个架次；源和目标均重新选机型/路径。
        for source_index, source in enumerate(solution):
            for box_id in source.box_ids:
                remaining_ids = tuple(
                    item for item in source.box_ids if item != box_id
                )
                if remaining_ids:
                    source_options = [
                        core.optimize_route(remaining_ids, uav_type, context.data)
                        for uav_type in core.UAV_TYPES
                    ]
                    source_options = [
                        item
                        for item in source_options
                        if item is not None
                        and repair_hard_deadline(item, context.data)
                    ]
                else:
                    source_options = [None]
                for target_index, target in enumerate(solution):
                    if target_index == source_index:
                        continue
                    target_ids = tuple(sorted((*target.box_ids, box_id)))
                    target_options = [
                        core.optimize_route(target_ids, uav_type, context.data)
                        for uav_type in core.UAV_TYPES
                    ]
                    target_options = [
                        item
                        for item in target_options
                        if item is not None
                        and repair_hard_deadline(item, context.data)
                    ]
                    for rebuilt_source in source_options:
                        for rebuilt_target in target_options:
                            candidate = [
                                sortie
                                for index, sortie in enumerate(solution)
                                if index not in (source_index, target_index)
                            ]
                            if rebuilt_source is not None:
                                candidate.append(rebuilt_source)
                            candidate.append(rebuilt_target)
                            signature = core.solution_signature(candidate)
                            if signature in seen_signatures:
                                continue
                            seen_signatures.add(signature)
                            evaluation = core.evaluate_solution(
                                candidate,
                                context.data,
                                (0.25, 0.25, 0.25, 0.25),
                                require_complete=True,
                            )
                            if evaluation.feasible:
                                options.append((evaluation, signature, candidate))

        improving = [
            item
            for item in options
            if item[0].z2_makespan_s < current.z2_makespan_s - TOLERANCE
        ]
        if not improving:
            break
        current, _, solution = min(
            improving,
            key=lambda item: (
                item[0].z2_makespan_s,
                item[0].z4_sorties,
                item[0].z3_energy_kwh,
                item[1],
            ),
        )
    return solution


def initialize_population(context: SearchContext) -> List[Individual]:
    """问题一、资源完工时间及六类启发式种子+随机个体。"""
    dimension = 3 * len(context.box_ids)
    feasible_reference = build_problem1_feasible_reference(context)
    makespan_references = [
        build_makespan_reference(context, sortie_overhead_s=overhead)
        for overhead in (0.0, 60.0, 120.0, 180.0, 240.0)
    ]
    improved_makespan_reference = improve_makespan_reference(
        makespan_references[3], context, max_iterations=16
    )
    merged_references = [
        merge_makespan_reference(improved_makespan_reference, target, context)
        for target in (25, 24, 23, 22, 21)
    ]
    feasible_variants = build_feasible_reference_variants(
        feasible_reference, context, limit=4
    )
    gene_candidates = [
        *(encode_reference_solution(reference, context) for reference in makespan_references),
        *(encode_reference_solution(reference, context) for reference in merged_references),
        encode_reference_solution(improved_makespan_reference, context),
        *(encode_reference_solution(variant, context) for variant in feasible_variants),
        problem1_seed_genes(context),
        *heuristic_seed_genes(context),
    ]
    population: List[Individual] = []
    decoded_signatures = set()
    attempts = 0

    while len(population) < context.config.population_size:
        if gene_candidates:
            genes = gene_candidates.pop(0)
        else:
            genes = np.array(
                [context.rng.random() for _ in range(dimension)], dtype=float
            )
        individual = Individual(genes=genes)
        evaluate_individual(individual, context)
        attempts += 1
        if individual.signature in decoded_signatures and attempts < 20 * context.config.population_size:
            continue
        decoded_signatures.add(individual.signature)
        population.append(individual)

    assign_rank_and_crowding(population)
    log(
        "初始种群: size={}, decoded unique={}, feasible={}".format(
            len(population),
            len(decoded_signatures),
            sum(individual.cv <= TOLERANCE for individual in population),
        )
    )
    return population


def create_offspring(
    population: Sequence[Individual], context: SearchContext
) -> List[Individual]:
    offspring: List[Individual] = []
    dimension = 3 * len(context.box_ids)
    mutation_probability = (
        context.config.mutation_probability
        if context.config.mutation_probability is not None
        else 1.0 / dimension
    )
    while len(offspring) < context.config.population_size:
        parent_a = tournament_selection(population, context.rng)
        parent_b = tournament_selection(population, context.rng)
        genes_a, genes_b = sbx_crossover(
            parent_a.genes,
            parent_b.genes,
            context.config.crossover_probability,
            context.config.eta_c,
            context.rng,
        )
        for genes in (genes_a, genes_b):
            genes = polynomial_mutation(
                genes,
                mutation_probability,
                context.config.eta_m,
                context.rng,
            )
            individual = Individual(genes=genes)
            evaluate_individual(individual, context)
            offspring.append(individual)
            if len(offspring) == context.config.population_size:
                break
    return offspring


def environmental_selection(
    combined: Sequence[Individual], population_size: int
) -> List[Individual]:
    fronts = assign_rank_and_crowding(combined)
    next_population: List[Individual] = []
    for front in fronts:
        if len(next_population) + len(front) <= population_size:
            next_population.extend(front)
        else:
            remaining = population_size - len(next_population)
            next_population.extend(
                sorted(
                    front,
                    key=lambda individual: (
                        -individual.crowding_distance,
                        individual.objectives,
                        individual.signature,
                    ),
                )[:remaining]
            )
            break
    assign_rank_and_crowding(next_population)
    return next_population


def generation_log_row(
    generation: int, population: Sequence[Individual]
) -> dict:
    fronts = assign_rank_and_crowding(population)
    feasible = [individual for individual in population if individual.cv <= TOLERANCE]
    source = feasible if feasible else list(population)
    return {
        "generation": generation,
        "feasible_count": len(feasible),
        "first_front_size": len(fronts[0]) if fronts else 0,
        "min_Z1": min(individual.objectives[0] for individual in source),
        "min_Z2": min(individual.objectives[1] for individual in source),
        "min_Z3": min(individual.objectives[2] for individual in source),
        "min_Z4": min(individual.objectives[3] for individual in source),
        "min_CV": min(individual.cv for individual in population),
        "decoded_cache_size": len({individual.signature for individual in population}),
    }


def nsga2(context: SearchContext) -> Tuple[List[Individual], pd.DataFrame]:
    """标准精英NSGA-II主循环。"""
    population = initialize_population(context)
    log_rows = [generation_log_row(0, population)]
    for generation in range(1, context.config.max_generations + 1):
        offspring = create_offspring(population, context)
        population = environmental_selection(
            list(population) + offspring,
            context.config.population_size,
        )
        row = generation_log_row(generation, population)
        log_rows.append(row)
        if generation % context.config.log_interval == 0:
            log(
                "generation {:3d}: feasible={:2d}, F1={:2d}, "
                "min=({:.3f}, {:.3f}, {:.6f}, {:.0f})".format(
                    generation,
                    row["feasible_count"],
                    row["first_front_size"],
                    row["min_Z1"],
                    row["min_Z2"],
                    row["min_Z3"],
                    row["min_Z4"],
                )
            )
    return population, pd.DataFrame(log_rows)


def extract_pareto_front(population: Sequence[Individual]) -> List[Individual]:
    """仅从CV=0个体中提取第一层，并去除重复目标向量或重复方案。"""
    feasible = [individual for individual in population if individual.cv <= TOLERANCE]
    unique: List[Individual] = []
    seen_objectives = set()
    seen_signatures = set()
    for individual in sorted(
        feasible, key=lambda item: (item.objectives, item.signature)
    ):
        objective_key = tuple(round(value, 9) for value in individual.objectives)
        if objective_key in seen_objectives or individual.signature in seen_signatures:
            continue
        seen_objectives.add(objective_key)
        seen_signatures.add(individual.signature)
        unique.append(individual)
    if not unique:
        return []
    fronts = fast_non_dominated_sort(unique)
    pareto = fronts[0]
    calculate_crowding_distance(pareto)
    return sorted(pareto, key=lambda item: (item.objectives, item.signature))


def select_knee_solution(pareto: Sequence[Individual]) -> Individual:
    """以归一化目标到理想点的最小Chebyshev遗憾选择可复现折衷点。"""
    if not pareto:
        raise RuntimeError("最终种群中没有CV=0的Pareto解")
    values = np.array([individual.objectives for individual in pareto], dtype=float)
    minimum = values.min(axis=0)
    maximum = values.max(axis=0)
    ranges = maximum - minimum
    normalized = np.divide(
        values - minimum,
        ranges,
        out=np.zeros_like(values),
        where=ranges > TOLERANCE,
    )
    candidates = []
    for index, individual in enumerate(pareto):
        chebyshev_regret = float(normalized[index].max())
        euclidean_distance = float(np.linalg.norm(normalized[index]))
        candidates.append(
            (
                chebyshev_regret,
                euclidean_distance,
                individual.objectives,
                individual.signature,
                individual,
            )
        )
    return min(candidates, key=lambda item: item[:-1])[-1]


def final_feasibility_check(
    representative: Individual, data: core.ProblemData
) -> pd.DataFrame:
    if representative.evaluation is None:
        raise RuntimeError("代表方案尚未评价")
    checks = core.check_global_feasibility(representative.evaluation, data).copy()
    checks.columns = ["Check_No", "Constraint_Check", "Result"]
    overall = "FEASIBLE" if (checks["Result"] == "PASS").all() else "INFEASIBLE"
    checks = pd.concat(
        [
            checks,
            pd.DataFrame(
                [
                    {
                        "Check_No": "OVERALL",
                        "Constraint_Check": "Final solution status",
                        "Result": overall,
                    }
                ]
            ),
        ],
        ignore_index=True,
    )
    return checks


def pareto_tables(
    pareto: Sequence[Individual], representative: Individual
) -> Tuple[pd.DataFrame, pd.DataFrame, Dict[Tuple[Tuple, ...], str]]:
    solution_ids = {
        individual.signature: f"NSGA2-{index:03d}"
        for index, individual in enumerate(pareto, start=1)
    }
    front_rows = []
    solution_rows = []
    for individual in pareto:
        solution_id = solution_ids[individual.signature]
        front_rows.append(
            {
                "Solution_ID": solution_id,
                "Z1_Timeliness": individual.objectives[0],
                "Z2_Makespan_s": individual.objectives[1],
                "Z3_Energy_kWh": individual.objectives[2],
                "Z4_Sorties": int(individual.objectives[3]),
                "CV": individual.cv,
                "Is_Representative": individual is representative,
            }
        )
        for sortie_index, sortie in enumerate(individual.solution, start=1):
            solution_rows.append(
                {
                    "Solution_ID": solution_id,
                    "Plan_Sortie_Index": sortie_index,
                    "UAV_Type": sortie.uav_type,
                    "Route": " -> ".join((core.DEPOT, *sortie.route, core.DEPOT)),
                    "Box_Count": len(sortie.box_ids),
                    "Box_IDs": ";".join(sortie.box_ids),
                }
            )
    return pd.DataFrame(front_rows), pd.DataFrame(solution_rows), solution_ids


def representative_tables(
    representative: Individual, data: core.ProblemData
) -> Tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    evaluation = representative.evaluation
    if evaluation is None:
        raise RuntimeError("代表方案尚未评价")
    sortie_rows = []
    box_rows = []
    uav_rows = []
    battery_rows = []

    uav_previous_end = defaultdict(float)
    battery_tasks: Dict[str, List[core.ScheduledSortie]] = defaultdict(list)
    for item in evaluation.scheduled:
        battery_tasks[item.battery_id].append(item)
    next_battery_use: Dict[Tuple[str, str], Optional[float]] = {}
    for battery_id, tasks in battery_tasks.items():
        ordered = sorted(tasks, key=lambda item: item.start_s)
        for index, item in enumerate(ordered):
            next_battery_use[(battery_id, item.sortie_id)] = (
                ordered[index + 1].start_s if index + 1 < len(ordered) else None
            )

    for item in sorted(evaluation.scheduled, key=lambda row: row.sortie_id):
        sortie = item.sortie
        metrics = item.metrics
        sortie_rows.append(
            {
                "Sortie_ID": item.sortie_id,
                "UAV_Type": sortie.uav_type,
                "UAV_ID": item.uav_id,
                "Battery_ID": item.battery_id,
                "Start_Time_s": item.start_s,
                "Takeoff_Time_s": item.departure_s,
                "Return_Time_s": item.completion_s,
                "Route": " -> ".join((core.DEPOT, *sortie.route, core.DEPOT)),
                "Box_Count": len(sortie.box_ids),
                "Box_IDs": ";".join(sortie.box_ids),
                "Initial_Payload_kg": metrics.total_weight_kg,
                "Volume_m3": metrics.total_volume_m3,
                "Energy_kWh": metrics.energy_kwh,
                "Return_SOC": metrics.return_soc,
                "Recharge_Time_s": item.charge_time_s,
            }
        )
        idle_time = item.start_s - uav_previous_end[item.uav_id]
        uav_rows.append(
            {
                "UAV_ID": item.uav_id,
                "UAV_Type": sortie.uav_type,
                "Sortie_ID": item.sortie_id,
                "Start_Time_s": item.start_s,
                "Return_Time_s": item.completion_s,
                "Idle_Time_Before_s": idle_time,
                "Duration_s": item.completion_s - item.start_s,
            }
        )
        uav_previous_end[item.uav_id] = item.completion_s
        battery_rows.append(
            {
                "Battery_ID": item.battery_id,
                "UAV_Type": sortie.uav_type,
                "Sortie_ID": item.sortie_id,
                "Use_Start_s": item.start_s,
                "Return_Time_s": item.completion_s,
                "Return_SOC": metrics.return_soc,
                "Charge_Start_s": item.completion_s,
                "Charge_End_s": item.charge_end_s,
                "Next_Available_s": item.charge_end_s,
                "Next_Use_s": next_battery_use[(item.battery_id, item.sortie_id)],
            }
        )

        delivery_relative = metrics.delivery_map
        for box_id in sortie.box_ids:
            box = data.boxes[box_id]
            delivery_time = item.start_s + delivery_relative[box_id]
            hard_deadline = min(box.hard_deadlines) if box.hard_deadlines else None
            box_rows.append(
                {
                    "Box_ID": box_id,
                    "Service_Zone": box.service_zone,
                    "Material_Type": box.material_type,
                    "Sortie_ID": item.sortie_id,
                    "UAV_Type": sortie.uav_type,
                    "Delivery_Time_s": delivery_time,
                    "Expected_Time_s": box.expected_time_s,
                    "First_Batch_Deadline_s": box.first_deadline_s,
                    "Hard_Deadline_s": hard_deadline,
                    "Delay_s": max(0.0, delivery_time - box.expected_time_s),
                    "Emergency_Coefficient": box.emergency_priority,
                    "Hard_Deadline_Satisfied": (
                        True
                        if hard_deadline is None
                        else delivery_time <= hard_deadline + TOLERANCE
                    ),
                }
            )

    box_table = (
        pd.DataFrame(box_rows)
        .set_index("Box_ID")
        .loc[list(data.box_order)]
        .reset_index()
    )
    return (
        pd.DataFrame(sortie_rows),
        box_table,
        pd.DataFrame(uav_rows).sort_values(["UAV_ID", "Start_Time_s"]),
        pd.DataFrame(battery_rows).sort_values(["Battery_ID", "Use_Start_s"]),
    )


def format_workbook(writer: pd.ExcelWriter) -> None:
    for worksheet in writer.book.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cells in worksheet.iter_cols():
            values = ["" if cell.value is None else str(cell.value) for cell in cells]
            width = min(max(max(map(len, values)) + 2, 12), 70)
            worksheet.column_dimensions[cells[0].column_letter].width = width
            for cell in cells[1:]:
                if isinstance(cell.value, float):
                    cell.number_format = "0.000000000"


def export_excel(
    pareto: Sequence[Individual],
    representative: Individual,
    data: core.ProblemData,
    nsga_log: pd.DataFrame,
    feasibility: pd.DataFrame,
    runtime_s: float,
) -> None:
    front_table, pareto_solutions, _ = pareto_tables(pareto, representative)
    sorties, boxes, uav_usage, battery_usage = representative_tables(
        representative, data
    )
    parameters = pd.DataFrame(
        [
            {
                "population_size": CONFIG.population_size,
                "max_generations": CONFIG.max_generations,
                "crossover_probability": CONFIG.crossover_probability,
                "mutation_probability": (
                    CONFIG.mutation_probability
                    if CONFIG.mutation_probability is not None
                    else 1.0 / (3 * len(data.boxes))
                ),
                "eta_c": CONFIG.eta_c,
                "eta_m": CONFIG.eta_m,
                "beta": CONFIG.beta,
                "random_seed": CONFIG.random_seed,
                "runtime_s": runtime_s,
                "knee_rule": "min normalized Chebyshev regret; Euclidean distance tie-break",
                "delivery_time_definition": "completion of all handover at visited service zone",
                "soft_timeliness_definition": "weighted tardiness of boxes without hard deadlines",
            }
        ]
    )
    with pd.ExcelWriter(assert_writable_output(OUTPUT_FILE), engine="openpyxl") as writer:
        front_table.to_excel(writer, sheet_name="Pareto_Front", index=False)
        pareto_solutions.to_excel(writer, sheet_name="Pareto_Solutions", index=False)
        sorties.to_excel(writer, sheet_name="Representative_Sorties", index=False)
        boxes.to_excel(writer, sheet_name="Representative_Boxes", index=False)
        uav_usage.to_excel(writer, sheet_name="UAV_Usage", index=False)
        battery_usage.to_excel(writer, sheet_name="Battery_Usage", index=False)
        nsga_log.to_excel(writer, sheet_name="NSGA2_Log", index=False)
        feasibility.to_excel(writer, sheet_name="Feasibility_Check", index=False)
        parameters.to_excel(writer, sheet_name="Parameters", index=False)
        format_workbook(writer)


def update_recorded_runtime(runtime_s: float) -> None:
    workbook = load_workbook(OUTPUT_FILE)
    worksheet = workbook["Parameters"]
    headers = {cell.value: cell.column for cell in worksheet[1] if cell.value}
    worksheet.cell(row=2, column=headers["runtime_s"], value=runtime_s)
    workbook.save(assert_writable_output(OUTPUT_FILE))


def setup_chinese_font() -> None:
    candidates = ["Microsoft YaHei", "SimHei", "SimSun", "Noto Sans CJK SC"]
    installed = {font.name for font in font_manager.fontManager.ttflist}
    for name in candidates:
        if name in installed:
            plt.rcParams["font.sans-serif"] = [name]
            break
    plt.rcParams["axes.unicode_minus"] = False


def plot_convergence(nsga_log: pd.DataFrame) -> None:
    fig, axes = plt.subplots(3, 2, figsize=(12, 12), dpi=180)
    generation = nsga_log["generation"]
    axes[0, 0].plot(generation, nsga_log["feasible_count"], marker="o")
    axes[0, 0].set_title("可行个体数量")
    axes[0, 1].plot(generation, nsga_log["first_front_size"], marker="o", color="#ef8a17")
    axes[0, 1].set_title("第一非支配层规模")
    labels = ["min Z1", "min Z2", "min Z3", "min Z4"]
    columns = ["min_Z1", "min_Z2", "min_Z3", "min_Z4"]
    colors = ["#2f78b7", "#ef8a17", "#3a9d5d", "#a05aa1"]
    for ax, label, column, color in zip(axes.flat[2:], labels, columns, colors):
        ax.plot(generation, nsga_log[column], marker=".", color=color)
        ax.set_title(label)
    for ax in axes.flat:
        ax.set_xlabel("Generation")
        ax.grid(alpha=0.25)
    fig.suptitle("NSGA-II 多目标搜索收敛过程", fontsize=16)
    fig.tight_layout()
    fig.savefig(assert_writable_output(CONVERGENCE_FIGURE), bbox_inches="tight")
    plt.close(fig)


def plot_pareto_projections(
    pareto: Sequence[Individual], representative: Individual
) -> None:
    values = np.array([item.objectives for item in pareto], dtype=float)
    pairs = [(0, 1), (0, 2), (1, 2), (1, 3)]
    labels = ["Z1 加权延迟", "Z2 Cmax(s)", "Z3 能耗(kWh)", "Z4 架次数"]
    representative_values = representative.objectives
    fig, axes = plt.subplots(2, 2, figsize=(11, 9), dpi=180)
    for ax, (x_index, y_index) in zip(axes.flat, pairs):
        ax.scatter(values[:, x_index], values[:, y_index], s=38, alpha=0.75, label="Pareto解")
        ax.scatter(
            representative_values[x_index],
            representative_values[y_index],
            marker="*",
            s=220,
            color="red",
            edgecolor="black",
            label="代表性knee方案",
            zorder=5,
        )
        ax.set_xlabel(labels[x_index])
        ax.set_ylabel(labels[y_index])
        ax.grid(alpha=0.25)
        ax.legend()
    fig.suptitle("NSGA-II Pareto前沿二维投影", fontsize=16)
    fig.tight_layout()
    fig.savefig(assert_writable_output(PARETO_FIGURE), bbox_inches="tight")
    plt.close(fig)


def plot_parallel_coordinates(
    pareto: Sequence[Individual], representative: Individual
) -> None:
    values = np.array([item.objectives for item in pareto], dtype=float)
    minimum = values.min(axis=0)
    ranges = values.max(axis=0) - minimum
    normalized = np.divide(
        values - minimum,
        ranges,
        out=np.zeros_like(values),
        where=ranges > TOLERANCE,
    )
    representative_index = next(
        index for index, item in enumerate(pareto) if item is representative
    )
    fig, ax = plt.subplots(figsize=(9, 5.5), dpi=180)
    x = np.arange(4)
    for index, row in enumerate(normalized):
        if index == representative_index:
            continue
        ax.plot(x, row, color="#2f78b7", alpha=0.35)
    ax.plot(
        x,
        normalized[representative_index],
        color="red",
        linewidth=2.8,
        marker="*",
        markersize=11,
        label="代表性knee方案",
    )
    ax.set_xticks(x, ["Z1", "Z2", "Z3", "Z4"])
    ax.set_ylabel("Pareto集内归一化目标值")
    ax.set_title("NSGA-II 四目标Pareto平行坐标图")
    ax.grid(alpha=0.25)
    ax.legend()
    fig.tight_layout()
    fig.savefig(assert_writable_output(PARALLEL_FIGURE), bbox_inches="tight")
    plt.close(fig)


def plot_representative_routes(
    evaluation: core.Evaluation, data: core.ProblemData
) -> None:
    colors = {"A": "#2f78b7", "B": "#ef8a17", "C": "#3a9d5d"}
    fig, ax = plt.subplots(figsize=(9, 7), dpi=180)
    for item in evaluation.scheduled:
        nodes = (core.DEPOT, *item.sortie.route, core.DEPOT)
        longitude = [data.coordinates[node][0] for node in nodes]
        latitude = [data.coordinates[node][1] for node in nodes]
        ax.plot(longitude, latitude, color=colors[item.sortie.uav_type], alpha=0.60, linewidth=1.3)
    for node, (longitude, latitude) in data.coordinates.items():
        ax.scatter(
            longitude,
            latitude,
            marker="*" if node == core.DEPOT else "o",
            s=150 if node == core.DEPOT else 32,
            color="black",
            zorder=5,
        )
        ax.text(longitude + 0.0007, latitude + 0.0007, node, fontsize=8)
    for uav_type, color in colors.items():
        ax.plot([], [], color=color, label=f"{uav_type}型架次")
    ax.set_xlabel("经度（°）")
    ax.set_ylabel("纬度（°）")
    ax.set_title("NSGA-II代表方案运输路线")
    ax.grid(alpha=0.2)
    ax.legend()
    fig.tight_layout()
    fig.savefig(assert_writable_output(ROUTE_FIGURE), bbox_inches="tight")
    plt.close(fig)


def plot_uav_gantt(evaluation: core.Evaluation, data: core.ProblemData) -> None:
    colors = {"A": "#2f78b7", "B": "#ef8a17", "C": "#3a9d5d"}
    uav_ids = [uav for uav_type in core.UAV_TYPES for uav in data.uav_ids_by_type[uav_type]]
    positions = {uav: index for index, uav in enumerate(uav_ids)}
    fig, ax = plt.subplots(figsize=(11, 5.5), dpi=180)
    for item in evaluation.scheduled:
        y = positions[item.uav_id]
        ax.barh(y, item.completion_s - item.start_s, left=item.start_s, height=0.55, color=colors[item.sortie.uav_type], edgecolor="black", linewidth=0.4)
        ax.text((item.start_s + item.completion_s) / 2, y, item.sortie_id, ha="center", va="center", fontsize=7, color="white")
    ax.set_yticks(range(len(uav_ids)), uav_ids)
    ax.set_xlabel("时刻（s）")
    ax.set_ylabel("实体无人机")
    ax.set_title("NSGA-II代表方案实体无人机甘特图")
    ax.grid(axis="x", alpha=0.25)
    fig.tight_layout()
    fig.savefig(assert_writable_output(UAV_GANTT_FIGURE), bbox_inches="tight")
    plt.close(fig)


def plot_battery_gantt(evaluation: core.Evaluation, data: core.ProblemData) -> None:
    colors = {"A": "#2f78b7", "B": "#ef8a17", "C": "#3a9d5d"}
    battery_ids = [battery for uav_type in core.UAV_TYPES for battery in data.battery_ids_by_type[uav_type]]
    positions = {battery: index for index, battery in enumerate(battery_ids)}
    fig, ax = plt.subplots(figsize=(12, 8), dpi=180)
    for item in evaluation.scheduled:
        y = positions[item.battery_id]
        ax.barh(y, item.completion_s - item.start_s, left=item.start_s, height=0.55, color=colors[item.sortie.uav_type], edgecolor="black", linewidth=0.4)
        ax.barh(y, item.charge_time_s, left=item.completion_s, height=0.55, color="#bdbdbd", edgecolor="black", linewidth=0.4, hatch="//")
        ax.text((item.start_s + item.completion_s) / 2, y, item.sortie_id, ha="center", va="center", fontsize=6, color="white")
    ax.set_yticks(range(len(battery_ids)), battery_ids)
    ax.set_xlabel("时刻（s）")
    ax.set_ylabel("共享电池")
    ax.set_title("NSGA-II代表方案共享电池任务/充电甘特图")
    ax.grid(axis="x", alpha=0.25)
    fig.tight_layout()
    fig.savefig(assert_writable_output(BATTERY_GANTT_FIGURE), bbox_inches="tight")
    plt.close(fig)


def plot_delivery_comparison(evaluation: core.Evaluation, data: core.ProblemData) -> None:
    ordered = sorted(
        data.box_order,
        key=lambda box_id: (
            data.boxes[box_id].expected_time_s,
            data.boxes[box_id].service_zone,
            box_id,
        ),
    )
    x = np.arange(len(ordered))
    actual = [evaluation.delivery_times[box_id] for box_id in ordered]
    expected = [data.boxes[box_id].expected_time_s for box_id in ordered]
    hard = [
        min(data.boxes[box_id].hard_deadlines)
        if data.boxes[box_id].hard_deadlines
        else np.nan
        for box_id in ordered
    ]
    fig, ax = plt.subplots(figsize=(13, 5.5), dpi=180)
    ax.scatter(x, actual, s=18, label="实际送达", color="#2f78b7")
    ax.plot(x, expected, linewidth=1.2, label="期望送达", color="#d64b40")
    ax.scatter(x, hard, marker="x", s=28, label="有效硬截止", color="black")
    ax.set_xlabel("按期望时刻排序的货箱序号")
    ax.set_ylabel("时刻（s）")
    ax.set_title("NSGA-II代表方案逐箱送达时间对比")
    ax.grid(alpha=0.2)
    ax.legend()
    fig.tight_layout()
    fig.savefig(assert_writable_output(DELIVERY_FIGURE), bbox_inches="tight")
    plt.close(fig)


def plot_results(
    pareto: Sequence[Individual],
    representative: Individual,
    data: core.ProblemData,
    nsga_log: pd.DataFrame,
) -> None:
    setup_chinese_font()
    if representative.evaluation is None:
        raise RuntimeError("代表方案尚未评价")
    plot_convergence(nsga_log)
    plot_pareto_projections(pareto, representative)
    plot_parallel_coordinates(pareto, representative)
    plot_representative_routes(representative.evaluation, data)
    plot_uav_gantt(representative.evaluation, data)
    plot_battery_gantt(representative.evaluation, data)
    plot_delivery_comparison(representative.evaluation, data)


def main() -> None:
    start = perf_counter()
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    log("Constraint-handled Random-Key NSGA-II 开始", reset=True)
    mutation_probability_text = (
        str(CONFIG.mutation_probability)
        if CONFIG.mutation_probability is not None
        else "1/(3n)"
    )
    log(
        "parameters: population={}, generations={}, pc={}, pm={}, eta_c={}, eta_m={}, beta={}, seed={}".format(
            CONFIG.population_size,
            CONFIG.max_generations,
            CONFIG.crossover_probability,
            mutation_probability_text,
            CONFIG.eta_c,
            CONFIG.eta_m,
            CONFIG.beta,
            CONFIG.random_seed,
        )
    )

    data = load_data()
    preprocess_static_data(data)
    context = SearchContext(
        data=data,
        config=CONFIG,
        rng=random.Random(CONFIG.random_seed),
        box_ids=data.box_order,
        box_index={box_id: index for index, box_id in enumerate(data.box_order)},
    )
    population, nsga_log = nsga2(context)
    pareto = extract_pareto_front(population)
    representative = select_knee_solution(pareto)
    feasibility = final_feasibility_check(representative, data)
    if feasibility.iloc[-1]["Result"] != "FEASIBLE":
        raise RuntimeError("NSGA-II代表方案未通过最终完整可行性检查")

    runtime_before_output = perf_counter() - start
    export_excel(
        pareto,
        representative,
        data,
        nsga_log,
        feasibility,
        runtime_before_output,
    )
    plot_results(pareto, representative, data, nsga_log)
    total_runtime = perf_counter() - start
    update_recorded_runtime(total_runtime)

    objective_array = np.array([item.objectives for item in pareto], dtype=float)
    log(f"最终可行Pareto解数量: {len(pareto)}")
    log(
        "Pareto ranges: Z1=[{:.3f}, {:.3f}], Z2=[{:.3f}, {:.3f}], "
        "Z3=[{:.6f}, {:.6f}], Z4=[{:.0f}, {:.0f}]".format(
            objective_array[:, 0].min(),
            objective_array[:, 0].max(),
            objective_array[:, 1].min(),
            objective_array[:, 1].max(),
            objective_array[:, 2].min(),
            objective_array[:, 2].max(),
            objective_array[:, 3].min(),
            objective_array[:, 3].max(),
        )
    )
    log(
        "Representative knee: Z1={:.3f}, Z2={:.3f}, Z3={:.6f}, Z4={:.0f}, CV={:.12f}".format(
            *representative.objectives,
            representative.cv,
        )
    )
    log("最终状态: FEASIBLE")
    log(f"总运行时间: {total_runtime:.3f} s")
    log(f"结果文件: {OUTPUT_FILE}")


if __name__ == "__main__":
    main()
