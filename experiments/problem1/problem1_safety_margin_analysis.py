"""问题1：返航安全余量对最大安全载荷与货箱组批的敏感性分析。"""

from dataclasses import replace
from pathlib import Path
import sys

# Keep historical module names while resolving all locations from this repository.
_RELEASE_ROOT = Path(__file__).resolve().parents[2]
if str(_RELEASE_ROOT) not in sys.path:
    sys.path.insert(0, str(_RELEASE_ROOT))
from project_paths import ROOT, bootstrap, raw_file, dem_file, input_file, output_dir, output_file, require_file, assert_writable_output

bootstrap()
from time import perf_counter
from typing import Dict, List, Tuple

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
import numpy as np
from openpyxl import load_workbook
import pandas as pd

from problem1_max_safe_payload import solve_max_safe_payload
from problem1_batch_optimization import (
    BOX_FILE,
    CAPACITY_TOLERANCE,
    DEPOT,
    MAX_PAYLOAD_FILE,
    OBJECTIVE_TOLERANCE_KWH,
    SEGMENT_FILE,
    SERVICE_ZONES,
    TIME_TOLERANCE_S,
    UAV_FILE,
    UAV_TYPES,
    UAVParameters,
    build_set_partitioning_model,
    candidate_energy,
    extract_solution,
    format_workbook,
    generate_all_candidates,
    load_box_data,
    load_segment_data,
    load_uav_parameters,
    solve_lexicographic,
    to_segment_zone,
    validate_solution,
)


BASELINE_OPTIMIZATION_FILE = input_file("problem1", "problem1_batch_optimization.xlsx")
RESULT_DIR = output_dir("problem1")
OUTPUT_FILE = RESULT_DIR / "problem1_safety_margin_analysis.xlsx"
QMAX_FIGURE = RESULT_DIR / "问题1_返航安全余量与最大安全载荷.png"
SORTIE_FIGURE = RESULT_DIR / "问题1_返航安全余量与最优总架次数.png"
UAV_MIX_FIGURE = RESULT_DIR / "问题1_不同安全余量下机型架次构成.png"
CANDIDATE_FIGURE = RESULT_DIR / "问题1_返航安全余量与候选方案数量.png"

RHO_VALUES = [0.00, 0.05, 0.10, 0.15, 0.20, 0.25, 0.30]
BASELINE_RHO = 0.20
MONOTONIC_TOLERANCE_KG = 1e-7
CHANGE_COLUMNS = [
    "rho_from",
    "rho_to",
    "change_type",
    "service_zone",
    "description",
]


def scenario_uav_parameters(
    base_uavs: Dict[str, UAVParameters], rho: float
) -> Dict[str, UAVParameters]:
    """仅替换返航安全余量，其余机型及能耗参数保持不变。"""
    return {
        uav_type: replace(
            uav,
            energy_model=replace(uav.energy_model, reserve_ratio=float(rho)),
        )
        for uav_type, uav in base_uavs.items()
    }


def calculate_qmax_for_rho(
    rho: float,
    uavs: Dict[str, UAVParameters],
    matrices: Dict[str, pd.DataFrame],
) -> Tuple[pd.DataFrame, Dict[Tuple[str, str], float]]:
    """复用基准程序的二分求解函数计算一个 rho 下的 3×15 个 qmax。"""
    distance = matrices["distance_matrix"]
    climb = matrices["climb_height_matrix"]
    records = []
    payloads: Dict[Tuple[str, str], float] = {}

    for uav_type in UAV_TYPES:
        uav = uavs[uav_type]
        model = uav.energy_model
        energy_limit = (1.0 - rho) * model.usable_energy_kwh
        for service_zone in SERVICE_ZONES:
            destination = to_segment_zone(service_zone)
            (
                max_payload,
                feasible,
                status,
                empty_energy,
                full_energy,
            ) = solve_max_safe_payload(
                model,
                float(distance.loc[DEPOT, destination]),
                float(distance.loc[destination, DEPOT]),
                float(climb.loc[DEPOT, destination]),
                float(climb.loc[destination, DEPOT]),
            )
            if not feasible:
                raise RuntimeError(
                    f"rho={rho:.0%}, {uav_type}-{service_zone} 空载也无法安全往返"
                )
            payloads[(uav_type, service_zone)] = float(max_payload)
            max_payload_energy = candidate_energy(
                service_zone, max_payload, uav, matrices
            ).total_energy_kwh
            records.append(
                {
                    "rho": rho,
                    "rho_percent": rho * 100.0,
                    "service_zone": service_zone,
                    "uav_type": uav_type,
                    "max_safe_payload_kg": max_payload,
                    "payload_limit_kg": model.payload_limit_kg,
                    "energy_limit_kwh": energy_limit,
                    "energy_at_qmax_kwh": max_payload_energy,
                    "empty_round_trip_energy_kwh": empty_energy.total_energy_kwh,
                    "full_round_trip_energy_kwh": full_energy.total_energy_kwh,
                    "status": status,
                }
            )
    return pd.DataFrame(records), payloads


def build_qmax_summary(qmax_detail: pd.DataFrame) -> pd.DataFrame:
    """汇总各 rho 下三种机型 qmax 的均值、最小值和最大值。"""
    rows = []
    for rho in RHO_VALUES:
        scenario = qmax_detail.loc[np.isclose(qmax_detail["rho"], rho)]
        row = {"rho": rho, "rho_percent": rho * 100.0}
        for uav_type in UAV_TYPES:
            values = scenario.loc[
                scenario["uav_type"] == uav_type, "max_safe_payload_kg"
            ]
            row[f"{uav_type}_mean_qmax"] = float(values.mean())
            row[f"{uav_type}_min_qmax"] = float(values.min())
            row[f"{uav_type}_max_qmax"] = float(values.max())
        rows.append(row)
    summary = pd.DataFrame(rows)
    ordered_columns = ["rho", "rho_percent"]
    ordered_columns += [f"{uav_type}_mean_qmax" for uav_type in UAV_TYPES]
    ordered_columns += [f"{uav_type}_min_qmax" for uav_type in UAV_TYPES]
    ordered_columns += [f"{uav_type}_max_qmax" for uav_type in UAV_TYPES]
    return summary[ordered_columns]


def validate_qmax_monotonicity(qmax_detail: pd.DataFrame) -> None:
    """检查每个机型—服务区的 qmax 随 rho 单调不增。"""
    violations = []
    for (uav_type, service_zone), group in qmax_detail.groupby(
        ["uav_type", "service_zone"], sort=False
    ):
        ordered = group.sort_values("rho")
        payloads = ordered["max_safe_payload_kg"].to_numpy(dtype=float)
        differences = np.diff(payloads)
        for position in np.flatnonzero(differences > MONOTONIC_TOLERANCE_KG):
            violations.append(
                {
                    "uav_type": uav_type,
                    "service_zone": service_zone,
                    "rho_from": ordered.iloc[position]["rho"],
                    "rho_to": ordered.iloc[position + 1]["rho"],
                    "qmax_from": payloads[position],
                    "qmax_to": payloads[position + 1],
                }
            )
    if violations:
        print(pd.DataFrame(violations).to_string(index=False))
        raise RuntimeError("最大安全载荷未通过随 rho 单调不增检查")
    print("[PASS] All 45 qmax series are non-increasing with rho.")


def validate_baseline_qmax(qmax_detail: pd.DataFrame) -> None:
    """检查 rho=20% 的 qmax 与已有基准结果一致。"""
    existing = pd.read_excel(require_file(MAX_PAYLOAD_FILE), sheet_name="max_safe_payload")
    existing = existing.loc[existing["feasible"]].copy()
    existing["service_zone"] = existing["service_zone"].map(
        lambda value: f"S{int(str(value)[1:]):03d}"
    )
    expected = existing.set_index(["uav_type", "service_zone"])[
        "max_payload_kg"
    ].sort_index()
    actual = qmax_detail.loc[np.isclose(qmax_detail["rho"], BASELINE_RHO)].set_index(
        ["uav_type", "service_zone"]
    )["max_safe_payload_kg"].sort_index()
    if not np.allclose(actual, expected, atol=MONOTONIC_TOLERANCE_KG, rtol=1e-10):
        raise RuntimeError("rho=20% 的 qmax 与已有最大安全载荷结果不一致")
    print("[PASS] rho=20% qmax matches problem1_max_safe_payload.xlsx.")


def normalized_batches(selected: pd.DataFrame) -> Dict[Tuple[str, Tuple[str, ...]], str]:
    """以服务区和排序后的箱号集合为键，避免依赖 candidate_id。"""
    normalized = {}
    for row in selected.itertuples(index=False):
        box_ids = tuple(sorted(str(row.box_ids).split(";")))
        normalized[(row.service_zone, box_ids)] = row.uav_type
    return normalized


def append_change(
    rows: List[dict],
    rho_from: float,
    rho_to: float,
    change_type: str,
    service_zone: str,
    description: str,
) -> None:
    rows.append(
        {
            "rho_from": rho_from,
            "rho_to": rho_to,
            "change_type": change_type,
            "service_zone": service_zone,
            "description": description,
        }
    )


def compare_two_solutions(
    rho_from: float,
    rho_to: float,
    selected_from: pd.DataFrame,
    selected_to: pd.DataFrame,
) -> List[dict]:
    """比较两个情景的标准化批次、机型及各区架次数。"""
    changes: List[dict] = []
    if len(selected_from) != len(selected_to):
        append_change(
            changes,
            rho_from,
            rho_to,
            "sortie_count_change",
            "ALL",
            f"Minimum sorties: {len(selected_from)} -> {len(selected_to)}",
        )

    overall_from = selected_from["uav_type"].value_counts().reindex(
        UAV_TYPES, fill_value=0
    )
    overall_to = selected_to["uav_type"].value_counts().reindex(
        UAV_TYPES, fill_value=0
    )
    if not overall_from.equals(overall_to):
        append_change(
            changes,
            rho_from,
            rho_to,
            "uav_type_change",
            "ALL",
            "UAV sorties A/B/C: "
            f"{tuple(overall_from.tolist())} -> {tuple(overall_to.tolist())}",
        )

    old_map = normalized_batches(selected_from)
    new_map = normalized_batches(selected_to)
    for service_zone in SERVICE_ZONES:
        old_zone = selected_from.loc[selected_from["service_zone"] == service_zone]
        new_zone = selected_to.loc[selected_to["service_zone"] == service_zone]
        if len(old_zone) != len(new_zone):
            append_change(
                changes,
                rho_from,
                rho_to,
                "sortie_count_change",
                service_zone,
                f"Zone sorties: {len(old_zone)} -> {len(new_zone)}",
            )

        old_mix = old_zone["uav_type"].value_counts().reindex(UAV_TYPES, fill_value=0)
        new_mix = new_zone["uav_type"].value_counts().reindex(UAV_TYPES, fill_value=0)
        if not old_mix.equals(new_mix):
            append_change(
                changes,
                rho_from,
                rho_to,
                "uav_type_change",
                service_zone,
                "Zone UAV sorties A/B/C: "
                f"{tuple(old_mix.tolist())} -> {tuple(new_mix.tolist())}",
            )

        old_groups = {
            boxes: uav
            for (zone, boxes), uav in old_map.items()
            if zone == service_zone
        }
        new_groups = {
            boxes: uav
            for (zone, boxes), uav in new_map.items()
            if zone == service_zone
        }
        common_groups = set(old_groups).intersection(new_groups)
        for boxes in sorted(common_groups):
            if old_groups[boxes] != new_groups[boxes]:
                append_change(
                    changes,
                    rho_from,
                    rho_to,
                    "uav_type_change",
                    service_zone,
                    f"Batch {';'.join(boxes)}: {old_groups[boxes]} -> {new_groups[boxes]}",
                )

        removed = set(old_groups).difference(new_groups)
        added = set(new_groups).difference(old_groups)
        for boxes in sorted(removed):
            append_change(
                changes,
                rho_from,
                rho_to,
                "batch_removed",
                service_zone,
                f"Removed {old_groups[boxes]} batch: {';'.join(boxes)}",
            )
        for boxes in sorted(added):
            append_change(
                changes,
                rho_from,
                rho_to,
                "batch_added",
                service_zone,
                f"Added {new_groups[boxes]} batch: {';'.join(boxes)}",
            )

        old_sets = {boxes: set(boxes) for boxes in old_groups}
        new_sets = {boxes: set(boxes) for boxes in new_groups}
        for old_boxes in sorted(removed):
            old_set = old_sets[old_boxes]
            overlapping_new = [
                new_set for new_set in new_sets.values() if old_set.intersection(new_set)
            ]
            if (
                len(overlapping_new) > 1
                and all(new_set.issubset(old_set) for new_set in overlapping_new)
                and set().union(*overlapping_new) == old_set
            ):
                append_change(
                    changes,
                    rho_from,
                    rho_to,
                    "batch_split",
                    service_zone,
                    f"Former batch split: {';'.join(old_boxes)}",
                )
        for new_boxes in sorted(added):
            new_set = new_sets[new_boxes]
            overlapping_old = [
                old_set for old_set in old_sets.values() if new_set.intersection(old_set)
            ]
            if len(overlapping_old) > 1:
                append_change(
                    changes,
                    rho_from,
                    rho_to,
                    "batch_recombination",
                    service_zone,
                    f"New batch draws boxes from former batches: {';'.join(new_boxes)}",
                )
    return changes


def build_solution_changes(
    selected_by_rho: Dict[float, pd.DataFrame]
) -> pd.DataFrame:
    """记录所有相邻 rho 实验点之间的实际方案变化。"""
    rows = []
    for rho_from, rho_to in zip(RHO_VALUES[:-1], RHO_VALUES[1:]):
        rows.extend(
            compare_two_solutions(
                rho_from,
                rho_to,
                selected_by_rho[rho_from],
                selected_by_rho[rho_to],
            )
        )
    return pd.DataFrame(rows, columns=CHANGE_COLUMNS)


def build_baseline_changes(
    selected_by_rho: Dict[float, pd.DataFrame]
) -> pd.DataFrame:
    """记录每个非基准情景相对 rho=20% 基准方案的变化。"""
    rows = []
    baseline = selected_by_rho[BASELINE_RHO]
    for rho in RHO_VALUES:
        if np.isclose(rho, BASELINE_RHO):
            continue
        rows.extend(compare_two_solutions(BASELINE_RHO, rho, baseline, selected_by_rho[rho]))
    return pd.DataFrame(rows, columns=CHANGE_COLUMNS)


def add_baseline_deltas(optimization_summary: pd.DataFrame) -> pd.DataFrame:
    """增加相对 20% 基准情形的目标值与机型架次变化。"""
    summary = optimization_summary.copy()
    baseline = summary.loc[np.isclose(summary["rho"], BASELINE_RHO)].iloc[0]
    summary["delta_sorties_vs_20pct"] = (
        summary["minimum_sorties"] - baseline["minimum_sorties"]
    )
    summary["delta_energy_vs_20pct"] = (
        summary["total_energy_kwh"] - baseline["total_energy_kwh"]
    )
    summary["delta_operation_time_vs_20pct"] = (
        summary["total_operation_time_s"] - baseline["total_operation_time_s"]
    )
    for uav_type in UAV_TYPES:
        summary[f"delta_{uav_type}_sorties_vs_20pct"] = (
            summary[f"{uav_type}_sorties"] - baseline[f"{uav_type}_sorties"]
        )
    summary["uav_structure_changed_vs_20pct"] = (
        summary[[f"delta_{uav_type}_sorties_vs_20pct" for uav_type in UAV_TYPES]]
        .ne(0)
        .any(axis=1)
    )
    return summary


def validate_baseline_optimization(optimization_summary: pd.DataFrame) -> None:
    """检查 20% 情景与已有基准组批目标值一致。"""
    existing = pd.read_excel(
        require_file(BASELINE_OPTIMIZATION_FILE), sheet_name="optimization_summary"
    ).iloc[0]
    baseline = optimization_summary.loc[
        np.isclose(optimization_summary["rho"], BASELINE_RHO)
    ].iloc[0]
    checks = [
        baseline["candidate_count"] == existing["candidate_count"],
        baseline["minimum_sorties"] == existing["minimum_sorties"],
        np.isclose(
            baseline["total_energy_kwh"],
            existing["minimum_energy_kwh"],
            atol=OBJECTIVE_TOLERANCE_KWH,
            rtol=1e-12,
        ),
        np.isclose(
            baseline["total_operation_time_s"],
            existing["minimum_operation_time_s"],
            atol=TIME_TOLERANCE_S,
            rtol=1e-12,
        ),
    ]
    if not all(checks):
        raise RuntimeError("rho=20% 的组批目标值与已有基准优化结果不一致")
    print("[PASS] rho=20% optimization matches problem1_batch_optimization.xlsx.")


def setup_chinese_font() -> None:
    """设置 matplotlib 中文字体。"""
    candidates = ["Microsoft YaHei", "SimHei", "SimSun", "Noto Sans CJK SC"]
    installed = {font.name for font in font_manager.fontManager.ttflist}
    for name in candidates:
        if name in installed:
            plt.rcParams["font.sans-serif"] = [name]
            break
    plt.rcParams["axes.unicode_minus"] = False


def add_baseline_reference(ax) -> None:
    ax.axvline(BASELINE_RHO * 100.0, color="#555555", linestyle="--", linewidth=1.2)
    ax.text(
        BASELINE_RHO * 100.0,
        0.98,
        "20% 基准",
        transform=ax.get_xaxis_transform(),
        ha="right",
        va="top",
        fontsize=9,
        color="#444444",
    )


def plot_sensitivity_results(
    qmax_summary: pd.DataFrame, optimization_summary: pd.DataFrame
) -> None:
    """绘制最大安全载荷、架次数、机型构成和候选数四张图。"""
    setup_chinese_font()
    colors = {"A": "#0072B2", "B": "#E69F00", "C": "#009E73"}
    x = optimization_summary["rho_percent"].to_numpy(dtype=float)

    fig, ax = plt.subplots(figsize=(8.2, 5.4), dpi=300)
    for uav_type in UAV_TYPES:
        ax.plot(
            qmax_summary["rho_percent"],
            qmax_summary[f"{uav_type}_mean_qmax"],
            marker="o",
            linewidth=2.0,
            color=colors[uav_type],
            label=f"{uav_type}型",
        )
    add_baseline_reference(ax)
    ax.set_xlabel("返航安全余量 ρ (%)")
    ax.set_ylabel("平均最大安全载荷 (kg)")
    # ax.set_title("返航安全余量与平均最大安全载荷")
    ax.grid(alpha=0.25)
    ax.legend()
    fig.tight_layout()
    fig.savefig(assert_writable_output(QMAX_FIGURE), bbox_inches="tight")
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(8.2, 5.4), dpi=300)
    ax.step(
        x,
        optimization_summary["minimum_sorties"],
        where="mid",
        marker="o",
        linewidth=2.0,
        color="#CC3311",
    )
    add_baseline_reference(ax)
    ax.set_xticks(x)
    ax.set_yticks(
        sorted(optimization_summary["minimum_sorties"].astype(int).unique())
    )
    ax.set_xlabel("返航安全余量 ρ (%)")
    ax.set_ylabel("最优总往返架次数")
    # ax.set_title("返航安全余量与最优总架次数")
    ax.grid(alpha=0.25)
    fig.tight_layout()
    fig.savefig(assert_writable_output(SORTIE_FIGURE), bbox_inches="tight")
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(8.2, 5.4), dpi=300)
    bottom = np.zeros(len(optimization_summary), dtype=float)
    for uav_type in UAV_TYPES:
        values = optimization_summary[f"{uav_type}_sorties"].to_numpy(dtype=float)
        ax.bar(
            x,
            values,
            bottom=bottom,
            width=3.4,
            color=colors[uav_type],
            label=f"{uav_type}型",
        )
        bottom += values
    add_baseline_reference(ax)
    ax.set_xticks(x)
    ax.set_xlabel("返航安全余量 ρ (%)")
    ax.set_ylabel("架次数")
    # ax.set_title("不同返航安全余量下的机型架次构成")
    ax.legend()
    ax.grid(axis="y", alpha=0.25)
    fig.tight_layout()
    fig.savefig(assert_writable_output(UAV_MIX_FIGURE), bbox_inches="tight")
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(8.2, 5.4), dpi=300)
    ax.plot(
        x,
        optimization_summary["candidate_count"],
        marker="o",
        linewidth=2.0,
        color="#6A3D9A",
    )
    add_baseline_reference(ax)
    ax.set_xticks(x)
    ax.set_xlabel("返航安全余量 ρ (%)")
    ax.set_ylabel("可行候选运输方案数量")
    # ax.set_title("返航安全余量与候选方案数量")
    ax.grid(alpha=0.25)
    fig.tight_layout()
    fig.savefig(assert_writable_output(CANDIDATE_FIGURE), bbox_inches="tight")
    plt.close(fig)


def save_results(
    qmax_detail: pd.DataFrame,
    qmax_summary: pd.DataFrame,
    optimization_summary: pd.DataFrame,
    selected_batches: pd.DataFrame,
    zone_summary: pd.DataFrame,
    solution_changes: pd.DataFrame,
    baseline_changes: pd.DataFrame,
) -> None:
    """保存敏感性分析的明细、汇总、批次与变化记录。"""
    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(assert_writable_output(OUTPUT_FILE), engine="openpyxl") as writer:
        qmax_detail.to_excel(writer, sheet_name="qmax_detail", index=False)
        qmax_summary.to_excel(writer, sheet_name="qmax_summary", index=False)
        optimization_summary.to_excel(
            writer, sheet_name="optimization_summary", index=False
        )
        selected_batches.to_excel(writer, sheet_name="selected_batches", index=False)
        zone_summary.to_excel(writer, sheet_name="zone_summary", index=False)
        solution_changes.to_excel(writer, sheet_name="solution_changes", index=False)
        baseline_changes.to_excel(
            writer, sheet_name="baseline_solution_changes", index=False
        )
        format_workbook(writer)


def save_analysis_summary(runtime_s: float) -> None:
    """在主结果首次写出后记录完整分析运行时间。"""
    workbook = load_workbook(OUTPUT_FILE)
    worksheet = workbook.create_sheet("analysis_summary")
    worksheet.append(
        ["rho_scenario_count", "baseline_rho", "total_runtime_s", "solver"]
    )
    worksheet.append([len(RHO_VALUES), BASELINE_RHO, runtime_s, "Gurobi"])
    worksheet.freeze_panes = "A2"
    for column in worksheet.columns:
        worksheet.column_dimensions[column[0].column_letter].width = 22
    worksheet.cell(row=2, column=3).number_format = "0.000000000"
    workbook.save(assert_writable_output(OUTPUT_FILE))


def main() -> None:
    total_start = perf_counter()
    boxes = load_box_data(BOX_FILE)
    base_uavs = load_uav_parameters(UAV_FILE)
    matrices = load_segment_data(SEGMENT_FILE)

    print(f"货箱文件: {BOX_FILE} [逐箱货箱清单]")
    print(f"无人机参数文件: {UAV_FILE} [数据]")
    print(
        f"航段文件: {SEGMENT_FILE} "
        "[distance_matrix, climb_height_matrix, descent_height_matrix]"
    )
    print(f"基准最大安全载荷文件: {MAX_PAYLOAD_FILE} [max_safe_payload]")
    print(f"基准组批结果文件: {BASELINE_OPTIMIZATION_FILE} [optimization_summary]")

    qmax_frames = []
    optimization_rows = []
    selected_frames = []
    zone_frames = []
    selected_by_rho: Dict[float, pd.DataFrame] = {}

    for rho in RHO_VALUES:
        scenario_start = perf_counter()
        is_baseline = np.isclose(rho, BASELINE_RHO)
        print("\n" + "=" * 60)
        print(
            f"Safety margin rho = {rho:.0%}"
            + (" [BASELINE]" if is_baseline else "")
        )
        print("=" * 60)

        uavs = scenario_uav_parameters(base_uavs, rho)
        qmax_detail, max_payloads = calculate_qmax_for_rho(
            rho, uavs, matrices
        )
        qmax_frames.append(qmax_detail)
        candidates = generate_all_candidates(boxes, uavs, max_payloads, matrices)
        (
            model,
            variables,
            sortie_expression,
            energy_expression,
            time_expression,
        ) = build_set_partitioning_model(boxes, candidates)
        result = solve_lexicographic(
            model,
            variables,
            sortie_expression,
            energy_expression,
            time_expression,
            len(candidates),
        )
        selected, assignment, service_summary = extract_solution(
            boxes, candidates, result
        )
        validate_solution(boxes, selected, assignment, uavs, matrices, result)

        scenario_runtime = perf_counter() - scenario_start
        type_counts = selected["uav_type"].value_counts().reindex(
            UAV_TYPES, fill_value=0
        )
        total_energy = float(selected["energy_kwh"].sum())
        total_time = float(selected["operation_time_s"].sum())
        print(f"\nCandidate count: {len(candidates)}")
        print(f"Optimal sorties: {result.minimum_sorties}")
        for uav_type in UAV_TYPES:
            print(f"{uav_type} sorties: {int(type_counts[uav_type])}")
        print(f"Total energy: {total_energy:.12f} kWh")
        print(f"Total operation time: {total_time:.9f} s")
        print(f"Scenario runtime: {scenario_runtime:.6f} s")

        optimization_rows.append(
            {
                "rho": rho,
                "rho_percent": rho * 100.0,
                "is_baseline": bool(is_baseline),
                "candidate_count": len(candidates),
                "minimum_sorties": result.minimum_sorties,
                "A_sorties": int(type_counts["A"]),
                "B_sorties": int(type_counts["B"]),
                "C_sorties": int(type_counts["C"]),
                "total_energy_kwh": total_energy,
                "total_operation_time_s": total_time,
                "runtime_s": scenario_runtime,
                "stage1_runtime_s": result.stage1_runtime_s,
                "stage2_runtime_s": result.stage2_runtime_s,
                "stage3_runtime_s": result.stage3_runtime_s,
            }
        )

        selected = selected.copy()
        selected.insert(0, "rho_percent", rho * 100.0)
        selected.insert(0, "rho", rho)
        selected_frames.append(selected.drop(columns="box_ids_tuple"))
        selected_by_rho[rho] = selected

        service_summary = service_summary.copy()
        service_summary.insert(0, "rho_percent", rho * 100.0)
        service_summary.insert(0, "rho", rho)
        zone_frames.append(service_summary)
        model.dispose()

    qmax_detail = pd.concat(qmax_frames, ignore_index=True)
    qmax_summary = build_qmax_summary(qmax_detail)
    optimization_summary = add_baseline_deltas(pd.DataFrame(optimization_rows))
    selected_batches = pd.concat(selected_frames, ignore_index=True)
    zone_summary = pd.concat(zone_frames, ignore_index=True)
    solution_changes = build_solution_changes(selected_by_rho)
    baseline_changes = build_baseline_changes(selected_by_rho)

    print("\n跨情景一致性检查:")
    validate_qmax_monotonicity(qmax_detail)
    validate_baseline_qmax(qmax_detail)
    validate_baseline_optimization(optimization_summary)

    for row in optimization_summary.itertuples(index=False):
        print(
            f"rho={row.rho:.0%}: candidates={row.candidate_count}, "
            f"Z1={row.minimum_sorties}, Z2={row.total_energy_kwh:.9f} kWh, "
            f"Z3={row.total_operation_time_s:.6f} s, "
            f"A/B/C={row.A_sorties}/{row.B_sorties}/{row.C_sorties}"
        )

    sortie_changes = solution_changes.loc[
        (solution_changes["change_type"] == "sortie_count_change")
        & (solution_changes["service_zone"] == "ALL")
    ]
    for row in sortie_changes.itertuples(index=False):
        print("\n[CHANGE]")
        print(f"Safety margin: {row.rho_from:.0%} -> {row.rho_to:.0%}")
        print(row.description)

    plot_sensitivity_results(qmax_summary, optimization_summary)
    save_results(
        qmax_detail,
        qmax_summary,
        optimization_summary,
        selected_batches,
        zone_summary,
        solution_changes,
        baseline_changes,
    )
    total_runtime = perf_counter() - total_start
    save_analysis_summary(total_runtime)
    print(f"\n总运行时间: {total_runtime:.6f} s")
    print(f"Excel 已保存: {OUTPUT_FILE}")
    for figure in (QMAX_FIGURE, SORTIE_FIGURE, UAV_MIX_FIGURE, CANDIDATE_FIGURE):
        print(f"图片已保存: {figure}")


if __name__ == "__main__":
    main()
