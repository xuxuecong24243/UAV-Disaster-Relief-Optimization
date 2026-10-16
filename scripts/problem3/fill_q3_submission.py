"""用Problem 3最终VNS解填写官方结果提交模板的两个Q3工作表。

本脚本不执行优化。它从problem3_vns_20cycles.xlsx恢复最终运输基因，
调用既有确定性decoder和标准通信链路函数生成连续通信保障区间，再用
openpyxl仅填写Q3_中继架次与Q3_通信保障。
"""

from __future__ import annotations

from collections import Counter, defaultdict
from copy import copy
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, List, Mapping, Sequence, Tuple
import hashlib
import math
import sys

import pandas as pd
from openpyxl import load_workbook


# Preserve historical module names used by trusted communication caches.
_PROJECT_ROOT = Path(__file__).resolve().parents[2]
if str(_PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(_PROJECT_ROOT))
import project_paths as paths
paths.bootstrap()
MODULE_DIR = Path(__file__).resolve().parent
PROJECT_DIR = paths.ROOT

import problem3_alns as alns  # noqa: E402
import problem3_baseline as baseline  # noqa: E402
import problem3_communication as communication  # noqa: E402
import problem3_core as core  # noqa: E402
import problem3_relay as relay  # noqa: E402


TEMPLATE_FILE = paths.template_file()
VNS_FILE = paths.input_file("problem3", "problem3_vns_20cycles.xlsx")
OUTPUT_FILE = paths.output_file("problem3", "结果提交模板_问题三已填写.xlsx")
TEMP_OUTPUT_FILE = paths.output_file("problem3", "结果提交模板_问题三已填写.tmp.xlsx")

RELAY_SHEET = "Q3_中继架次"
COMMUNICATION_SHEET = "Q3_通信保障"
TARGET_SHEETS = {RELAY_SHEET, COMMUNICATION_SHEET}
WEIGHTS = (0.55, 0.15, 0.15, 0.15)
EXPECTED_OBJECTIVES = (0.0, 6985.310067776591, 68.77109563851582, 23.0)
EXPECTED_OBJECTIVE = 0.4499192949212766
NUMERIC_TOLERANCE = 1e-6
INTERVAL_TOLERANCE_S = communication.STATE_TRANSITION_TOLERANCE_S

EXPECTED_RELAY_HEADERS = (
    "中继架次编号",
    "中继无人机编号",
    "能源组件编号",
    "开始时刻（s）",
    "悬停经度（°）",
    "悬停纬度（°）",
    "悬停海拔（m）",
    "建链完成时刻（s）",
    "服务结束时刻（s）",
    "返回O01时刻（s）",
    "架次能耗（kWh）",
)
EXPECTED_COMMUNICATION_HEADERS = (
    "运输架次编号",
    "通信阶段",
    "开始时刻（s）",
    "结束时刻（s）",
    "保障方式",
    "中继架次编号",
)


@dataclass(frozen=True)
class CommunicationSource:
    method: str
    relay_sortie_id: str = ""


@dataclass(frozen=True)
class CommunicationInterval:
    transport_sortie_id: str
    phase: str
    start_time_s: float
    end_time_s: float
    source: CommunicationSource

    @property
    def duration_s(self) -> float:
        return self.end_time_s - self.start_time_s


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def sheet_fingerprint(sheet) -> Tuple[object, ...]:
    """比较非目标Sheet的值、公式、样式、尺寸和合并区域。"""
    cells = tuple(
        (
            cell.coordinate,
            cell.value,
            cell.data_type,
            cell.style_id,
            cell.number_format,
            cell.hyperlink.target if cell.hyperlink else None,
            cell.comment.text if cell.comment else None,
        )
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None or cell.has_style or cell.hyperlink or cell.comment
    )
    columns = tuple(
        sorted(
            (
                key,
                value.width,
                value.hidden,
                value.outlineLevel,
                value.bestFit,
            )
            for key, value in sheet.column_dimensions.items()
        )
    )
    rows = tuple(
        sorted(
            (
                key,
                value.height,
                value.hidden,
                value.outlineLevel,
            )
            for key, value in sheet.row_dimensions.items()
        )
    )
    return (
        sheet.title,
        sheet.sheet_state,
        sheet.max_row,
        sheet.max_column,
        tuple(str(value) for value in sheet.merged_cells.ranges),
        str(sheet.freeze_panes),
        sheet.auto_filter.ref,
        str(sheet.print_area),
        columns,
        rows,
        cells,
    )


def validate_vns_workbook() -> Mapping[str, object]:
    summary = pd.read_excel(paths.require_file(VNS_FILE), sheet_name="Summary")
    values = dict(zip(summary["Metric"], summary["Value"]))
    required = {
        "Completed_Cycles": 20,
        "Final_Transport_Sorties": 20,
        "Final_Relay_Sorties": 3,
        "Final_Uncovered_Communication_Duration_s": 0.0,
        "LEVEL2_Executed": True,
        "LEVEL2_Passed": True,
    }
    for key, expected in required.items():
        if key not in values or values[key] != expected:
            raise AssertionError(f"VNS Summary不一致: {key}={values.get(key)!r}")
    for key, expected in zip(
        ("Final_Z1", "Final_Z2_s", "Final_Z3_kWh", "Final_Z4"),
        EXPECTED_OBJECTIVES,
    ):
        if not math.isclose(float(values[key]), expected, abs_tol=NUMERIC_TOLERANCE):
            raise AssertionError(f"VNS目标不一致: {key}={values[key]!r}")
    if not math.isclose(
        float(values["Final_Formal_F"]), EXPECTED_OBJECTIVE, abs_tol=1e-12
    ):
        raise AssertionError("VNS综合目标不一致")
    feasibility = pd.read_excel(paths.require_file(VNS_FILE), sheet_name="Final_Feasibility_Check")
    if len(feasibility) != 13 or not feasibility["Result"].eq("PASS").all():
        raise AssertionError("VNS LEVEL 2可行性不是13/13 PASS")
    return values


def reconstruct_final_solution() -> alns.Problem3Solution:
    table = pd.read_excel(paths.require_file(VNS_FILE), sheet_name="Transport_Sorties").sort_values(
        "Schedule_Order"
    )
    if len(table) != 20:
        raise AssertionError(f"最终运输架次数不是20: {len(table)}")
    sorties = tuple(
        alns.p2_core.Sortie(
            box_ids=tuple(str(row.Box_IDs).split(";")),
            route=core._parse_route(str(row.Route)),
            uav_type=str(row.UAV_Type),
        )
        for row in table.itertuples(index=False)
    )
    return alns.Problem3Solution(sorties)


def assert_close(actual: float, expected: float, label: str) -> None:
    if not math.isclose(float(actual), float(expected), abs_tol=NUMERIC_TOLERANCE):
        raise AssertionError(f"{label}不一致: decoded={actual}, workbook={expected}")


def validate_decoded_against_workbook(decoded: alns.Problem3DecodedState) -> None:
    assert decoded.feasible
    assert len(decoded.transport_evaluation.scheduled) == 20
    assert len(decoded.relay_sorties) == 3
    for index, (actual, expected) in enumerate(
        zip(decoded.objectives, EXPECTED_OBJECTIVES), start=1
    ):
        assert_close(actual, expected, f"Z{index}")
    assert_close(decoded.objective, EXPECTED_OBJECTIVE, "F")

    transport = pd.read_excel(paths.require_file(VNS_FILE), sheet_name="Transport_Sorties")
    workbook_transport = {str(row.Sortie_ID): row for row in transport.itertuples()}
    for scheduled in decoded.transport_evaluation.scheduled:
        row = workbook_transport[scheduled.sortie_id]
        if scheduled.uav_id != str(row.UAV_ID) or scheduled.battery_id != str(row.Battery_ID):
            raise AssertionError(f"{scheduled.sortie_id}运输资源调度不一致")
        assert_close(scheduled.start_s, row.Start_Time_s, "运输开始时刻")
        assert_close(scheduled.completion_s, row.Return_Time_s, "运输返回时刻")

    relay_table = pd.read_excel(paths.require_file(VNS_FILE), sheet_name="Relay_Sorties")
    workbook_relay = {str(row.Sortie_ID): row for row in relay_table.itertuples()}
    for scheduled in decoded.relay_sorties:
        row = workbook_relay[scheduled.sortie_id]
        checks = (
            (scheduled.relay_uav_id, str(row.Relay_UAV_ID), "中继无人机"),
            (scheduled.energy_component_id, str(row.Energy_Component_ID), "能源组件"),
            (scheduled.position_id, str(row.Position_ID), "中继位置"),
        )
        if any(actual != expected for actual, expected, _ in checks):
            raise AssertionError(f"{scheduled.sortie_id}中继资源/位置不一致")
        for actual, expected, label in (
            (scheduled.prepare_start_s, row.Prepare_Start_s, "准备开始"),
            (scheduled.setup_end_s, row.Setup_End_s, "建链完成"),
            (scheduled.service_start_s, row.Service_Start_s, "服务开始"),
            (scheduled.service_end_s, row.Service_End_s, "服务结束"),
            (scheduled.return_end_s, row.Return_End_s, "返回O01"),
            (scheduled.total_energy_kwh, row.Total_Energy_kWh, "架次能耗"),
        ):
            assert_close(actual, expected, f"{scheduled.sortie_id} {label}")


class StandardSourceEvaluator:
    """完全复用正式UG/UR/RG链路函数的通信来源判定。"""

    def __init__(
        self,
        decoded: alns.Problem3DecodedState,
        context: alns.PhaseAContext,
    ) -> None:
        self.decoded = decoded
        self.context = context
        self.positions = {row.position_id: row for row in decoded.relay_positions}
        self.rg = {
            position_id: communication.evaluate_rg_link(
                position.endpoint,
                context.gateway.point,
                context.communication_parameters,
                context.dem,
            )
            for position_id, position in self.positions.items()
        }
        if not all(value.available for value in self.rg.values()):
            raise AssertionError("最终中继位置存在RG不可用")
        self.ug_cache: Dict[Tuple[float, float, float], object] = {}
        self.ur_cache: Dict[Tuple[str, float, float, float], object] = {}

    @staticmethod
    def point_key(point) -> Tuple[float, float, float]:
        return (
            round(point.longitude_deg, 10),
            round(point.latitude_deg, 10),
            round(point.altitude_m, 6),
        )

    def evaluate(
        self, trajectory: core.TransportTrajectory, time_s: float
    ) -> Tuple[CommunicationSource, bool, bool, bool]:
        point = trajectory.phase_at(time_s).position_at(time_s)
        point_key = self.point_key(point)
        if point_key not in self.ug_cache:
            self.ug_cache[point_key] = communication.evaluate_ug_link(
                point,
                self.context.gateway.point,
                self.context.communication_parameters,
                self.context.dem,
            )
        ug = self.ug_cache[point_key]
        if ug.available:
            return CommunicationSource("G01直连"), True, False, False
        active = sorted(
            (
                row
                for row in self.decoded.relay_sorties
                if row.service_start_s - core.TIME_TOLERANCE_S
                <= time_s
                <= row.service_end_s + core.TIME_TOLERANCE_S
            ),
            key=lambda row: (row.relay_uav_id, row.sortie_id),
        )
        for relay_sortie in active:
            key = (relay_sortie.position_id, *point_key)
            if key not in self.ur_cache:
                self.ur_cache[key] = communication.evaluate_ur_link(
                    point,
                    self.positions[relay_sortie.position_id].endpoint,
                    self.context.communication_parameters,
                    self.context.dem,
                )
            ur = self.ur_cache[key]
            rg = self.rg[relay_sortie.position_id]
            if ur.available and rg.available:
                return (
                    CommunicationSource("中继通信", relay_sortie.sortie_id),
                    False,
                    True,
                    True,
                )
        return CommunicationSource("UNAVAILABLE"), False, False, False


def timeline_source(row) -> CommunicationSource:
    if str(row.Communication_Source) == "G01":
        return CommunicationSource("G01直连")
    relay_id = str(row.Relay_Sortie_ID)
    if relay_id and relay_id.lower() != "nan":
        return CommunicationSource("中继通信", relay_id)
    return CommunicationSource("UNAVAILABLE")


def refine_source_transition(
    evaluator: StandardSourceEvaluator,
    trajectory: core.TransportTrajectory,
    left_s: float,
    right_s: float,
    left_source: CommunicationSource,
) -> float:
    while right_s - left_s > INTERVAL_TOLERANCE_S:
        midpoint = (left_s + right_s) / 2.0
        source, _, _, _ = evaluator.evaluate(trajectory, midpoint)
        if source == left_source:
            left_s = midpoint
        else:
            right_s = midpoint
    return (left_s + right_s) / 2.0


def append_interval(
    rows: List[CommunicationInterval],
    sortie_id: str,
    phase: str,
    start_s: float,
    end_s: float,
    source: CommunicationSource,
) -> None:
    if end_s - start_s <= 1e-10:
        return
    if source.method == "UNAVAILABLE":
        raise AssertionError(
            f"通信不可用: sortie={sortie_id}, phase={phase}, [{start_s}, {end_s}]"
        )
    if (
        rows
        and rows[-1].transport_sortie_id == sortie_id
        and rows[-1].phase == phase
        and rows[-1].source == source
        and abs(rows[-1].end_time_s - start_s) <= 1e-8
    ):
        previous = rows[-1]
        rows[-1] = CommunicationInterval(
            previous.transport_sortie_id,
            previous.phase,
            previous.start_time_s,
            end_s,
            previous.source,
        )
    else:
        rows.append(
            CommunicationInterval(sortie_id, phase, start_s, end_s, source)
        )


def build_communication_intervals(
    decoded: alns.Problem3DecodedState,
    context: alns.PhaseAContext,
) -> Tuple[Tuple[CommunicationInterval, ...], baseline.CommunicationCheckResult]:
    """标准LEVEL 2复核后，将采样时间轴压缩为连续来源区间。"""
    positions = {row.position_id: row for row in decoded.relay_positions}
    exact = baseline.verify_continuous_communication(
        decoded.trajectories,
        decoded.demands,
        decoded.relay_sorties,
        positions,
        context.communication_parameters,
        context.gateway,
        context.dem,
    )
    if exact.uncovered_duration_s != 0.0 or not exact.all_samples_available:
        raise AssertionError(
            f"标准通信复核失败: uncovered={exact.uncovered_duration_s}"
        )
    evaluator = StandardSourceEvaluator(decoded, context)
    service_boundaries = {
        round(value, 9)
        for row in decoded.relay_sorties
        for value in (row.service_start_s, row.service_end_s)
    }
    intervals: List[CommunicationInterval] = []
    for trajectory in sorted(
        decoded.trajectories, key=lambda row: row.record.schedule_order
    ):
        timeline = exact.timeline.loc[
            exact.timeline["Transport_Sortie_ID"].eq(trajectory.record.sortie_id)
        ].sort_values("Time_s")
        source_by_time = {
            round(float(row.Time_s), 9): timeline_source(row)
            for row in timeline.itertuples(index=False)
        }
        timeline_times = tuple(float(value) for value in timeline["Time_s"])
        for phase in trajectory.communication_phases:
            if phase.duration_s <= core.TIME_TOLERANCE_S:
                continue
            times = sorted(
                {
                    phase.start_time_s,
                    phase.end_time_s,
                    *(
                        value
                        for value in timeline_times
                        if phase.start_time_s < value < phase.end_time_s
                    ),
                }
            )
            phase_rows: List[CommunicationInterval] = []
            for left_s, right_s in zip(times, times[1:]):
                width = right_s - left_s
                left_source = source_by_time.get(round(left_s, 9))
                right_source = source_by_time.get(round(right_s, 9))
                epsilon = min(1e-4, width / 4.0)
                if round(left_s, 9) in service_boundaries:
                    left_source = evaluator.evaluate(
                        trajectory, left_s + epsilon
                    )[0]
                if round(right_s, 9) in service_boundaries:
                    right_source = evaluator.evaluate(
                        trajectory, right_s - epsilon
                    )[0]
                if left_source is None:
                    left_source = evaluator.evaluate(trajectory, left_s)[0]
                if right_source is None:
                    right_source = evaluator.evaluate(trajectory, right_s)[0]
                if left_source == right_source:
                    append_interval(
                        phase_rows,
                        trajectory.record.sortie_id,
                        phase.phase,
                        left_s,
                        right_s,
                        left_source,
                    )
                    continue
                transition = refine_source_transition(
                    evaluator,
                    trajectory,
                    left_s,
                    right_s,
                    left_source,
                )
                midpoint_source = evaluator.evaluate(trajectory, transition)[0]
                if midpoint_source not in (left_source, right_source):
                    raise AssertionError("单个1s区间内出现超过一次通信来源切换")
                append_interval(
                    phase_rows,
                    trajectory.record.sortie_id,
                    phase.phase,
                    left_s,
                    transition,
                    left_source,
                )
                append_interval(
                    phase_rows,
                    trajectory.record.sortie_id,
                    phase.phase,
                    transition,
                    right_s,
                    right_source,
                )
            if not phase_rows:
                raise AssertionError(f"非零通信阶段未生成区间: {phase.phase_id}")
            if not math.isclose(
                phase_rows[0].start_time_s,
                phase.start_time_s,
                abs_tol=1e-8,
            ) or not math.isclose(
                phase_rows[-1].end_time_s,
                phase.end_time_s,
                abs_tol=1e-8,
            ):
                raise AssertionError(f"阶段端点覆盖不完整: {phase.phase_id}")
            for first, second in zip(phase_rows, phase_rows[1:]):
                if not math.isclose(
                    first.end_time_s, second.start_time_s, abs_tol=1e-8
                ):
                    raise AssertionError(f"阶段存在通信空洞/重叠: {phase.phase_id}")
            intervals.extend(phase_rows)
    return tuple(intervals), exact


def audit_intervals(
    intervals: Sequence[CommunicationInterval],
    decoded: alns.Problem3DecodedState,
    context: alns.PhaseAContext,
) -> None:
    by_sortie_phase: Dict[Tuple[str, float, float], List[CommunicationInterval]] = (
        defaultdict(list)
    )
    trajectories = {row.record.sortie_id: row for row in decoded.trajectories}
    for trajectory in decoded.trajectories:
        for phase in trajectory.communication_phases:
            if phase.duration_s > core.TIME_TOLERANCE_S:
                key = (
                    trajectory.record.sortie_id,
                    phase.start_time_s,
                    phase.end_time_s,
                )
                by_sortie_phase[key] = []
    for row in intervals:
        trajectory = trajectories[row.transport_sortie_id]
        phase = trajectory.phase_at((row.start_time_s + row.end_time_s) / 2.0)
        key = (row.transport_sortie_id, phase.start_time_s, phase.end_time_s)
        by_sortie_phase[key].append(row)

    expected_duration = 0.0
    actual_duration = 0.0
    for (sortie_id, start_s, end_s), rows in by_sortie_phase.items():
        if not rows:
            raise AssertionError(f"阶段没有通信记录: {sortie_id} [{start_s},{end_s}]")
        rows.sort(key=lambda value: value.start_time_s)
        if abs(rows[0].start_time_s - start_s) > 1e-8:
            raise AssertionError("通信记录阶段起点空洞")
        if abs(rows[-1].end_time_s - end_s) > 1e-8:
            raise AssertionError("通信记录阶段终点空洞")
        for first, second in zip(rows, rows[1:]):
            if abs(first.end_time_s - second.start_time_s) > 1e-8:
                raise AssertionError("通信记录存在空洞或重叠")
        expected_duration += end_s - start_s
        actual_duration += sum(row.duration_s for row in rows)
    assert_close(actual_duration, expected_duration, "通信记录总覆盖时长")

    relay_map = {row.sortie_id: row for row in decoded.relay_sorties}
    evaluator = StandardSourceEvaluator(decoded, context)
    for row in intervals:
        midpoint = (row.start_time_s + row.end_time_s) / 2.0
        trajectory = trajectories[row.transport_sortie_id]
        source, ug, ur, rg = evaluator.evaluate(trajectory, midpoint)
        if source != row.source:
            raise AssertionError(f"区间中点通信来源不一致: {row}")
        if row.source.method == "G01直连":
            if not ug or row.source.relay_sortie_id:
                raise AssertionError("G01直连区间链路审计失败")
        else:
            relay_sortie = relay_map[row.source.relay_sortie_id]
            if not (
                relay_sortie.service_start_s - INTERVAL_TOLERANCE_S
                <= midpoint
                <= relay_sortie.service_end_s + INTERVAL_TOLERANCE_S
                and ur
                and rg
            ):
                raise AssertionError("中继通信区间SERVICE/UR/RG审计失败")


def copy_row_style(sheet, source_row: int, target_row: int) -> None:
    if source_row == target_row:
        return
    for column in range(1, sheet.max_column + 1):
        source = sheet.cell(source_row, column)
        target = sheet.cell(target_row, column)
        if source.has_style:
            target._style = copy(source._style)
        if source.number_format:
            target.number_format = source.number_format
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)
    source_dimension = sheet.row_dimensions[source_row]
    target_dimension = sheet.row_dimensions[target_row]
    target_dimension.height = source_dimension.height
    target_dimension.hidden = source_dimension.hidden


def clear_values(sheet, first_column: int, last_column: int) -> None:
    for row in range(2, sheet.max_row + 1):
        for column in range(first_column, last_column + 1):
            sheet.cell(row, column).value = None


def fill_template(
    decoded: alns.Problem3DecodedState,
    intervals: Sequence[CommunicationInterval],
) -> bool:
    paths.require_file(TEMPLATE_FILE)
    paths.ensure_output_path(TEMP_OUTPUT_FILE)
    paths.ensure_output_path(OUTPUT_FILE)
    template_hash_before = file_sha256(TEMPLATE_FILE)
    workbook = load_workbook(TEMPLATE_FILE, data_only=False)
    if tuple(workbook.sheetnames) != (
        "Q1_单点组批",
        "Q2_运输架次",
        "Q2_逐箱交付",
        RELAY_SHEET,
        COMMUNICATION_SHEET,
        "Q4_分区配置",
    ):
        raise AssertionError("官方模板Sheet名称或顺序不一致")
    other_before = {
        name: sheet_fingerprint(workbook[name])
        for name in workbook.sheetnames
        if name not in TARGET_SHEETS
    }
    relay_sheet = workbook[RELAY_SHEET]
    communication_sheet = workbook[COMMUNICATION_SHEET]
    if tuple(relay_sheet.cell(1, column).value for column in range(1, 12)) != EXPECTED_RELAY_HEADERS:
        raise AssertionError("Q3_中继架次真实表头与预期不一致")
    if tuple(
        communication_sheet.cell(1, column).value for column in range(1, 7)
    ) != EXPECTED_COMMUNICATION_HEADERS:
        raise AssertionError("Q3_通信保障真实表头与预期不一致")

    clear_values(relay_sheet, 1, 11)
    positions = {row.position_id: row for row in decoded.relay_positions}
    for row_index, sortie in enumerate(decoded.relay_sorties, start=2):
        position = positions[sortie.position_id]
        values = (
            sortie.sortie_id,
            sortie.relay_uav_id,
            sortie.energy_component_id,
            sortie.prepare_start_s,
            position.longitude_deg,
            position.latitude_deg,
            position.hover_absolute_altitude_m,
            sortie.setup_end_s,
            sortie.service_end_s,
            sortie.return_end_s,
            sortie.total_energy_kwh,
        )
        for column, value in enumerate(values, start=1):
            relay_sheet.cell(row_index, column).value = value

    original_comm_max_row = communication_sheet.max_row
    clear_values(communication_sheet, 1, 6)
    for row_index, interval in enumerate(intervals, start=2):
        if row_index > original_comm_max_row:
            copy_row_style(
                communication_sheet, original_comm_max_row, row_index
            )
        values = (
            interval.transport_sortie_id,
            interval.phase,
            interval.start_time_s,
            interval.end_time_s,
            interval.source.method,
            interval.source.relay_sortie_id or None,
        )
        for column, value in enumerate(values, start=1):
            communication_sheet.cell(row_index, column).value = value

    workbook.save(TEMP_OUTPUT_FILE)
    reloaded = load_workbook(TEMP_OUTPUT_FILE, data_only=False)
    other_unchanged = all(
        sheet_fingerprint(reloaded[name]) == signature
        for name, signature in other_before.items()
    )
    if not other_unchanged:
        raise AssertionError("保存后发现非Q3 Sheet的值或格式发生变化")
    if file_sha256(TEMPLATE_FILE) != template_hash_before:
        raise AssertionError("原始结果提交模板被修改")
    TEMP_OUTPUT_FILE.replace(OUTPUT_FILE)
    return other_unchanged


def main() -> None:
    paths.require_file(TEMPLATE_FILE)
    validate_vns_workbook()
    solution = reconstruct_final_solution()
    context = alns.load_phase_a_context()
    geometry_cache = alns.GeometryCommunicationCache(context)
    baseline_demands, _ = relay.load_step3_demands()
    elite_pool = alns.load_saved_elite_pool(
        context,
        geometry_cache,
        baseline_demands,
        alns.ELITE_POOL_FILE,
    )
    if not alns.attach_persistent_segment_cache_if_available(
        context, geometry_cache, elite_pool
    ):
        raise RuntimeError("持久化通信缓存不存在")
    result = alns.evaluate_problem3_solution(
        solution,
        context,
        geometry_cache,
        elite_pool,
        WEIGHTS,
        verify_exact=False,
    )
    if not result.feasible or result.decoded_state is None:
        raise RuntimeError(f"最终VNS解确定性重建失败: {result.infeasibility_reason}")
    decoded = result.decoded_state
    validate_decoded_against_workbook(decoded)
    intervals, exact = build_communication_intervals(decoded, context)
    audit_intervals(intervals, decoded, context)
    other_unchanged = fill_template(decoded, intervals)

    method_counts = Counter(row.source.method for row in intervals)
    phase_counts = Counter(row.phase for row in intervals)
    covered_sorties = {row.transport_sortie_id for row in intervals}
    relay_ids = sorted(
        {
            row.source.relay_sortie_id
            for row in intervals
            if row.source.relay_sortie_id
        }
    )
    print("========== Q3 SUBMISSION AUDIT ==========", flush=True)
    print("Q3_中继架次:", flush=True)
    print(f"relay_sorties_written = {len(decoded.relay_sorties)}", flush=True)
    print("expected = 3", flush=True)
    print("", flush=True)
    print("Q3_通信保障:", flush=True)
    print(
        f"transport_sorties_covered = {len(covered_sorties)} / 20", flush=True
    )
    print(f"communication_intervals_written = {len(intervals)}", flush=True)
    print(f"direct_intervals = {method_counts['G01直连']}", flush=True)
    print(f"relay_intervals = {method_counts['中继通信']}", flush=True)
    print("", flush=True)
    print("relay IDs used:", flush=True)
    print(", ".join(relay_ids), flush=True)
    print("", flush=True)
    print("communication phases:", flush=True)
    for phase in (core.CLIMB, core.CRUISE, core.DESCENT, core.HANDOVER):
        print(f"{phase} = {phase_counts[phase]}", flush=True)
    print("", flush=True)
    print(f"uncovered_duration = {exact.uncovered_duration_s:.9f}", flush=True)
    print(f"template_other_sheets_unchanged = {other_unchanged}", flush=True)
    print("", flush=True)
    print("Final consistency:", flush=True)
    print("transport sorties = 20", flush=True)
    print("relay sorties = 3", flush=True)
    print(f"Z1 = {decoded.objectives[0]:.12f}", flush=True)
    print(f"Z2 = {decoded.objectives[1]:.12f}", flush=True)
    print(f"Z3 = {decoded.objectives[2]:.12f}", flush=True)
    print(f"Z4 = {decoded.objectives[3]:.0f}", flush=True)
    print("LEVEL2 = PASS", flush=True)
    print("feasibility = 13/13", flush=True)
    print("=========================================", flush=True)
    print(f"output = {OUTPUT_FILE}", flush=True)


if __name__ == "__main__":
    main()
